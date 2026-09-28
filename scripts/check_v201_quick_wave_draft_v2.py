from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "RTDS_V2.01快速仿真波形确认稿.docx"
doc = Document(DOCX)
texts=[]
for p in doc.paragraphs:
    texts.append(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            texts.extend(p.text for p in cell.paragraphs)
full="\n".join(texts)

for forbidden in ("VCD波形重绘图","结论与边界","封存","发布","消费","。。","&#x20;"):
    assert forbidden not in full, forbidden
for number in range(1,18):
    assert f"3.{number} " in full, number
assert len(doc.inline_shapes)==33, len(doc.inline_shapes)
expected=int(17.2/2.54*914400)
assert all(abs(s.width-expected)<5000 for s in doc.inline_shapes)
assert "63×32ns=2.016µs" in full
assert "0x180FF8" in full and "AWLEN=255" in full and "AWLEN=6" in full
assert "WSTRB=E" in full and "CPU_WR_DONE_EVENT" in full

with ZipFile(DOCX) as zf:
    assert zf.testzip() is None
    xml=etree.fromstring(zf.read("word/document.xml"))
    rel=etree.fromstring(zf.read("word/_rels/document.xml.rels"))
    ns={"a":"http://schemas.openxmlformats.org/drawingml/2006/main",
        "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "pr":"http://schemas.openxmlformats.org/package/2006/relationships"}
    ids=xml.xpath(".//a:blip/@r:embed",namespaces=ns)
    assert len(ids)==33
    image_rels={x.get("Id"):x.get("Target") for x in rel.xpath(".//pr:Relationship",namespaces=ns) if x.get("Type","").endswith("/image")}
    assert len({image_rels[x] for x in ids})==33
    toc=" ".join(xml.xpath(".//w:instrText/text()",namespaces=ns))
    assert "TOC" in toc

print(f"DOCX_V2_STRUCTURE_PASS figures={len(doc.inline_shapes)} tables={len(doc.tables)} sections={len(doc.sections)}")
