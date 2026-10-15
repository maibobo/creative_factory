from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Xilinx_FPGA工程知识库使用指南.docx"
MANIFEST = ROOT / "index" / "manifest.json"
REPORT = ROOT / "index" / "build_report.json"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
MARGIN_DXA = 1440
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
CAUTION = "7A5A00"
CAUTION_FILL = "FFF8E1"
RISK = "9B1C1C"
RISK_FILL = "FDECEC"
POSITIVE = "1F3A5F"
WHITE = "FFFFFF"
BLACK = "000000"


def set_run_font(run, *, latin="Calibri", east_asia="Microsoft YaHei", size=None,
                 bold=None, italic=None, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, *, latin="Calibri", east_asia="Microsoft YaHei", size=11,
                   bold=False, color=BLACK):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_keep_next(style):
    ppr = style.element.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_style_font(title, size=30, bold=True, color=INK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, size=14, color=DARK_BLUE)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        set_style_font(style, size=size, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        set_keep_next(style)

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(code, latin="Consolas", east_asia="DengXian", size=9.2, color=INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.08)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.05

    small = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(small, size=9.5, color=MUTED)
    small.paragraph_format.space_before = Pt(2)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.15

    table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(table_text, size=9.5, color=BLACK)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.15

    header_text = styles.add_style("Table Header Text", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(header_text, size=9.5, bold=True, color=INK)
    header_text.paragraph_format.space_before = Pt(0)
    header_text.paragraph_format.space_after = Pt(0)
    header_text.paragraph_format.line_spacing = 1.1

    caption = styles["Caption"]
    set_style_font(caption, size=9.5, bold=True, color=MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_field(paragraph, instruction: str, display: str = ""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Xilinx / FPGA 工程知识库")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("使用指南  |  第 ")
    set_run_font(r, size=8.5, color=MUTED)
    page_run = add_field(p, "PAGE", "1")
    set_run_font(page_run, size=8.5, color=MUTED)
    r = p.add_run(" 页")
    set_run_font(r, size=8.5, color=MUTED)


def add_page_numbering_restart(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def add_numbering(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)

    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        if kind == "bullet":
            lvl_text.set(qn("w:val"), ["•", "–", "○"][level])
        else:
            lvl_text.set(qn("w:val"), f"%{level + 1}.")
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(540 + level * 540))
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(540 + level * 540))
        ind.set(qn("w:hanging"), "270")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), "Calibri")
        rfonts.set(qn("w:hAnsi"), "Calibri")
        rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        rpr.append(rfonts)
        lvl.append(rpr)
        abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text: str, num_id: int, level=0, *, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    ppr.append(num_pr)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def shade_paragraph(paragraph, fill, border_color=None):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border_color:
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "16")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        pbdr.append(left)


def add_callout(doc, label: str, text: str, *, kind="note"):
    palette = {
        "note": (CALLOUT, BLUE),
        "caution": (CAUTION_FILL, CAUTION),
        "risk": (RISK_FILL, RISK),
        "positive": (LIGHT_BLUE, POSITIVE),
    }
    fill, accent = palette[kind]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.06)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, fill, accent)
    r = p.add_run(label + "：")
    set_run_font(r, bold=True, color=accent)
    r = p.add_run(text)
    set_run_font(r, color=BLACK)
    return p


def add_code(doc, text: str):
    p = doc.add_paragraph(style="Code Block")
    shade_paragraph(p, LIGHT_GRAY, BLUE)
    lines = text.splitlines() or [""]
    for i, line in enumerate(lines):
        r = p.add_run(line)
        set_run_font(r, latin="Consolas", east_asia="DengXian", size=9.2, color=INK)
        if i != len(lines) - 1:
            r.add_break()
    return p


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS.items():
        tag = "w:start" if side == "start" else "w:end" if side == "end" else f"w:{side}"
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_text(cell, text, *, header=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = "Table Header Text" if header else "Table Text"
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_run_font(r, size=9.5, bold=header, color=INK if header else BLACK)


def add_table(doc, headers, rows, widths, *, aligns=None, caption=None):
    if caption:
        doc.add_paragraph(caption, style="Caption")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_header(table.rows[0])
    for i, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], text, header=True,
                      align=(aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT))
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            set_cell_text(cells[i], value, align=(aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT))
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, *, bold_lead=None, style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_cover(doc: Document, stats):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(54)
    r = p.add_run("本地工程参考手册")
    set_run_font(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    r = p.add_run("Xilinx / FPGA 工程知识库\n使用指南")
    set_run_font(r, size=30, bold=True, color=INK)

    p = doc.add_paragraph(style="Subtitle")
    r = p.add_run("范围说明、检索方法、证据规则与维护操作")
    set_run_font(r, size=14, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("给第一次使用知识库的人")
    set_run_font(r, size=11, italic=True, color=MUTED)

    add_table(
        doc,
        ["项目", "当前值"],
        [
            ["知识库目录", str(ROOT)],
            ["去重条目", f"{stats['documents']} 个"],
            ["全文字符", f"{stats['chars']:,} 字符"],
            ["构建状态", f"抽取错误 {stats['errors']}；缺失来源 {stats['missing']}"],
            ["指南日期", date.today().isoformat()],
        ],
        [2700, 6660],
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_callout(doc, "使用前先记住", "知识库给出的是可追溯证据，不会自动替你判断所有文档都已验证。看到“PASS”“已实现”或“已有 bitstream”时，仍要确认它属于哪一个验证层级。", kind="caution")
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "目录", 1)
    p = doc.add_paragraph(style="Small Note")
    r = p.add_run("在 Microsoft Word 中打开后，如果页码没有自动更新，请选中目录并按 F9。")
    set_run_font(r, size=9.5, color=MUTED)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "目录将在 Word 中更新")
    p.paragraph_format.space_after = Pt(10)
    doc.add_page_break()


def build_document():
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    report = json.loads(REPORT.read_text(encoding="utf-8"))
    by_category = {}
    for entry in manifest:
        info = by_category.setdefault(entry["category"], {"documents": 0, "chars": 0})
        info["documents"] += 1
        info["chars"] += int(entry.get("chars", 0))
    stats = {
        "documents": len(manifest),
        "chars": sum(int(x.get("chars", 0)) for x in manifest),
        "errors": len(report.get("errors", [])),
        "missing": len(report.get("missing_sources", [])),
        "duplicates": len(report.get("duplicates", [])),
        "candidates": int(report.get("candidate_count", 0)),
    }

    doc = Document()
    section = doc.sections[0]
    configure_page(section)
    configure_header_footer(section)
    configure_styles(doc)
    bullet_num = add_numbering(doc, "bullet")
    decimal_num = add_numbering(doc, "decimal")

    props = doc.core_properties
    props.title = "Xilinx / FPGA 工程知识库使用指南"
    props.subject = "知识库范围、检索、证据规则与维护操作"
    props.author = "Xilinx 工程知识库"
    props.keywords = "Xilinx, FPGA, Vivado, Aurora, PCIe, XDMA, DDR4, MIG, Ethernet, RGMII"
    props.comments = "由本地知识库清单与主题页生成"

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    add_cover(doc, stats)
    add_toc(doc)

    add_heading(doc, "1. 一分钟上手", 1)
    add_para(doc, "知识库已经放在本机目录 `E:\\centOS\\xilinx_knowledge_base`。日常使用不需要打开数据库，也不需要写 Python。打开 PowerShell，进入目录，然后运行搜索命令即可。")
    add_heading(doc, "1.1 第一次搜索", 2)
    add_code(doc, "cd E:\\centOS\\xilinx_knowledge_base\n.\\search.ps1 'Aurora lane_up channel_up'")
    add_para(doc, "结果会列出命中的文档标题、分类、原始文件路径、知识库文本副本路径和命中片段。先看主题速查页，再回到原始文档核对细节。")
    add_heading(doc, "1.2 只查一个专题", 2)
    add_code(doc, ".\\search.ps1 '4KB 边界' --category pcie\n.\\search.ps1 'init_calib_complete' --category ddr --limit 15\n.\\search.ps1 '风暴窗口 DROP_DUP_CRC' --category ethernet")
    add_heading(doc, "1.3 看当前规模", 2)
    add_code(doc, ".\\search.ps1 --stats")
    add_callout(doc, "推荐习惯", "查询词不要写成完整问题。优先使用信号名、寄存器名、错误码、IP 名和关键行为，例如 `usr_irq_ack CQ sink`、`ui_clk_sync_rst`、`rx_dst_rdy overflow`。", kind="positive")

    add_heading(doc, "2. 这个知识库是什么", 1)
    add_para(doc, "它是一个完全本地、可重建的 Xilinx/FPGA 技术资料库。原始文件仍留在各自工程目录；知识库负责抽取可检索文本、按 SHA-256 去重、记录来源，并生成 SQLite 全文索引。")
    add_list_item(doc, "官方资料用于确认协议、IP 参数和工具方法。", bullet_num)
    add_list_item(doc, "已验证工程用于确认实际配置、脚本、约束和验证结果。", bullet_num)
    add_list_item(doc, "当前 RTL、测试平台和工程说明用于定位实现行为。", bullet_num)
    add_list_item(doc, "人工主题页把容易混淆的结论、验证边界和使用入口整理成速查说明。", bullet_num)
    add_callout(doc, "它不是什么", "它不是 Vivado 工程替代品，也不是自动判定所有资料正确的问答系统。检索结果可能同时命中官方指南、旧设计稿和当前 RTL；最终判断要遵守证据优先级。", kind="caution")

    add_heading(doc, "3. 当前收录规模", 1)
    cat_labels = {
        "foundation": "通用基础",
        "aurora": "Aurora",
        "pcie": "PCIe / XDMA / NVMe",
        "ddr": "DDR / MIG",
        "ethernet": "以太网 / lwIP",
        "cross_ip": "跨 IP 集成",
        "guide": "人工主题指南",
    }
    cat_desc = {
        "foundation": "Vivado、Tcl、综合、仿真、时序、AXI、AXI-Stream、FIFO、CDC、时钟、DMA",
        "aurora": "64B/66B、8B/10B、LocalLink、GT/复位、接口适配、验证工程",
        "pcie": "PCIe Gen3、XDMA、Root Port、H2C/C2H、MSI、FDMA、NVMe 数据通路",
        "ddr": "DDR4 MIG、AXI BIST、校准、复位、实现与板级验证边界",
        "ethernet": "RGMII/GMII、帧解析、CRC、MAC/ACL、风暴过滤、分类 DMA、lwIP",
        "cross_ip": "Aurora/PCIe/DDR/NVMe 之间的 CDC、FIFO、回压、描述符和集成验证",
        "guide": "范围说明、四大专题速查和跨 IP 检查表",
    }
    rows = []
    order = ["foundation", "aurora", "pcie", "ddr", "ethernet", "cross_ip", "guide"]
    for key in order:
        item = by_category[key]
        rows.append([cat_labels[key], key, str(item["documents"]), cat_desc[key]])
    add_table(doc, ["中文范围", "分类参数", "条目", "主要内容"], rows,
              [1700, 1450, 900, 5310],
              aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                      WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
              caption="表 1  当前分类与内容范围")
    add_para(doc, f"本轮从 {stats['candidates']} 个候选文件中形成 {stats['documents']} 个去重条目，合并 {stats['duplicates']} 个重复来源。全文约 {stats['chars']:,} 字符，构建错误 {stats['errors']}，缺失来源 {stats['missing']}。")

    add_heading(doc, "4. 明确不收录的内容", 1)
    add_para(doc, "知识库按白名单组织，不是把整个磁盘扫进数据库。下面这些内容已经排除：")
    exclusions = [
        "HDMI、DisplayPort、USB、MIPI、摄像头和其他显示接口。",
        "图像/视频处理，以及 SIFT、CNN、YOLO 等算法或工程。",
        "Vivado 大体量生成目录，例如 `.Xil`、`.cache`、`.runs`、`.sim`、`.gen`、`ip_user_files`。",
        "波形、日志、渲染页和仅有图片、无法纯文本审计的中间截图。",
        "商业报价、采购执行包等与技术设计无关的材料。",
    ]
    for text in exclusions:
        add_list_item(doc, text, bullet_num)
    add_callout(doc, "加密文件", "两个密码保护的 Aurora Office 文件只保留路径、SHA-256 和重复来源，正文未解密，也不作为技术结论。获得密码或未加密副本后，再运行增量构建即可。", kind="note")

    add_heading(doc, "5. 资料来自哪里", 1)
    add_table(doc, ["位置", "用途"], [
        [r"C:\Users\MaiBo\Documents\XilinxDocs", "用户指定的 XilinxDocs 主目录；当前精选时钟类文档。"],
        [r"F:\Xiliinx文档", "Vivado、AXI、FIFO、DMA、时序和 lwIP 等官方资料。"],
        ["R:\\", "远端映射盘上的 Aurora、DDR/MIG、PCIe/XDMA、NVMe/DDR 已验证或当前工程。"],
        [r"E:\aurora_proc、E:\ddr_proc", "本机 Aurora 与 DDR 工程、脚本、RTL、约束和阶段性资料。"],
        [r"E:\PCIE-CPU、E:\NVME_PRO", "PCIe/XDMA、FDMA/Aurora、NVMe/DDR 集成设计和回归材料。"],
        [r"E:\eth_storm_project", "RGMII 风暴过滤、分类 FIFO、DataMover DMA 和验证规划。"],
        [r"F:\lwip_echo_server", "精选 lwIP 应用层源码。"],
    ], [3100, 6260], caption="表 2  主要来源目录")
    add_para(doc, "完整白名单、显式文件和排除规则位于：")
    add_code(doc, r"E:\centOS\xilinx_knowledge_base\config\sources.json")

    add_heading(doc, "6. 目录结构怎么读", 1)
    add_table(doc, ["目录或文件", "作用"], [
        ["README.md", "最短入口，列出主题页和常用命令。"],
        ["knowledge/", "人工整理的结论型主题页。适合先读，能快速识别验证状态和关键风险。"],
        ["corpus/", "原始 PDF、DOC/DOCX、XLSX、Markdown、RTL、Tcl、XDC、XCI 等抽取出的文本副本。"],
        ["index/xilinx_kb.sqlite3", "SQLite FTS5 全文索引；日常通过 search.ps1 使用，不必手工打开。"],
        ["index/manifest.json / .csv", "条目来源、SHA-256、修改时间、分类、抽取器、别名和文本路径。"],
        ["index/build_report.json", "最近一次构建统计、重复来源、错误和缺失来源。"],
        ["config/sources.json", "来源白名单、分类、标签、文件大小限制和排除规则。"],
        ["cache/extracted/", "按文件哈希缓存抽取结果，加快增量重建。"],
        ["tools/", "构建、检索和本指南生成脚本。"],
        ["build.ps1 / search.ps1", "用户直接运行的 PowerShell 入口。"],
    ], [3000, 6360], caption="表 3  知识库目录说明")

    add_heading(doc, "7. 检索命令完整说明", 1)
    add_heading(doc, "7.1 基本格式", 2)
    add_code(doc, ".\\search.ps1 '查询词' [--category 分类] [--limit 数量] [--json]")
    add_table(doc, ["参数", "含义", "示例"], [
        ["查询词", "中文、英文、信号名、路径片段、错误码均可。", "'usr_irq_ack CQ sink'"],
        ["--category", "只查一个分类。", "--category pcie"],
        ["--limit", "限制返回条数。", "--limit 15"],
        ["--stats", "显示各分类条目数与字符数，不需要查询词。", ".\\search.ps1 --stats"],
        ["--json", "以 JSON 输出，便于脚本继续处理。", ".\\search.ps1 'channel_up' --json"],
        ["--help", "显示命令帮助。", ".\\search.ps1 --help"],
    ], [1700, 4200, 3460], caption="表 4  search.ps1 参数")
    add_heading(doc, "7.2 可用分类", 2)
    add_code(doc, "foundation  aurora  pcie  ddr  ethernet  cross_ip  guide")
    add_para(doc, "如果不知道应该选哪一类，先不加 `--category`。看一次结果后，再用分类收窄。")
    add_heading(doc, "7.3 查询词怎么写", 2)
    query_tips = [
        "优先写 2～5 个辨识度高的词，例如 `Aurora rx_dst_rdy overflow`。",
        "协议问题把接口名和行为放在一起，例如 `AW W independent handshake`。",
        "找验证结果时加入 `PASS`、`未执行`、`errors=0`、`VALIDATION_SUMMARY`。",
        "找 RTL 时直接用模块名、状态名、端口名或错误编码。",
        "中文描述搜不到时，换成源码里的英文信号名；反过来也一样。",
    ]
    for text in query_tips:
        add_list_item(doc, text, bullet_num)
    add_callout(doc, "不要一上来写长句", "`为什么 Aurora 收端不能反压并且 FIFO 多深才够` 不如拆成两次：先查 `Aurora RX 无反压 rx_dst_rdy`，再查 `FIFO overflow waterline`。", kind="note")

    add_heading(doc, "8. 如何看懂检索结果", 1)
    add_para(doc, "每条结果至少包含四部分：")
    add_list_item(doc, "标题与分类：先判断它是官方基础资料、专题指南、当前工程还是跨 IP 资料。", decimal_num)
    add_list_item(doc, "source：原始文件绝对路径。需要引用、修改或核对上下文时回到这里。", decimal_num)
    add_list_item(doc, "corpus：知识库生成的文本副本。适合快速阅读和全文搜索，不替代原始文件。", decimal_num)
    add_list_item(doc, "命中片段：只展示附近文字。重要结论必须打开 source 或 corpus 查看完整上下文。", decimal_num)
    add_heading(doc, "8.1 证据优先级", 2)
    add_table(doc, ["优先级", "证据", "用法"], [
        ["1", "官方 PG / UG / 协议规范", "确认接口、参数和协议定义。"],
        ["2", "*_verified 工程与验证摘要", "确认实际 XCI/XDC、脚本和已完成的验证。"],
        ["3", "当前 RTL、测试平台、设计说明", "确认代码实际行为和当前集成状态。"],
        ["4", "历史示例、过程记录、对话材料", "解释背景；不能覆盖前三级证据。"],
    ], [1000, 3100, 5260],
       aligns=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
       caption="表 5  证据优先级")
    add_heading(doc, "8.2 状态词不能混用", 2)
    add_table(doc, ["状态", "含义"], [
        ["已通过", "有明确的仿真、编译、实现或板级结果，并写清具体层级。"],
        ["已实现", "RTL、测试平台或脚本已经存在，但不能据此推断运行通过。"],
        ["待验证", "目标已定义，当前证据中没有完成结果。"],
        ["设计边界", "当前模型故意不覆盖，例如 GT、信号完整性、操作系统 ISR。"],
        ["文档与 RTL 不一致", "以当前 RTL 为行为基准，差异列为待收口项。"],
    ], [2000, 7360], caption="表 6  知识库状态词")

    add_heading(doc, "9. 各专题具体能查到什么", 1)
    add_heading(doc, "9.1 通用基础 foundation", 2)
    add_para(doc, "这一类用于解决跨项目都会遇到的问题：Vivado 工程方法、Tcl、综合、仿真、静态时序、AXI/AXI-Stream、FIFO、CDC、时钟、AXI DMA 和 DataMover。")
    add_code(doc, ".\\search.ps1 'report_timing_summary WNS TNS' --category foundation\n.\\search.ps1 'AXI AWVALID WREADY 4KB boundary' --category foundation\n.\\search.ps1 'FIFO Generator FWFT prog_full' --category foundation")

    add_heading(doc, "9.2 Aurora", 2)
    add_para(doc, "包括 Aurora 64B/66B 10.3125G、Aurora 8B/10B 1.25G、LocalLink/AXI 适配、GT 与协议复位、时钟、回压、CRC，以及 verified 工程的配置和验证摘要。")
    add_list_item(doc, "10.3125G、64-bit 用户接口的用户时钟约为 161.1328125 MHz；156.25 MHz 不能直接当作这个用户时钟。", bullet_num)
    add_list_item(doc, "TX 可以反压，握手前 data/SOF/EOF/rem 必须保持；RX 没有 ready，后级需要 FIFO 和 overflow 策略。", bullet_num)
    add_list_item(doc, "64B/66B 当前有接口审计、XCI 编译/展开和接口仿真证据，但不能据此宣称实现、时序或上板通过。", bullet_num)
    add_code(doc, ".\\search.ps1 '161.1328125 user_clk_out 64B66B' --category aurora\n.\\search.ps1 'rx_dst_rdy no backpressure overflow' --category aurora\n.\\search.ps1 'VALIDATION_SUMMARY synthesis NOT RUN' --category aurora")

    add_heading(doc, "9.3 PCIe / XDMA / NVMe", 2)
    add_para(doc, "包括 PCIe Gen3 x2、XDMA Endpoint、官方 Root Port/TPI、BAR、H2C/C2H、MSI/user IRQ、CQ Host-memory sink、FDMA/Aurora 适配，以及 NVMe/DDR 数据管线。")
    add_list_item(doc, "RTDS_test-1ch 的 G2 最小闭环已有 512 B 双向数据、MSI、read-to-clear 和日志门禁 PASS 证据。", bullet_num)
    add_list_item(doc, "G3 压力回归和 G4 板卡/软件验证未宣称完成。`usr_irq_ack` 证明 MSI 已发送，不证明 OS ISR 已运行。", bullet_num)
    add_list_item(doc, "官方 RP 的 DATA_STORE 是 H2C/描述符读取的数据源，不是 C2H 通用 Host RAM；C2H 应在 CQ 接口接收。", bullet_num)
    add_code(doc, ".\\search.ps1 'RTDS_ROOTPORT_PASS errors=0' --category pcie\n.\\search.ps1 'CQ Host memory sink DATA_STORE C2H' --category pcie\n.\\search.ps1 'usr_irq_req usr_irq_ack msi_enable' --category pcie")

    add_heading(doc, "9.4 DDR / MIG", 2)
    add_para(doc, "包括 KU040 DDR4 MIG、AXI BIST、校准与 ui_clk 复位、AW/W 独立握手、行为仿真、实现时序和板级验证边界。")
    add_list_item(doc, "AXI BIST 行为级仿真已通过；MIG + DDR4 器件模型联合仿真报告仍未收口。", bullet_num)
    add_list_item(doc, "2026-07-27 的实现产物显示 route 成功、WNS=0.010 ns、TNS=0，bitstream 已生成；这不等于板卡校准或真实 DDR 数据测试通过。", bullet_num)
    add_list_item(doc, "实现日志仍有 CFGBVS/CONFIG_VOLTAGE 和 no-routable-load 两类 DRC warning 待处理。", bullet_num)
    add_code(doc, ".\\search.ps1 'init_calib_complete ui_clk_sync_rst axi_aresetn' --category ddr\n.\\search.ps1 'AW W independent handshake BRESP RRESP' --category ddr\n.\\search.ps1 'WNS 0.010 CFGBVS CONFIG_VOLTAGE' --category guide")

    add_heading(doc, "9.5 以太网 / RGMII / 风暴过滤 / DMA", 2)
    add_para(doc, "包括 RGMII 到 GMII、前导码/SFD、VLAN 与 FCS 剥离、CRC、目的 MAC、ACL、组播 Hash、帧数/字节数风暴、重复帧、GOOSE/SV/Other 分类和 DataMover S2MM 写 DMA。")
    add_list_item(doc, "当前有接收过滤顶层和写 DMA 顶层，但缺少最终系统顶层、寄存器封装、读 DMA、部分 Vivado IP/约束和正式 testbench。", bullet_num)
    add_list_item(doc, "cfg_filter_en=0 只旁路策略过滤；RXER、长度、FIFO 资源和独立 CRC 检查仍可能丢帧。", bullet_num)
    add_list_item(doc, "验证 XLSX 定义了 22 个测试点和 17 个归并用例，这是测试规划，不是执行结果。", bullet_num)
    add_code(doc, ".\\search.ps1 'DROP_CRC_ERR DROP_DST_MISS 丢弃优先级' --category ethernet\n.\\search.ps1 '风暴窗口 frame byte threshold epoch' --category ethernet\n.\\search.ps1 'GOOSE SV Other DataMover S2MM' --category ethernet")

    add_heading(doc, "9.6 跨 IP 集成 cross_ip", 2)
    add_para(doc, "用于查 Aurora、PCIe、DDR、NVMe 之间的接口转换、CDC、FIFO、描述符、地址窗口、回压和回归平台。很多系统问题不属于单一 IP，这一类通常比只查 `aurora` 或 `pcie` 更容易命中。")
    add_code(doc, ".\\search.ps1 'FDMA Aurora bridge FIFO overflow' --category cross_ip\n.\\search.ps1 'descriptor payload commit pointer' --category cross_ip\n.\\search.ps1 '4KB boundary burst split' --category cross_ip")

    add_heading(doc, "10. 常见任务怎么查", 1)
    add_table(doc, ["你想做什么", "建议第一条查询", "下一步"], [
        ["第一次接手 Aurora 工程", "'Aurora 工程速查' --category guide", "再查具体速率、XCI、时钟和验证摘要。"],
        ["确认 PCIe 大仿真是否真的通过", "'RTDS_ROOTPORT_PASS errors=0'", "打开问题台账和验证环境说明，核对 G2/G3/G4。"],
        ["定位 DMA 数据没有到 Host", "'CQ sink DATA_STORE C2H' --category pcie", "检查 CQ handshake、地址、tkeep 和 scoreboard。"],
        ["确认 DDR 是仿真通过还是上板通过", "'DDR 验证层级 已通过 未执行' --category guide", "分别核对行为仿真、实现、联合仿真和板卡证据。"],
        ["查以太网某帧为什么丢弃", "'drop_reason 丢弃优先级' --category ethernet", "再加 CRC、MAC、ACL、storm 或 FIFO 关键词。"],
        ["检查跨时钟域风险", "'CDC toggle sticky clear async FIFO'", "同时查 foundation 和 cross_ip。"],
        ["找 Vivado 命令", "'目标命令名' --category foundation", "先查 UG835，再看工程 Tcl 的实际调用。"],
    ], [2600, 3200, 3560], caption="表 7  常见任务查询入口")

    add_heading(doc, "11. 推荐的工作流程", 1)
    add_heading(doc, "11.1 回答一个技术问题", 2)
    for text in [
        "先查 guide，确认主题结论、边界和术语。",
        "再查对应 category，找到官方资料、当前 RTL 和验证记录。",
        "按证据优先级解决冲突。文档与 RTL 不一致时，先说明差异，不要混成一个结论。",
        "打开 source 阅读完整上下文，记录文件路径和关键状态。",
    ]:
        add_list_item(doc, text, decimal_num)
    add_heading(doc, "11.2 调试一个工程问题", 2)
    for text in [
        "用端口名、状态机名、错误码搜索，不要只用现象描述。",
        "同时搜索接口两端。例如 valid/ready 错位时，把源模块和目标模块名放在同一查询里。",
        "先确认方向、时钟、复位和尾包语义，再看数据内容。",
        "查找已有 testbench、assertion、scoreboard 和问题台账，避免重复走旧路径。",
        "最终以新日志和可重复用例收口，不用旧截图代替。",
    ]:
        add_list_item(doc, text, decimal_num)
    add_heading(doc, "11.3 写设计或验证文档", 2)
    add_para(doc, "先用主题页建立章节框架，再从 corpus 和 source 提取证据。引用验证结果时必须带层级，例如“行为级仿真 PASS”“实现时序满足”“板卡未执行”，不能只写一个笼统的“验证通过”。")

    add_heading(doc, "12. 更新与重建知识库", 1)
    add_heading(doc, "12.1 增量重建", 2)
    add_para(doc, "远端或本机资料有新增、修改后，优先运行增量构建。未变化文件会使用哈希缓存。")
    add_code(doc, "cd E:\\centOS\\xilinx_knowledge_base\n.\\build.ps1")
    add_heading(doc, "12.2 干净重建", 2)
    add_para(doc, "来源规则大改、索引异常或需要重新生成全部 corpus 时使用：")
    add_code(doc, ".\\build.ps1 -Clean")
    add_heading(doc, "12.3 查看详细构建过程", 2)
    add_code(doc, ".\\build.ps1 -VerboseReport")
    add_heading(doc, "12.4 修改收录范围", 2)
    add_para(doc, "编辑 `config/sources.json`。新增目录应写明 category、priority、tags 和 max_size_mb；排除项应放在目录名、后缀或路径关键词规则中。修改后先做增量构建，再检查统计和 build_report。")
    add_callout(doc, "不要直接扫整个磁盘", "知识库的价值来自范围受控和证据可追溯。新增来源时尽量指向 docs、rtl、sim、scripts、constraints 等明确目录，不要把完整 Vivado 工程生成树作为根目录。", kind="risk")

    add_heading(doc, "13. 如何确认构建是否健康", 1)
    add_code(doc, ".\\search.ps1 --stats")
    add_para(doc, "然后检查下列文件：")
    add_list_item(doc, "index/build_report.json：errors 和 missing_sources 应为 0；duplicates 是正常去重记录。", bullet_num)
    add_list_item(doc, "index/manifest.json：检查新增文件的 source_path、category、extractor、chars 和 aliases。", bullet_num)
    add_list_item(doc, "corpus/<category>/：确认新增条目有可读的 Markdown 文本副本。", bullet_num)
    add_list_item(doc, "运行四类代表性查询，确认中文、英文信号名和分类过滤都能命中。", bullet_num)
    add_callout(doc, "当前健康状态", f"本指南生成时：documents={stats['documents']}，duplicates={stats['duplicates']}，errors={stats['errors']}，missing_sources={stats['missing']}。", kind="positive")

    add_heading(doc, "14. 常见问题", 1)
    add_table(doc, ["现象", "原因与处理"], [
        ["搜索没有结果", "缩短查询词；去掉 category；改用信号名、模块名或英文术语。"],
        ["结果太多", "增加两个高辨识度关键词，或使用 --category 和 --limit。"],
        ["结果中文显示乱码", "从 search.ps1 运行，不要绕过其中的 UTF-8 环境设置。"],
        ["source 路径打不开", "远端 R: 可能未挂载；先恢复映射盘。corpus 文本仍可用于定位。"],
        ["同一资料出现多个位置", "构建器按 SHA-256 合并正文，其他路径记录在 aliases/duplicate sources。"],
        ["加密 Office 只有占位文字", "需要文档密码或未加密副本；占位条目不会参与技术结论。"],
        ["看到 bitstream，能否认为板卡通过", "不能。bitstream 只证明构建完成，仍需板卡校准、数据和长稳证据。"],
        ["设计文档与 RTL 冲突", "以当前 RTL 行为为基准，同时记录差异并查验证用例；不要静默改写结论。"],
        ["重建很慢", "第一次抽取大 PDF 较慢；后续增量构建会使用 cache/extracted。"],
        ["担心误收 HDMI/USB/图像资料", "检查 config/sources.json 排除规则和 manifest 路径审计；当前验收为 0 命中。"],
    ], [2800, 6560], caption="表 8  常见问题与处理")

    add_heading(doc, "15. 命令速查", 1)
    add_code(doc, "# 进入知识库\ncd E:\\centOS\\xilinx_knowledge_base\n\n# 全库搜索\n.\\search.ps1 '查询词'\n\n# 分类搜索\n.\\search.ps1 '查询词' --category aurora\n.\\search.ps1 '查询词' --category pcie\n.\\search.ps1 '查询词' --category ddr\n.\\search.ps1 '查询词' --category ethernet\n.\\search.ps1 '查询词' --category cross_ip\n\n# 限制条数 / JSON 输出\n.\\search.ps1 '查询词' --limit 15\n.\\search.ps1 '查询词' --json\n\n# 查看统计\n.\\search.ps1 --stats\n\n# 增量 / 干净 / 详细重建\n.\\build.ps1\n.\\build.ps1 -Clean\n.\\build.ps1 -VerboseReport")

    add_heading(doc, "16. 最后检查清单", 1)
    checklist = [
        "我先看了 guide，再回到原始 source。",
        "我知道当前结论属于仿真、实现、板卡还是软件层级。",
        "我没有把“已实现”写成“已通过”。",
        "我核对了接口方向、时钟、复位、回压和尾包语义。",
        "我记录了原始文件路径，而不是只复制命中片段。",
        "资料更新后，我重建并检查了 build_report。",
    ]
    for item in checklist:
        add_list_item(doc, "□ " + item, bullet_num)
    add_callout(doc, "一句话原则", "先用知识库缩小范围，再用原始工程和验证证据下结论。", kind="positive")

    add_heading(doc, "附录 A：主题页入口", 1)
    for path, desc in [
        ("knowledge\\00_范围与来源.md", "范围、来源、排除项和状态词。"),
        ("knowledge\\01_Aurora_工程速查.md", "Aurora 速率、时钟、LocalLink、复位与验证边界。"),
        ("knowledge\\02_PCIe_XDMA_工程速查.md", "XDMA Root Port、H2C/C2H、MSI 和大仿真状态。"),
        ("knowledge\\03_DDR_MIG_工程速查.md", "DDR4 MIG、BIST、实现和板卡边界。"),
        ("knowledge\\04_以太网_工程速查.md", "RGMII、风暴过滤、分类 FIFO 和 DMA。"),
        ("knowledge\\05_跨IP_集成检查表.md", "CDC、FIFO、描述符、内存一致性和分层验收。"),
    ]:
        add_para(doc, f"{path}  —  {desc}", style="Small Note")

    add_heading(doc, "附录 B：关键索引文件", 1)
    for path in [
        r"E:\centOS\xilinx_knowledge_base\index\xilinx_kb.sqlite3",
        r"E:\centOS\xilinx_knowledge_base\index\manifest.json",
        r"E:\centOS\xilinx_knowledge_base\index\manifest.csv",
        r"E:\centOS\xilinx_knowledge_base\index\build_report.json",
        r"E:\centOS\xilinx_knowledge_base\config\sources.json",
    ]:
        add_code(doc, path)

    # Cover uses the same section; restart visible numbering for the full document.
    add_page_numbering_restart(doc.sections[0], 1)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
