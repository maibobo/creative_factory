from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "Xilinx_FPGA工程知识库使用指南.docx"
REQUIRED = {
    "[Content_Types].xml",
    "_rels/.rels",
    "word/document.xml",
    "word/styles.xml",
    "word/numbering.xml",
    "word/settings.xml",
    "word/header1.xml",
    "word/footer1.xml",
    "docProps/core.xml",
}


def attr_int(element, name):
    return int(element.get(qn(name)))


def pt_value(value):
    return 0.0 if value is None else float(value.pt)


def verify():
    assert DOCX.is_file() and DOCX.stat().st_size > 30_000, "DOCX missing or unexpectedly small"
    with zipfile.ZipFile(DOCX) as archive:
        bad = archive.testzip()
        assert bad is None, f"corrupt ZIP member: {bad}"
        names = set(archive.namelist())
        missing = REQUIRED - names
        assert not missing, f"missing OOXML members: {sorted(missing)}"
        assert not [n for n in names if n.startswith("word/media/")], "unexpected embedded media"
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                ET.fromstring(archive.read(name))
        settings = ET.fromstring(archive.read("word/settings.xml"))
        update_fields = settings.find(qn("w:updateFields"))
        assert update_fields is not None and update_fields.get(qn("w:val")) == "true"
        document_bytes = archive.read("word/document.xml")
        document_xml = document_bytes.decode("utf-8")
        footer_xml = archive.read("word/footer1.xml").decode("utf-8")
        assert "TOC \\o" in document_xml, "missing TOC field"
        assert ">PAGE<" in footer_xml, "missing PAGE field"
        document_root = ET.fromstring(document_bytes)
        assert document_root.find(".//" + qn("w:ins")) is None
        assert document_root.find(".//" + qn("w:del")) is None

    doc = Document(DOCX)
    assert doc.core_properties.title == "Xilinx / FPGA 工程知识库使用指南"
    styles = doc.styles
    expected_styles = {
        "Normal": (11, None, 0, 6, 1.25),
        "Heading 1": (16, "2E74B5", 18, 10, 1.0),
        "Heading 2": (13, "2E74B5", 14, 7, 1.0),
        "Heading 3": (12, "1F4D78", 10, 5, 1.0),
    }
    for name, (size, color, before, after, line) in expected_styles.items():
        style = styles[name]
        assert round(style.font.size.pt, 2) == size, f"{name} size"
        if color:
            assert str(style.font.color.rgb) == color, f"{name} color"
        assert round(pt_value(style.paragraph_format.space_before), 2) == before, f"{name} before"
        assert round(pt_value(style.paragraph_format.space_after), 2) == after, f"{name} after"
        assert round(float(style.paragraph_format.line_spacing), 2) == line, f"{name} line spacing"
        rfonts = style.element.get_or_add_rPr().rFonts
        assert rfonts is not None and rfonts.get(qn("w:eastAsia")) == "Microsoft YaHei", f"{name} East Asian font"

    headings = [p for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    assert len(headings) >= 35, f"too few headings: {len(headings)}"
    assert len(doc.tables) == 9, f"unexpected table count: {len(doc.tables)}"

    list_paragraphs = 0
    for p in doc.paragraphs:
        ppr = p._p.pPr
        if ppr is not None and ppr.numPr is not None:
            list_paragraphs += 1
    assert list_paragraphs >= 45, f"too few real numbered/list paragraphs: {list_paragraphs}"
    fake_lists = [p.text for p in doc.paragraphs if p.text.startswith(("- ", "• ")) and
                  (p._p.pPr is None or p._p.pPr.numPr is None)]
    assert not fake_lists, f"fake list paragraphs: {fake_lists[:3]}"

    full_text = "\n".join(p.text for p in doc.paragraphs)
    required_text = [
        "396", "17,054,012", "errors=0", "missing_sources=0",
        "search.ps1", "build.ps1", "foundation", "cross_ip",
        "Aurora", "PCIe / XDMA / NVMe", "DDR / MIG", "风暴过滤",
    ]
    for marker in required_text:
        assert marker in full_text, f"missing required guide content: {marker}"

    for idx, table in enumerate(doc.tables, start=1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        layout = tbl_pr.find(qn("w:tblLayout"))
        assert tbl_w is not None and attr_int(tbl_w, "w:w") == 9360, f"table {idx} tblW"
        assert tbl_w.get(qn("w:type")) == "dxa", f"table {idx} tblW type"
        assert tbl_ind is not None and attr_int(tbl_ind, "w:w") == 120, f"table {idx} tblInd"
        assert layout is not None and layout.get(qn("w:type")) == "fixed", f"table {idx} layout"
        grid_widths = [attr_int(x, "w:w") for x in table._tbl.tblGrid.findall(qn("w:gridCol"))]
        assert sum(grid_widths) == 9360, f"table {idx} grid sum {sum(grid_widths)}"
        first_tr_pr = table.rows[0]._tr.trPr
        assert first_tr_pr is not None and first_tr_pr.find(qn("w:tblHeader")) is not None, f"table {idx} header row"
        for row in table.rows:
            for cell, expected in zip(row.cells, grid_widths):
                tc_w = cell._tc.tcPr.find(qn("w:tcW"))
                assert tc_w is not None and attr_int(tc_w, "w:w") == expected, f"table {idx} tcW"
                assert cell._tc.tcPr.find(qn("w:tcMar")) is not None, f"table {idx} cell margins"

    section = doc.sections[0]
    assert round(section.page_width.inches, 3) == 8.5
    assert round(section.page_height.inches, 3) == 11.0
    for margin in [section.top_margin, section.bottom_margin, section.left_margin, section.right_margin]:
        assert round(margin.inches, 3) == 1.0

    print(f"STRUCTURE_OK file={DOCX}")
    print(f"size={DOCX.stat().st_size} headings={len(headings)} tables={len(doc.tables)} lists={list_paragraphs}")


if __name__ == "__main__":
    verify()
