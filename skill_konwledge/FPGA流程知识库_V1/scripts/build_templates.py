#!/usr/bin/env python3
"""从指定Word前置样式及Excel表格样式生成独立的流程模板，不修改来源。"""
import argparse
from copy import deepcopy, copy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORD_SOURCE = Path(
    r"R:\fdma_aurora_brg_aurora_loop-validate\docs"
    r"\fdma_aurora_brg_IP总体设计与验证概述-7-6(1).docx"
)
EXCEL_SOURCE = Path(
    r"F:\RTDS_test_DDR4_64bit_batch_dev\docs"
    r"\RTDS_设计规格_测试点_测试用例_精炼版_20261010.xlsx"
)
FIELDS = {
    "设计规格": ["编号", "模块", "功能类别", "功能名称", "设计要求", "备注"],
    "测试点": [
        "编号", "类别", "子类别", "测试点内容", "对应验证用例名称", "设计规格反标",
    ],
    "测试用例": [
        "编号", "名称", "验证层级", "测试点反标", "准备与输入", "预期结果", "执行结果",
    ],
}
WIDTHS = {
    "设计规格": [15, 25, 24, 32, 66, 32],
    "测试点": [26, 25, 26, 56, 38, 66],
    "测试用例": [16, 36, 20, 52, 66, 62, 38],
}


def replace_text(element, value):
    """保留来源段落和run格式，仅替换文字。"""
    nodes = element.xpath(".//w:t")
    if nodes:
        for node in nodes:
            node.text = ""
        nodes[0].text = value
        nodes[0].set(qn("xml:space"), "preserve")
    else:
        paragraph = (
            element if element.tag == qn("w:p")
            else element.find(qn("w:p"))
        )
        if paragraph is None:
            paragraph = OxmlElement("w:p")
            element.append(paragraph)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = value
        run.append(text)
        paragraph.append(run)


def toc(doc):
    toc_style = next(s for s in doc.styles if s.name.lower() == "toc 1")
    paragraph = doc.add_paragraph(style=toc_style)
    for kind in ["begin", "instruction", "separate", "text", "end"]:
        run = paragraph.add_run()
        if kind == "instruction":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = r' TOC \o "1-3" \h \z \u '
        elif kind == "text":
            run.text = "[在Word中更新目录]"
            continue
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)


def front_document(source, title):
    doc = Document()
    # 仅复制字体/段落/编号样式及前置内容，不携带原业务正文和图片。
    doc.part._styles_part._element = deepcopy(source.styles._element)
    doc.part.numbering_part._element = deepcopy(source.part.numbering_part._element)
    for relation in [RT.THEME, RT.FONT_TABLE]:
        try:
            doc.part.part_related_by(relation)._blob = source.part.part_related_by(relation).blob
        except KeyError:
            pass
    body = doc._element.body
    for element in list(body):
        body.remove(element)
    section = deepcopy(source._element.body.sectPr)
    # 来源页眉页脚为空；保留页设置，去掉指向来源文档的关系ID。
    for element in list(section):
        if element.tag in [qn("w:headerReference"), qn("w:footerReference")]:
            section.remove(element)
    body.append(section)
    style_names = {s.style_id: s.name.lower() for s in source.styles}
    front = []
    for element in source._element.body:
        if element.tag == qn("w:p"):
            node = element.find(qn("w:pPr"))
            style = node.find(qn("w:pStyle")) if node is not None else None
            if style is not None and style_names.get(style.get(qn("w:val"))) == "heading 1":
                break
        if element.tag != qn("w:sectPr"):
            front.append(element)
    found_toc = False
    for original in front:
        element = deepcopy(original)
        instructions = original.xpath(".//w:instrText")
        if any("TOC" in (node.text or "") for node in instructions):
            # 保留目录标题的格式，丢弃原目录缓存和业务章节。
            inside = False
            for child in list(element):
                if child.tag == qn("w:pPr"):
                    continue
                if child.xpath(".//w:fldChar") or child.xpath(".//w:instrText"):
                    inside = True
                if inside:
                    element.remove(child)
            if "".join(element.xpath(".//w:t/text()")).strip():
                section.addprevious(element)
            toc(doc)
            found_toc = True
            break
        section.addprevious(element)
    if not found_toc:
        doc.add_paragraph("目录", style="Heading 1")
        toc(doc)
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "封面文档标题" and paragraph.text.strip():
            replace_text(paragraph._p, title)
            break
    cover, approval, revisions = doc.tables[:3]
    replace_text(cover.cell(1, 0)._tc, "[项目名称]")
    replace_text(cover.cell(1, 1)._tc, "[密级]")
    replace_text(cover.cell(2, 1)._tc, "Total [页数] pages 共 [页数] 页")
    if len(cover.rows) > 3:
        replace_text(cover.cell(3, 0)._tc, "[产品版本]")
    for row in approval.rows:
        replace_text(row.cells[1]._tc, "[姓名]")
        replace_text(row.cells[3]._tc, "[日期]")
    for row in revisions.rows[1:]:
        for cell, value in zip(row.cells, ["[版本]", "[修改内容]", "[姓名]", "[日期]"]):
            replace_text(cell._tc, value)
    update = doc.settings._element.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        doc.settings._element.append(update)
    update.set(qn("w:val"), "true")
    core = doc.core_properties
    core.title = title
    core.subject = "FPGA设计验证流程模板"
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.keywords = ""
    core.created = datetime.now(timezone.utc)
    core.modified = datetime.now(timezone.utc)
    return doc


def heading(doc, title, level=1):
    style = next(s for s in doc.styles if s.name.lower() == f"heading {level}")
    p = doc.add_paragraph(title, style=style)
    # 样式原有自动编号；正文采用明确章节号，防止双重编号。
    properties = p._p.get_or_add_pPr()
    numbers = OxmlElement("w:numPr")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), "0")
    numbers.append(number)
    properties.append(numbers)
    return p


def instruction(doc, value):
    return doc.add_paragraph("【填写说明】" + value, style="Normal")


def placeholder(doc, value):
    return doc.add_paragraph("[" + value + "]", style="Normal")


def table(doc, headers, values, weights=None):
    output = doc.add_table(rows=1, cols=len(headers))
    grid = next((s for s in doc.styles if s.name.lower() == "table grid"), None)
    if grid is not None:
        output.style = grid
    output.autofit = False
    width = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    weights = weights or [1] * len(headers)
    for i, (cell, label) in enumerate(zip(output.rows[0].cells, headers)):
        cell.text = label
        for run in cell.paragraphs[0].runs:
            run.bold = True
        output.columns[i].width = int(width * weights[i] / sum(weights))
        cell.width = output.columns[i].width
    cells = output.add_row().cells
    for i, (cell, value) in enumerate(zip(cells, values)):
        cell.text = value
        cell.width = output.columns[i].width
    return output


def design(source):
    doc = front_document(source, "[项目/模块名称]设计方案")
    heading(doc, "1 概述")
    instruction(doc, "说明系统目标及当前功能范围，采用工程师提供的需求初定值。")
    heading(doc, "2 模块架构")
    placeholder(doc, "插入架构图，标明顶层、子模块职责及主要连接")
    table(doc, ["模块", "功能职责", "模块之间的交互"], ["[模块名称]", "[实现什么功能]", "[输入、输出及关系]"], [2, 4, 5])
    heading(doc, "3 分模块详细设计")
    for index, module in [(1, "顶层/模块名称"), (2, "子模块名称")]:
        heading(doc, f"3.{index} [{module}]", 2)
        heading(doc, f"3.{index}.1 功能规格", 3)
        instruction(doc, "列本模块功能规格及要实现的功能，可结合实现展开；与Excel规格同步。")
        table(doc, ["规格编号", "功能类别", "要实现的功能"], ["[编号]", "[类别]", "[功能行为及明确的数值要求]"], [2, 2, 7])
        heading(doc, f"3.{index}.2 信号列表", 3)
        table(
            doc, ["信号名", "中文名称", "方向", "位宽", "含义"],
            ["[signal_name]", "[中文名称]", "[I/O]", "[位宽]", "[功能语义；必要时写时钟域和有效条件]"],
            [3, 2, 1, 1, 5],
        )
        heading(doc, f"3.{index}.3 时序图", 3)
        placeholder(doc, "插入相关功能的关键时序图，说明正常交互及必要的等待、边界或异常")
        heading(doc, f"3.{index}.4 功能实现", 3)
        instruction(doc, "按功能点说明内部逻辑和数据/控制过程；有关状态机、交互、跨域或异常处理随具体功能展开。")
    return doc


def verification(source):
    doc = front_document(source, "[项目/模块名称]验证方案")
    heading(doc, "1 验证目的")
    instruction(doc, "说明本轮要验证的功能、目标及边界。")
    heading(doc, "2 验证方法")
    heading(doc, "2.1 UT", 2)
    instruction(doc, "说明单模块验证对象、功能范围、激励和检查方法。")
    heading(doc, "2.2 IT", 2)
    instruction(doc, "说明多模块组合及接口、总线、跨时钟交互；交代实际模块与模型的范围。")
    heading(doc, "2.3 回归安排及完成条件", 2)
    instruction(doc, "按用例表行序安排定向及适用的随机场景，明确Seed、超时、每日回归集合/入口和功能完成条件。")
    heading(doc, "3 环境组件介绍")
    instruction(doc, "选基本TB、轻量SV、完整SV或轻量UVM；按实际功能复杂度说明选择并列真实组件。")
    table(doc, ["组件", "职责", "参与功能"], ["[组件名称]", "[组件作用]", "[服务的功能/测试点]"], [2, 4, 5])
    heading(doc, "4 环境架构与数据流")
    placeholder(doc, "插入环境架构图，标明激励、观察、检查路径及DUT边界")
    instruction(doc, "说明事务、数据和控制事件在各组件间如何流动，何处驱动、采样、建模和检查。")
    heading(doc, "5 环境组件功能设计")
    components = [
        ("激励", "说明事务类型、核心参数、序列和时序变化，如何形成相关功能及边界/异常场景。"),
        ("driver/BFM", "说明事务如何转换成接口时序，如何处理握手、等待及所服务的测试点。"),
        ("monitor", "说明采样条件、关联字段、完成识别及数据送往何处。"),
        ("RM", "说明建模功能、输入/输出关系、必要的时序约定和结果去向。"),
        ("checker/scoreboard", "按测试点说明检查数据、顺序、状态和时序的方法，以及失败及完成条件。"),
        ("其他组件", "按需说明coverage、interface、env/config等；未使用的组件小节删除。"),
    ]
    for index, (name, explanation) in enumerate(components, 1):
        heading(doc, f"5.{index} {name}", 2)
        instruction(doc, explanation)
    instruction(doc, "非标准SV/UVM环境按真实组成调整上述小节，说明各部分功能及关系。")
    heading(doc, "6 用例实例说明")
    heading(doc, "6.1 [代表用例名称]", 2)
    heading(doc, "6.1.1 测试点反标", 3)
    placeholder(doc, "测试点编号＋测试点内容")
    heading(doc, "6.1.2 涉及组件与架构", 3)
    placeholder(doc, "列本例涉及的组件，并插入本例的组件/架构示意")
    heading(doc, "6.1.3 数据流演示", 3)
    placeholder(doc, "按触发、激励、传递、观察和检查顺序演示本例数据流")
    heading(doc, "6.1.4 检查现象", 3)
    placeholder(doc, "说明各检查点应观察到的现象及判定方法，与Excel用例相符")
    return doc


def report(source):
    doc = front_document(source, "[项目/模块名称]仿真报告")
    heading(doc, "1 测试版本")
    table(doc, ["对象", "版本/位置", "说明"], ["[RTL、验证代码、工具或配置]", "[版本/位置]", "[本轮使用内容]"], [2, 4, 4])
    heading(doc, "2 仿真与回归结果概览")
    instruction(doc, "概括本轮仿真及每日回归的范围、执行情况、通过/失败/待执行项和关键限制，引用测试表及结果目录。")
    table(doc, ["集合/范围", "执行情况", "总体结果", "说明"], ["[集合]", "[次数/范围]", "[真实结果]", "[日志/结果位置及限制]"], [2, 2, 2, 5])
    heading(doc, "3 功能与用例测试结果")
    for index in [1, 2]:
        heading(doc, f"3.{index} [功能/用例名称]", 2)
        heading(doc, f"3.{index}.1 测试点与配置", 3)
        placeholder(doc, "测试点编号＋内容，以及本例有关配置/Seed")
        heading(doc, f"3.{index}.2 测试现象", 3)
        table(doc, ["观察点", "实际现象", "必要证据位置"], ["[观察点]", "[实际现象]", "[日志/关键波形位置]"], [2, 5, 4])
        heading(doc, f"3.{index}.3 结论", 3)
        placeholder(doc, "依据相关功能判据给出结论；必要时说明问题修复与复测结果，未运行如实记录")
    heading(doc, "4 验证结论")
    heading(doc, "4.1 功能验证结论", 2)
    placeholder(doc, "汇总本轮功能及测试点的验证结论")
    heading(doc, "4.2 implement及时序收敛结果", 2)
    placeholder(doc, "记录最终实现检查状态、时序结论及相关修改后的必要回归结果；未执行标明状态")
    heading(doc, "4.3 未覆盖内容与遗留问题", 2)
    placeholder(doc, "说明未覆盖功能、遗留问题和后续工作")
    return doc


def spreadsheet(source_path, target):
    source = load_workbook(source_path, read_only=False)
    book = Workbook()
    book.remove(book.active)
    comments = {
        ("设计规格", "设计要求"): "按顶层/子模块和功能类别逐条写。有关定值来源于人工系统需求。",
        ("设计规格", "备注"): "补充说明；不以备注替代功能要求。",
        ("测试点", "编号"): "项目/IP简称．三位功能类别序号．三位类别内序号，例如DDR.001.001。",
        ("测试点", "对应验证用例名称"): "写用例名称，不只列用例编号。",
        ("测试点", "设计规格反标"): "规格编号＋规格具体描述。",
        ("测试用例", "测试点反标"): "测试点编号＋测试点内容。",
        ("测试用例", "准备与输入"): "通常约3–30行，聚焦本例核心变量与激励。写中文含义、来源、选用原因及所覆盖功能。",
        ("测试用例", "预期结果"): "可检查的现象、值或时序。不再另设执行过程。",
        ("测试用例", "执行结果"): "填实际结论及必要日志/Seed位置；未执行填待执行。",
    }
    for name, fields in FIELDS.items():
        sheet = book.create_sheet(name)
        original = source[name]
        for col, field in enumerate(fields, 1):
            cell = sheet.cell(1, col, field)
            cell.fill = copy(original.cell(1, 1).fill)
            cell.border = copy(original.cell(1, 1).border)
            cell.font = Font(
                name="Microsoft YaHei", size=13, bold=True,
                color=copy(original.cell(1, 1).font.color),
            )
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if (name, field) in comments:
                cell.comment = Comment(comments[(name, field)], "模板填写说明")
            sheet.column_dimensions[get_column_letter(col)].width = WIDTHS[name][col - 1]
            for row in range(2, 12):
                body_cell = sheet.cell(row, col)
                body_cell.font = Font(name="Microsoft YaHei", size=12)
                body_cell.border = copy(original.cell(2, 1).border)
                body_cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[1].height = 38
        # 不锁定正文行高，填写后可自动适配多行文本。
        sheet.sheet_format.defaultRowHeight = 19
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}11"
        sheet.print_title_rows = "1:1"
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    source.close()
    book.save(target)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--word-source", type=Path, default=WORD_SOURCE)
    parser.add_argument("--excel-source", type=Path, default=EXCEL_SOURCE)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    output = ROOT / "templates"
    output.mkdir(parents=True, exist_ok=True)
    names = ["设计方案模板.docx", "验证方案模板.docx", "仿真报告模板.docx", "设计规格_测试点_测试用例模板.xlsx"]
    if not args.force and any((output / name).exists() for name in names):
        raise SystemExit("模板已存在；明确需要重建时使用--force，避免覆盖人工填写内容。")
    source = Document(args.word_source)
    for name, factory in [
        ("设计方案模板.docx", design), ("验证方案模板.docx", verification),
        ("仿真报告模板.docx", report),
    ]:
        factory(source).save(output / name)
        print(name)
    spreadsheet(args.excel_source, output / names[-1])
    print(names[-1])


if __name__ == "__main__":
    main()
