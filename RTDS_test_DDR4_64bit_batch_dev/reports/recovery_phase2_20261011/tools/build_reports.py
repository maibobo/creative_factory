from pathlib import Path
from collections import Counter
import json,re,hashlib,shutil,difflib,zipfile,posixpath,xml.etree.ElementTree as ET,datetime
from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
R=Path('F:/RTDS_test_DDR4_64bit_batch_dev');S=Path(__file__).parent
D=R/'docs';F=D/'RTDS_functional_20261011';REP=R/'reports/recovery_phase2_20261011';REP.mkdir(parents=True,exist_ok=True)
V=R/'verification_20261011';BACK=R/'backups/recovery_phase2_20261011'
DATE='2026-10-12';STAMP=datetime.datetime.now().isoformat(timespec='seconds')
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def readjson(p,default=None):return json.loads(p.read_text('utf-8')) if p.exists() else default
seeds=readjson(S/'seeds.json')['seeds'];bookjson=readjson(F/'content_and_evidence.json');catalog={c[0]:c for c in bookjson['cases']}
groups={
 **{f'TC-F{i:02}':['fast_base'] for i in range(1,9)},
 'TC-M01':['1g'],'TC-M02':['ring'],'TC-M03':['axi'],'TC-M04':['axi_tail'],
 'TC-M05':['regs'],'TC-M06':['rx_fifo'],'TC-M07':['tx'],'TC-M08':['sync'],'TC-M09':['period_100ms'],
 'TC-F09':['byte1_1g'],'TC-F10':['tx_boundaries'],'TC-F11':['ready_recovery'],
 'TC-F12':['regs_multiblock'],'TC-F13':['backend_read','backend_edge_recovery'],
 'TC-F14':['backend_write','backend_edge_recovery','writer_error_recovery'],
 'TC-F15':['recovery_cross_variants'],'TC-R01':['rootport'],'TC-F16':['real_mig']}
directories=['random_regression','random_additional','random_final','random_backend_final','random_resume_20261012']
allrows=[]
for name in directories:
    for r in readjson(V/name/'summary.json',[]):
        r=dict(r);r['evidence_directory']=name;r['evidence_log']='verification_20261011/'+name+'/'+r['log'];allrows.append(r)
def selected(group):
    for name in reversed(directories):
        rows=[r for r in allrows if r['case']==group and r.get('seed') in seeds and r['evidence_directory']==name]
        if len(rows)==len(seeds):return rows
    return [r for r in allrows if r['case']==group]
root_dir='rootport_resume_20261012' if (V/'rootport_resume_20261012/summary_current.json').exists() else 'rootport_current_o0'
root=readjson(V/root_dir/'summary_current.json',{})
mig=readjson(V/'real_mig_recovery/summary.json',{})
resume_mig=[r for r in allrows if r['case']=='real_mig' and r['evidence_directory']=='random_resume_20261012']
if len(resume_mig)==3:
    mig=dict(mig,runs=resume_mig,status='PASS' if all(r['status']=='PASS' for r in resume_mig) else 'FAIL')
selected_groups={g:selected(g) for gs in groups.values() for g in gs if g not in ['rootport','real_mig']}
for retry in readjson(V/'random_period_retry/summary.json',[]):
    retry=dict(retry,evidence_log='verification_20261011/random_period_retry/'+retry['log'],evidence_directory='random_period_retry')
    selected_groups['period_100ms']=[r for r in selected_groups['period_100ms'] if r.get('seed')!=retry['seed']]+[retry]
details={
 'TC-M01':'活动模块为两拍6 Byte；目标四拍16 Byte检查失败。',
 'TC-F09':'adapter配置链路进入实际活动1G模块；四拍输出检查失败，不能沿用归档四拍通过。',
 'TC-M08':'仅独立RTDS_SYNC_TIMEBASE模块：250MHz、500/12500拍；当前活动TOP未接入同源节拍。',
 'TC-M09':'生产250MHz、PERIOD_CYCLES=25000000；每轮400ms、四次到期、三个100ms间隔；分别检查空块、在途B、非空块及奇数尾字的可读先后。',
 'TC-F11':'真实INFO_EXCH/adapter/CONV_DMA/后端，RAM和LocalLink邻居为模型；恢复成功路径不使用H.reset。',
 'TC-F13':'基础读事务1/2/256拍、ID=0/1/31及随机ID、8/64/400拍保持检查三种子通过；新增同拍/重复中断检查原始FAIL，预期边界需复核，不据此直接判定RTL缺陷。',
 'TC-F14':'基础AW/W/B收尾及真实writer/CONV_DMA错误双字、尾字计数检查三种子通过；新增同拍/重复中断检查原始FAIL，判据待复核。DDR物理数据保留不在模型结论内。',
 'TC-F15':'100次循环×三个种子；覆盖PCIe、DDR、channel、lane四种失效与六个位置的组合，用户停钟/慢钟后恢复；新任务4拍、2帧、2个Byte1、1完成。',
 'TC-R01':'真实XDMA/Root Port、当前PCIe与共享design_1；后端DDR4_SIM_RAM，封块周期25000拍（100us）。本轮工具展开状态与历史功能结果分列。',
 'TC-F16':'单独真实MIG/Clock Converter/DDR颗粒模型只是子场景；真实XDMA+共享互连+MIG同时中断恢复及板级项目仍缺组合证据。'}
cases=[]
for cid,row in catalog.items():
    gs=groups[cid];runs=[]
    if cid=='TC-R01':
        for rr in root.get('runs',[]):runs.append(dict(rr,evidence_log='verification_20261011/'+root_dir+'/'+rr['log']))
        rs=root.get('status','运行中')
        status='通过' if rs=='PASS' else ('失败' if rs=='FAIL' else '未执行（展开暂不追踪）')
    elif cid=='TC-F16':
        for rr in mig.get('runs',[]):runs.append(dict(rr,evidence_log=rr.get('evidence_log','verification_20261011/real_mig_recovery/'+rr['log'])))
        status='部分执行' if runs else '未完成'
        details[cid]+='真实MIG子场景状态：'+mig.get('status','运行中')+'。'
    else:
        for g in gs:runs+=selected_groups[g]
        status='通过' if len(runs)==3*len(gs) and all(r['status']=='PASS' for r in runs) else ('失败' if any(r['status']=='FAIL' for r in runs) else '未完成')
        if cid in ['TC-F13','TC-F14']:
            core=[rr for g in gs if g!='backend_edge_recovery' for rr in selected_groups[g]]
            extra=selected_groups.get('backend_edge_recovery',[])
            if len(core)==3*(len(gs)-1) and all(rr['status']=='PASS' for rr in core) and any(rr['status']=='FAIL' for rr in extra):
                status='核心通过／补充判据待复核'
        if cid=='TC-M09' and len(runs)==3:
            functional=True
            for rr in runs:
                text=(R/rr['evidence_log']).read_text('utf-8',errors='replace')
                rr['functional_status']='PASS' if 'TEST PASS tb_period_100ms' in text and not re.search(r'(?m)^\s*(Fatal:|Error:|ERROR:)',text) else 'FAIL'
                functional= functional and rr['functional_status']=='PASS'
            if functional and any(rr.get('exit_code')==-999 for rr in runs):status='功能通过／工具退出超时'
    dates=sorted({str(r.get('finished',r.get('started','2026-10-11')))[:10] for r in runs})
    actual_date='～'.join(dates) if dates else DATE
    conclusion=f'{status}｜{actual_date}｜{row[2]}；'
    if cid=='TC-R01' and not runs:conclusion+='当前真实XDMA展开未完成；历史16 Byte功能结果不代替本轮。'
    elif cid=='TC-F16':conclusion+='真实MIG子场景见报告；共享路径完整故障恢复及板级未完成。'
    else:conclusion+=f'三种子；'+details.get(cid,'数据、地址、长度及用例判据见对应日志。')
    if cid=='TC-M09' and status=='功能通过／工具退出超时':conclusion+='三轮功能检查均结束；种子750318462退出超时，工具门禁未通过。'
    cases.append({'id':cid,'name':row[1],'level':row[2],'status':status,'conclusion':conclusion,'groups':gs,'runs':runs,'note':details.get(cid,'')})
counts=dict(Counter(c['status'] for c in cases))
elabfile=V/'elaborate_final/rtl_elaboration_results.txt'
elabtext=elabfile.read_text('utf-8',errors='replace') if elabfile.exists() else ''
elabpass=bool(re.search(r'^rtds_ddr4_backend rc=0',elabtext,re.M)) and bool(re.search(r'^TOP rc=0',elabtext,re.M))
rollup={'date':DATE,'generated':STAMP,'root':str(R),'seeds':seeds,'case_count':len(cases),'counts':counts,'cases':cases,'selected_groups':selected_groups,'rootport':root,'real_mig':mig,'rtl_elaboration':{'status':'PASS' if elabpass else '未完成','log':'verification_20261011/elaborate_final/vivado.log','note':'synth_design -rtl；IP为stub/blackbox，不构成逻辑综合、IP内部仿真、完整CDC或时序签核'},'synthesis':False,'implementation':False,'board':False}
(REP/'regression_summary.json').write_text(json.dumps(rollup,ensure_ascii=False,indent=2),encoding='utf-8')
issues=readjson(S/'issue_register.json');hist=readjson(S/'history_candidates.json',[])
issues.append({'id':'H22','priority':'P2','state':'新增测试判据待复核，未确认RTL缺陷','topic':'同拍故障补充测试的响应保存边界','finding':'新增tb_backend_edge_recovery三种子均在k=2返回OKAY时失败。测试预期从该拍开始全部SLVERR，但DDR复位请求与后端storage_ready实际下降之间存在同步延迟；此前已经展示的响应按设计必须保留。基础读写恢复三种子均通过。本轮未继续修改判据或RTL，原始FAIL日志保留。','impact':'不能将新增补充检查记为通过，也不能仅凭这条判据认定改后RTL功能失败。','evidence':['verification_20261011/random_resume_20261012/backend_edge_recovery_seed_1417633351.simulate.log','verification_20261011/random_resume_20261012/backend_edge_recovery_seed_750318462.simulate.log','verification_20261011/random_resume_20261012/backend_edge_recovery_seed_360207181.simulate.log','regression/recovery_phase2_20261011/tb_backend_edge_recovery.sv'],'cases':['TC-F13','TC-F14'],'next':'按storage_ready实际下降时刻记录已展示/已接受R拍，核对独立参考队列后修正判据并重跑；如仍违反保持或计数规则再修改RTL。','acceptance':'同拍接受与重复故障检查通过，原始失败及判据变更原因完整保留。'})
closed=[
 ('C01','旧读响应进入新任务及恢复后重新启动失败','已修复并由TC-F11/F15模型回归验证。用户域停止确认、busy清零及8个安静FDMA拍共同决定重新启动；不能据此关闭完整TOP CDC。'),
 ('C02','DDR中断后原事务永久等待或响应在反压期间变化','后端保存ID和长度，读补齐SLVERR、写接收并丢弃剩余W后给错误B，已展示响应保持；TC-F13/F14验证，真实MIG与组合路径边界另列。'),
 ('C03','阶段一闲置fdma_rready导致假超时','回归发现ready=1且无valid被CONV_DMA计作等待；改为无有效任务时仅在fdma_rbusy=1排空。修复前0x2C=0x10证据保留，TC-F11修复后通过。'),
 ('C04','正常完成清理期间紧邻合法任务被拒绝','增加仅正常清理时的一条暂存命令；异常恢复提交仍拒绝。三种子快速集成连续16B/1KiB/32KiB任务通过。'),
 ('C05','随机TX测试台空拍重复上一字','修正测试驱动在每次R接受后撤销valid，不降低数据/地址/帧边界及反压判据；失败日志与修复后结果分别保留。'),
 ('C06','历史真实Root Port的16B被旧512B日志门禁误判','历史功能PASS/errors=0与脚本FAIL分别记录；当前运行使用明确检查标记和错误扫描，展开超时不能归入数据功能失败。'),
 ('C07','旧Aurora OOC DCP导致lane0 UCIO/LOC冲突及遗漏RegisterSlice','历史重新生成IP/修正引用后关闭，当前RTL展开再次核对；其余方法学告警仍需当前实现检查。'),
 ('C08','仿真兼容问题：路径、VIP头文件、forward声明和种子命令','本轮独立工作目录重建源列表，补include、前移声明并使用XSim options文件传递SEED；测试台模块实例由interface改为module。生产端口未因工具兼容问题改变。')]
history_source_map={
 'H01':['新功能','规格','大仿真'],'H02':['SpyGlass'],'H03':['SpyGlass','上板'],'H04':['大仿真','DDR分块'],'H05':['大仿真'],'H06':['SpyGlass','DDR分块'],'H07':['DDR分块','DDR方案'],'H08':['DDR分块','DDR方案'],
 'H09':['DDR分块','DDR方案'],'H10':['DDR分块','新功能'],'H11':['新功能','DDR分块'],'H12':['规格'],'H13':['新功能','Aurora'],'H14':['新功能','上板'],'H15':['大仿真','DDR分块'],
 'H16':['上板','时序'],'H17':['上板','时序'],'H18':['BIST','bist','DDR工程'],'H19':['大仿真'],'H20':['大仿真','DDR分块'],'H21':['新功能']}
for issue in issues:
    # This list identifies related history sources; it does not promote keyword hits to findings.
    issue['history_sources']=[{'path':h['path'],'sha256':h['sha256']} for h in hist if any(k in Path(h['path']).name for k in history_source_map.get(issue['id'],[]))]
    issue['evidence']=[x for x in issue['evidence'] if (Path(x).exists() if re.match(r'^[A-Za-z]:',x) else (R/x).exists())]
    if issue['id']=='H05':issue['current_run']={'rootport':root.get('status','运行中'),'real_mig':mig.get('status','运行中')}
(REP/'historical_issue_register.json').write_text(json.dumps({'date':DATE,'issues':issues,'closed':closed,'history_sources':hist},ensure_ascii=False,indent=2),encoding='utf-8')
mods=[]
for rel in ['RTDS_test_DDR4_64bit.srcs/sources_1/new/pcie_fdma_aurora_brg_adapter.v','rtl/ddr4/rtds_ddr4_backend.sv']:
    p=R/rel;before=R/'backups/recovery_phase1_20261011'/rel
    diff=''.join(difflib.unified_diff(before.read_text('utf-8').splitlines(True),p.read_text('utf-8').splitlines(True),fromfile='phase1_before/'+rel,tofile='current/'+rel))
    (REP/(p.name+'.diff')).write_text(diff,encoding='utf-8')
    mods.append({'path':rel,'before_sha256':sha(before),'after_sha256':sha(p),'diff':p.name+'.diff'})
(REP/'modified_sources.json').write_text(json.dumps(mods,ensure_ascii=False,indent=2),encoding='utf-8')
def source_lines(p,term):
    path=R/p
    return [str(i) for i,l in enumerate(path.read_text('utf-8',errors='replace').splitlines(),1) if term in l][:8] if path.exists() else []
change_sections=[
 ('1. 为什么需要修改','2026-10-10 TC-F11 的六项恢复检查失败：旧任务的底层读事务没有随命令状态结束，恢复后的响应和迟到完成事件可能进入下一任务。阶段一以当前F盘工程修改adapter和DDR后端，采用不复位INFO_EXCH/CONV_DMA/队列的协议收尾；当时仅做静态语法与接口核对，没有仿真或elaborate。本轮按三随机种子重跑现有及新增用例，并保留修改前失败证据。'),
 ('2. 下行恢复如何改变','adapter中断后进入tx_recovering：停止新读请求、禁止旧数据送TX FIFO或产生Byte1事件，继续排空旧读响应。只有fdma_rbusy清零、用户域桥停止确认同步返回、连续8个FDMA时钟没有旧完成事件后重新开放任务。用户时钟暂停时继续等待真实确认，不能用固定延时假设已经停止。0x20清零只表示旧命令结束，不能等同底层事务收尾完成。'),
 ('3. DDR收尾如何改变','每个方向最多一笔在途突发，AR/AW握手时保存ID和LEN+1，读剩余计数为9bit以支持256拍。DDR失效后，读尚未取得的拍返回零数据SLVERR，原ID不变且仅最后一拍RLAST；写继续接收并丢弃剩余W，全部原W收到后给一次SLVERR B。已展示的R/B在READY=0期间保持原数据/ID/RESP/LAST，不能事后把已展示OKAY改成错误。'),
 ('4. 真实MIG接口复位与就绪','主机复位、DDR系统复位或UI接口不可用共同断言转换器两侧及MIG AXI复位；每侧连续运行16个本域时钟后释放，MIG侧释放确认再同步回业务域。业务事务收尾记录不随存储侧复位丢失。0x2C bit1保持校准含义，bit0还要求两侧接口可用且旧错误事务收尾完成，允许暂时不同。0x10只清诊断，不释放块或清掉恢复事务。'),
 ('5. 本轮新增发现及修正','第一处是空闲fdma_rready恒1：CONV_DMA把ready且无valid计为等待，导致未传输时0x2C超时位置位。fallback改为fdma_rbusy，旧事务仍可排空，闲置不再假等待。第二处是正常结束后的清理窗口：紧邻的合法命令原本被拒绝。本轮仅在正常清理中暂存一条地址/长度，清理完成再启动；异常恢复期命令仍拒绝，断链取消暂存。保存参数防止后续寄存器变化覆盖任务。'),
 ('6. 测试台与运行修正','加入真实种子激励：xorshift32改变数据、相位、响应空拍和反压；快速BFM的原生随机通过XSim --sv_seed指定。保留250MHz/25000000生产周期，不以缩时替代。修正随机TX空拍中valid未撤销导致的重复数据，以及2020.2不允许interface内实例化module的测试环境封装；修正forward声明和include，options文件解决Windows批处理拆分SEED=参数的问题。所有通过必须同时满足明确PASS标记、错误扫描与实际检查，不仅依赖进程退出码。'),
 ('7. 代价与功能边界','正常R/B增加一级业务域寄存器，后端每方向并发度限制为一笔突发；共享通路吞吐和板级时序需后续实测。正常清理仅一条暂存，不是通用命令队列。物理DDR重新校准不保证历史数据保留，部分物理写不能回滚，已发往10G对端的部分帧不能撤回。ready始终高但永久无R/B的情况，目前超时只诊断，不能宣称任意停滞均自动恢复。'),
 ('8. 当前活动版本与1G端口说明','活动rtl/rtds_1g_periodic_tx.v仍为两拍6 Byte，XPR引用该活动文件；TOP未连接tick_toggle_2us。external_snapshot的Verilog/SystemVerilog中没有这个端口或该模块定义。四拍16 Byte、tick_toggle_2us接口位于evidence/local_wave_sessions/byte1_1g/sources/rtl归档版本。本次未自动将归档模块覆盖到活动工程，因此TC-M01与TC-F09保留当前失败。独立TC-M08通过不证明TOP已集成2us/50us节拍。')]
def title(doc,text):doc.add_heading(text,0)
def table(doc,heads,rows):
    t=doc.add_table(rows=1, cols=len(heads));t.style='Table Grid'
    for c,s in zip(t.rows[0].cells,heads):c.text=str(s)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    return t
def newdoc(name):
    doc=Document();doc.styles['Normal'].font.name='Microsoft YaHei';doc.styles['Normal'].font.size=Pt(10)
    title(doc,name);doc.add_paragraph(f'核查日期：{DATE}；结果生成：{STAMP}。基线：{R}。')
    return doc
def write_md_doc(stem,md,doc):
    (D/(stem+'.md')).write_text(md,encoding='utf-8');doc.save(D/(stem+'.docx'))
summary_line=f'现有26条用例的当前归档状态：'+ '、'.join(f'{k}{v}条' for k,v in counts.items())+'。随机种子：'+ '、'.join(map(str,seeds))+'。'
doc=newdoc('RTDS恢复修复：修改因果、随机回归与RTL展开')
doc.add_paragraph(summary_line);doc.add_paragraph('通过的适用范围以各条用例模型边界为准。没有运行综合、实现、bitstream生成或板级测试。')
doc.add_paragraph('快速阅读：本次恢复修复的TC-F11、基础TC-F13/F14、TC-F15和连续下行集成检查在三个种子下通过，独立真实MIG恢复也在三个种子下通过。当前1G两条四拍失败属于活动两拍版本差异。新增同拍补充检查的原始FAIL保留，响应展示边界判据待复核，不能宣称全部场景已经关闭。按用户要求，归档后暂停，Root Port展开问题本轮暂不追踪。')
md=['# RTDS恢复修复：修改因果、随机回归与RTL展开','',summary_line,'','基线：`'+str(R)+'`。没有运行综合、实现、bitstream生成或板级测试。','']
for h,text in change_sections:
    doc.add_heading(h,1);doc.add_paragraph(text);md+=['## '+h,'',text,'']
doc.add_heading('9. 用例回归结果与证据',1)
table(doc,['用例','名称 / 层级','结论','直接证据'],[[c['id'],c['name']+' / '+c['level'],c['conclusion'],'\n'.join(r['evidence_log'] for r in c['runs']) or 'verification_20261011/rootport_current_o0/summary_current.json（或子验证记录）'] for c in cases])
md+=['## 9. 用例回归结果与证据','']
for c in cases:
    md+=['### '+c['id']+' '+c['name'],'',c['conclusion'],'']
    for run in c['runs']:md+=['- 种子 '+str(run.get('seed'))+'：'+run['status']+'；`'+run['evidence_log']+'`。']
    md+=['']
doc.add_heading('10. RTL展开与检查边界',1)
doc.add_paragraph('当前最终源码RTL展开：'+rollup['rtl_elaboration']['status']+'。命令为synth_design -rtl，分别展开rtds_ddr4_backend和TOP，Vivado 2020.2 / xcku040-ffva1156-2-i / DDR4_REAL；结果文件：verification_20261011/elaborate_final/rtl_elaboration_results.txt。只读工程中IP以stub/blackbox绑定，未做逻辑综合，不能替代真实IP内部功能、完整TOP CDC和当前版时序检查。')
md+=['## 10. RTL展开与检查边界','','状态：'+rollup['rtl_elaboration']['status']+'。`synth_design -rtl`分别展开后端和TOP；只读IP stub/blackbox绑定。未做逻辑综合，不能替代真实IP内部功能、完整TOP CDC和当前版时序检查。','']
doc.add_heading('11. 文件、备份与复现',1)
table(doc,['源码','修改前SHA256','最终SHA256'],[[m['path'],m['before_sha256'],m['after_sha256']] for m in mods])
for text in ['业务源码修改前副本：backups/recovery_phase1_20261011；阶段一至阶段二过渡副本：backups/recovery_phase2_20261011。差异和本轮摘要：reports/recovery_phase2_20261011。','各次构建独立目录保留source_hashes.json、编译、elaborate、逐种子日志与summary；修复前失败不删除、不挪作本轮通过。main构建中的adapter早于最终正常清理暂存修正，该目录只用于与该修改无关的独立模块及生产周期结果。','新增测试与可复现脚本：regression/recovery_phase2_20261011。波形及匹配源码：evidence/recovery_phase2_20261011。原始WDB和VCD范围见波形manifest，VCD没有记录的数组对象不能从查看器推断历史值。','证据对应与完整清单：reports/recovery_phase2_20261011/regression_summary.json、modified_sources.json、historical_issue_register.json、delivery_checks.json。']:
    doc.add_paragraph(text);md+=[text,'']
write_md_doc('RTDS_恢复修复_修改因果_随机回归与展开_20261012','\n'.join(md),doc)
doc=newdoc('RTDS历史遗留问题与后续工作清单')
intro='结合F盘活动工程、工作区资料与E:\\codex_history_md\\PCIE-CPU的37份历史Markdown记录整理。按语义区分当前缺陷、设计检查、证据缺口、参数限制、可选扩展和已关闭事项；未将所有历史提案视为当前要求。状态以本轮可核实的证据为准。'
doc.add_paragraph(intro);md=['# RTDS历史遗留问题与后续工作清单','',intro,'']
table(doc,['编号','优先级','事项','当前状态'],[[i['id'],i['priority'],i['topic'],i['state']] for i in issues])
for i in issues:
    doc.add_heading(i['id']+' '+i['topic'],1);md+=['## '+i['id']+' '+i['topic'],'']
    fields=[('优先级 / 当前状态',i['priority']+' / '+i['state']),('核查发现',i['finding']),('影响',i['impact']),('后续处理',i['next']),('关闭条件',i['acceptance']),('关联用例','、'.join(i['cases']) or '独立设计检查或扩展，不计入现有用例通过'),('直接资料','\n'.join(i['evidence']) or '关联历史记录及源码清单见末尾')]
    for label,text in fields:doc.add_paragraph(label+'：'+text);md+=[label+'：'+text,'']
    for h in i.get('history_sources',[])[:6]:doc.add_paragraph('关联历史：'+h['path']);md+=['关联历史：`'+h['path']+'`。']
    md+=['']
doc.add_heading('已关闭与已明确的历史事项',1);md+=['## 已关闭与已明确的历史事项','']
for cid,h,text in closed:doc.add_heading(cid+' '+h,2);doc.add_paragraph(text);md+=['### '+cid+' '+h,'',text,'']
doc.add_heading('建议后续顺序与完成判据',1)
priority_text='先统一1G活动版本和TOP节拍/GPIO接口；再关闭完整TOP的复位、IRQ及GT边界CDC检查，并完成真实共享XDMA/AXI/MIG的故障恢复组合验证。随后对DDR复位数据失效策略、普通读错误传播、永久停滞恢复和CPU确认协议作明确设计决定。最后按最终源码完成新综合/实现/时序及板级验证。高地址与多子机扩展按产品需要单列，不能用默认配置测试替代。'
doc.add_paragraph(priority_text);md+=['## 建议后续顺序与完成判据','',priority_text,'']
doc.add_heading('资料来源与追溯',1)
table(doc,['资料','SHA256 / 说明'],[[h['path'],h['sha256']] for h in hist])
md+=['## 资料来源与追溯','','历史来源的文件SHA256、问题到用例对应及各次构建证据见 `reports/recovery_phase2_20261011/historical_issue_register.json`。','']
for h in hist:md+=['- `'+h['path']+'`；SHA256='+h['sha256']+'。']
write_md_doc('RTDS_历史遗留问题与后续工作清单_20261012','\n'.join(md),doc)
def backup(p):
    b=BACK/p.relative_to(R)
    if not b.exists():b.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,b)
xp=F/'RTDS_功能规格_测试点_详细用例_20261011.xlsx';backup(xp);w=load_workbook(xp)
rows={r[0].value:r for r in w['测试用例'].iter_rows(min_row=2)}
for c in cases:
    rows[c['id']][7].value=c['conclusion']
    if c['id']=='TC-R01':rows[c['id']][1].value='真实Root Port与XDMA当前回归'
specrows={r[0].value:r for r in w['设计规格'].iter_rows(min_row=2)}
specrows['S-897'][4].value='异常恢复期命令拒绝并结束；正常完成清理中允许暂存一条合法下一命令，清理完成按锁定参数启动；断链取消暂存。0x20=0不等同恢复完成。2026-10-11三种子模型回归，完整共享路径边界见TC-F16。'
specrows['S-15'][4].value='2026-10-11已修改adapter与DDR后端；TC-F11/F13/F14/F15结果见本轮回归报告，历史2026-10-10失败保留。真实IP与完整TOP CDC、板级边界单列。'
pointrows={r[0].value:r for r in w['测试点'].iter_rows(min_row=2)}
pointrows['P-32'][4].value='未就绪不发新访问；0x20结束旧命令后旧读继续排空；恢复完成的新任务数据/地址/长度/计数准确，无旧Byte1或完成事件。三种子结果见TC-F11/F15。'
pointrows['P-47'][4].value='异常清理无新请求和Byte1，期间命令拒绝；新32Byte任务4拍、2帧、1完成。正常完成清理允许暂存一条合法下一命令；完成恢复条件后再启动。'
w.save(xp);backup(F/'content_and_evidence.json')
for key,sheet in [('specs','设计规格'),('points','测试点'),('cases','测试用例')]:bookjson[key]=[list(r) for r in list(w[sheet].values)[1:]]
bookjson['phase2_revision']={'date':DATE,'seeds':seeds,'counts':counts,'source_manifest':'reports/recovery_phase2_20261011/modified_sources.json','evidence':'reports/recovery_phase2_20261011/regression_summary.json','full_top_cdc':'尚未关闭','synthesis':False,'implementation':False,'board':False}
(F/'content_and_evidence.json').write_text(json.dumps(bookjson,ensure_ascii=False,indent=2),encoding='utf-8')
for p in [D/'RTDS_主机FPGA_PCIe_DDR4_Aurora_详细设计方案_V1.0_20261002.docx',F/'RTDS_功能规格与验证核查报告_20261011.docx']:
    backup(p);old=Document(p)
    # Replace only phase-two addendum on regeneration; earlier reports stay historical.
    marker='2026-10-11～12第二阶段：随机回归与展开记录'
    remove=False
    for node in list(old.element.body):
        if node.tag.endswith('}sectPr'):continue
        if marker in ''.join(node.itertext()):remove=True
        if remove:old.element.body.remove(node)
    if p.parent==D:
        for para in old.paragraphs:
            if para.text.startswith('2026-10-11修订：RD_START'):
                para.text='2026-10-11第二阶段修订：空闲且tx_recovering=0、fdma_rbusy=0、链路及参数有效时接受RD_START。正常完成的清理窗口可暂存一条合法下一命令并锁定地址与长度；异常恢复期命令仍拒绝，断链取消暂存。旧读排空，fdma_rready在非有效任务时仅随busy允许排空，避免闲置假超时；旧清理期间禁止Byte1。0x20清零不代表底层恢复完成。'
            if para.text.startswith('工程基线为') or para.text.startswith('原V1.0依据'):
                para.text='原V1.0依据2026-10-02远程工程编写。本版保留原方案结构，2026-10-11以F:\\RTDS_test_DDR4_64bit_batch_dev活动工程修订恢复逻辑，并完成随机回归与RTL展开，结果和实际边界见末尾第二阶段记录。本轮未执行新综合、实现或板级验证。'
            if para.text.startswith('对应S-15、S-897、P-23/P-32/P-47'):
                para.text='对应S-15、S-897、P-23/P-32/P-47及TC-F11/TC-F15：中断或正常结束进入桥停止清理，旧R按busy排空且不进入TX/Byte1。正常清理暂存一条合法下一命令，异常恢复期提交拒绝；完成用户域停止确认、busy清零和8个安静FDMA拍后启动。原外部端口和CONV_DMA/INFO_EXCH/BD/XPR保持。'
    old.add_heading(marker,1);old.add_paragraph(summary_line)
    old.add_paragraph('新增恢复检查结果由本轮verification_20261011各构建日志支持。真实Root Port：'+root.get('status','未完成')+'；真实MIG独立子验证：'+mig.get('status','未完成')+'；TC-F16完整共享故障恢复及板级结果未完成。完整TOP CDC仍沿用10月10日未关闭记录，不能用功能仿真替代。')
    old.add_paragraph('当前最终RTL展开：'+rollup['rtl_elaboration']['status']+'；synth_design -rtl分别展开TOP与后端，IP stub/blackbox边界见elaborate_final。未运行新逻辑综合或布局布线；历史bitstream及时序结果不适用于本轮修改后的RTL。')
    old.add_paragraph('正常清理单条暂存、空闲ready假超时修正及全部因果见docs/RTDS_恢复修复_修改因果_随机回归与展开_20261012.docx。22项后续事项（含本轮新增判据复核项）及8项历史关闭/明确记录见docs/RTDS_历史遗留问题与后续工作清单_20261012.docx。对应摘要和源码哈希见reports/recovery_phase2_20261011。归档后按用户要求暂停。')
    old.save(p)
checks=[]
for p in [D/'RTDS_恢复修复_修改因果_随机回归与展开_20261012.docx',D/'RTDS_历史遗留问题与后续工作清单_20261012.docx',D/'RTDS_主机FPGA_PCIe_DDR4_Aurora_详细设计方案_V1.0_20261002.docx',F/'RTDS_功能规格与验证核查报告_20261011.docx']:
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None;names=set(z.namelist())
        for n in names:
            if n.endswith(('.xml','.rels')):ET.fromstring(z.read(n))
            if n.endswith('.rels'):
                base=posixpath.dirname(posixpath.dirname(n))
                for rel in ET.fromstring(z.read(n)):
                    if rel.attrib.get('TargetMode')=='External':continue
                    t=rel.attrib['Target'];target=posixpath.normpath(posixpath.join(base,t)) if not t.startswith('/') else t[1:]
                    assert target in names,(n,target)
        oldp=BACK/p.relative_to(R)
        if oldp.exists():
            with zipfile.ZipFile(oldp) as oldzip:
                for n in oldzip.namelist():
                    if n.startswith('word/media/') and not n.endswith('/'):assert z.read(n)==oldzip.read(n)
    document=Document(p);checks.append({'path':str(p.relative_to(R)),'zip_xml_relationships':'PASS','tables':len(document.tables),'rendered':False,'sha256':sha(p)})
ids={sheet.title:[r[0] for r in list(sheet.values)[1:] if r[0]] for sheet in w}
for title,idlist in ids.items():assert len(idlist)==len(set(idlist)),title
for r in list(w['测试点'].values)[1:]:
    for value in re.findall(r'S-\d+',str(r[5])):assert value in ids['设计规格'],value
    for value in re.findall(r'TC-[A-Z]\d+',str(r[6])):assert value in ids['测试用例'],value
delivery={'date':DATE,'docx':checks,'workbook_semantic_references':'PASS','row_counts':{k:len(v) for k,v in ids.items()},'format_checks':False,'pdf_png_rendering':False,'regression_counts':counts,'elaborate':rollup['rtl_elaboration'],'open_issues':len(issues)}
(REP/'delivery_checks.json').write_text(json.dumps(delivery,ensure_ascii=False,indent=2),encoding='utf-8')
for name in ['run_random.py','run_rootport.py','run_real_mig.py','seeds.json']:
    shutil.copy2(S/name,R/'regression/recovery_phase2_20261011'/name)
print(json.dumps({'counts':counts,'documents':[c['path'] for c in checks],'elaborate':rollup['rtl_elaboration']['status']},ensure_ascii=False,indent=2))
