from pathlib import Path
import json,zipfile,xml.etree.ElementTree as ET,shutil,hashlib
from docx import Document
from docx.shared import Pt
R=Path('F:/RTDS_test_DDR4_64bit_batch_dev');D=R/'docs';P=R/'reports/recovery_phase2_20261011';S=Path(__file__).parent
j=json.loads((D/'RTDS_functional_20261011/content_and_evidence.json').read_text('utf-8'))
summary=json.loads((P/'regression_summary.json').read_text('utf-8'))
title='RTDS 最新设计方案（精简版）'
sections=[
 ('1. 基线与实现边界','日期：2026-10-12。以F:\\RTDS_test_DDR4_64bit_batch_dev活动XPR和源码为准；保留原V1.0详细方案。本版概括当前实现，全部端口、位宽和逐项规格仍以RTDS_functional_20261011工作簿为准。当前1G活动版本为两拍6 Byte、周期参数63个用户时钟；四拍16 Byte、Byte6=0xF0、同源2 us和50 us GPIO属于待统一的目标，当前TOP未接入tick_toggle_2us。'),
 ('2. 数据路径','下行：CPU完成H2C写DDR后，通过INFO_EXCH设置地址/长度并写0x20=1；adapter经CONV_DMA读取，发送桥通过异步FIFO送Aurora 10G。每16 Byte为两拍64bit，首拍SOF、末拍EOF、REM=7，以valid/ready推进。首次合法Byte1编号1～10用于1G参数配置；后续配置不覆盖首次锁定值。上行：每个10G有效拍保存低32bit，高32bit不进入业务数据；先到字置64bit低半，后到字置高半，经writer写DDR，再由CPU逐块C2H读取和确认。'),
 ('3. 地址、容量与块规则','下行窗口0x00008000～0x0000FFFF，地址16 Byte对齐，长度为正16 Byte倍数且不超过32 KiB；整个请求必须落在窗口内。上行窗口0x00100000～0x002FFFFF，8个256 KiB环形块；未确认的可读块不得覆盖。RX FIFO生产深度2048×64bit，线路不能靠业务ready无限回压，拥塞时统计丢失的32bit字。生产封块周期为250 MHz、25000000拍，即100 ms；周期到期不等于块可读，必须等成功写响应与奇数尾字完成。正常双字WSTRB=0xFF、有效长度+8；尾字WSTRB=0x0F、长度+4；错误写不增加有效长度。'),
 ('4. 时钟、复位与事务恢复','业务/FDMA时钟250 MHz；Aurora各自用户时钟由协议核输出，跨域采用FIFO或明确同步结构。1G初始化由100 MHz经BUFGCE_DIV四分频得到25 MHz。DDR接口复位共同覆盖AXI转换器两侧及MIG AXI，各侧运行16个本域时钟后同步释放，MIG侧完成再同步回业务域；校准与业务ready分开。后端每方向最多一笔在途突发，记录ID和LEN+1：存储失效时，读补齐尚未取得的零数据SLVERR，写接收并丢弃剩余W后给一次SLVERR B。已经展示且被反压的R/B保持原字段，不能因失效改写；正常R/B各增加一级寄存。'),
 ('5. 命令与跨任务隔离','下行结束或中断后停止新读请求，旧读按busy排空且不进入TX/Byte1。重新启动要求读busy清零、用户域桥停止确认返回、连续8个FDMA拍无旧完成事件。正常完成的清理窗口允许暂存一条合法下一命令并锁定地址/长度；异常恢复期命令拒绝，断链取消暂存。空闲fdma_rready不再恒1，避免无事务时假超时。0x20清零只代表命令结束，不等于DDR写成功、块可读、CPU读完或下游执行完成。'),
 ('6. CPU使用顺序','下行：H2C完成→写0x24/0x28→写0x20=1→等待命令结束并检查错误状态。上行：先看0x30数量，读取0x18/0x1C/0x34队头信息，完成对应C2H后仅写一次0x14=1。读0x10只清通知和相关诊断，不释放块；0x38累计丢字不随读清。接口接受、DDR成功提交、块可读、CPU确认是四个不同阶段。'),
 ('7. 本次验证与剩余边界','本次恢复修复的TC-F11、基础TC-F13/F14、TC-F15及连续下行集成检查三个种子通过；独立真实MIG/Clock Converter/DDR颗粒模型恢复也三个种子通过。TOP与后端RTL展开通过，使用工程现有IP stub/DCP，不构成新综合或实现。新增同拍/重复中断测试有原始FAIL，已展示响应与storage_ready下降的判据边界待复核；当前不据此认定RTL缺陷，也不记补充检查通过。1G两条四拍失败属于活动版本差异。完整TOP CDC、真实XDMA+共享互连+MIG组合故障恢复、物理时序和板级结果尚未关闭。DDR物理复位不保证旧块数据保留；永久无响应而ready仍高时，超时目前仅诊断。')]
regs=[]
for row in j['specs']:
    if row[0] in [f'S-{n}' for n in range(32,47)]:regs.append([row[2].split(' ')[0],row[2].split(' ',1)[1],row[3]])
doc=Document();doc.styles['Normal'].font.name='Microsoft YaHei';doc.styles['Normal'].font.size=Pt(10);doc.add_heading(title,0)
md=['# '+title,'']
for heading,text in sections:
    doc.add_heading(heading,1);doc.add_paragraph(text);md+=['## '+heading,'',text,'']
    if heading.startswith('6.'):
        t=doc.add_table(rows=1,cols=3);t.style='Table Grid'
        for c,s in zip(t.rows[0].cells,['偏移','寄存器','核心语义']):c.text=s
        md+=['| 偏移 | 寄存器 | 核心语义 |','|---|---|---|']
        for row in regs:
            for c,s in zip(t.add_row().cells,row):c.text=s
            md+=['| '+' | '.join(row)+' |']
        md+=['']
doc.add_heading('资料入口',1)
for text in ['详细方案：RTDS_主机FPGA_PCIe_DDR4_Aurora_详细设计方案_V1.0_20261002.docx。','规格/测试点/用例：RTDS_functional_20261011/RTDS_功能规格_测试点_详细用例_20261011.xlsx。','修改与验证：RTDS_恢复修复_修改因果_随机回归与展开_20261012.docx。','后续问题：RTDS_历史遗留问题与后续工作清单_20261012.docx；日志与源码哈希见../reports/recovery_phase2_20261011/regression_summary.json。']:
    doc.add_paragraph(text);md+=[text,'']
stem='RTDS_最新设计方案_精简版_20261012';doc.save(D/(stem+'.docx'));(D/(stem+'.md')).write_text('\n'.join(md),encoding='utf-8')
with zipfile.ZipFile(D/(stem+'.docx')) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith(('.xml','.rels')):ET.fromstring(z.read(name))
index='''# 仿真与设计资料入口（2026-10-12）

资料直接保存在工程目录，没有制作压缩包。当前工作按用户要求暂停。

## 优先阅读

1. [最新设计方案·精简版](RTDS_最新设计方案_精简版_20261012.docx)：当前实现、寄存器与恢复规则。
2. [修改因果与随机回归](RTDS_恢复修复_修改因果_随机回归与展开_20261012.docx)：通过结果、原始失败和模型边界。
3. [历史遗留与后续清单](RTDS_历史遗留问题与后续工作清单_20261012.docx)：22项后续事项，含新增判据复核。
4. [规格、测试点与用例](RTDS_functional_20261011/RTDS_功能规格_测试点_详细用例_20261011.xlsx)：编号不变，执行结果已更新。
5. [原V1.0详细方案](RTDS_主机FPGA_PCIe_DDR4_Aurora_详细设计方案_V1.0_20261002.docx)：保留原结构，追加本轮修订。

上述新报告另有同名Markdown，便于检索。历史文档保留，最新结果以上述入口为准。

## 证据与复现

- `../reports/recovery_phase2_20261011/regression_summary.json`：26条用例到逐种子日志的对应。
- `../reports/recovery_phase2_20261011/modified_sources.json`：两个业务RTL的前后哈希与差异。
- `../verification_20261011/random_resume_20261012`：补充运行；基础读写及真实MIG通过，新增同拍检查原始FAIL保留。
- `../verification_20261011/elaborate_final`：最终源码RTL展开结果。
- `../regression/recovery_phase2_20261011`：新增测试台与运行脚本。
- `../evidence/local_wave_sessions`：已有波形及其配套源码；本轮新增修复波形尚未补齐，不能把旧归档波形当作本轮源码证据。

## 如何理解结果

恢复修复的核心检查三个种子通过。1G四拍两条失败是活动两拍版本与设计目标的差异。生产100ms的三轮功能检查均通过，其中一轮工具退出超时单列。新增同拍补充检查的判据待复核，未确认RTL缺陷；不能宣称所有场景全部关闭。Root Port展开暂不追踪，完整CDC与板级验证仍待后续。
'''
(D/'仿真与设计资料入口_20261012.md').write_text(index,encoding='utf-8')
tools=P/'tools';tools.mkdir(exist_ok=True)
manifest=[]
for p in S.iterdir():
    if p.is_file():
        dest=tools/p.name;shutil.copy2(p,dest)
        assert hashlib.sha256(p.read_bytes()).digest()==hashlib.sha256(dest.read_bytes()).digest()
        manifest.append({'name':p.name,'sha256':hashlib.sha256(dest.read_bytes()).hexdigest()})
(tools/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(P/'pause_status.json').write_text(json.dumps({'date':'2026-10-12','status':'paused_at_user_request','native_simulation_processes_expected':0,'documents_in_docs':True,'archive_created':False,'pending':['新增同拍判据复核','1G版本统一','完整TOP CDC','共享真实IP组合验证','新波形','板级验证']},ensure_ascii=False,indent=2),encoding='utf-8')
print('Created '+str(D/(stem+'.docx')));print('Created docs/仿真与设计资料入口_20261012.md');print('Preserved '+str(len(manifest))+' helper files in reports/recovery_phase2_20261011/tools')
