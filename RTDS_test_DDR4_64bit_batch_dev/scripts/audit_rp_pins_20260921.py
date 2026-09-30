"""Compare existing placed IO report with local RP board pin references."""
from pathlib import Path
import csv
import json
import re
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
out = root / 'reports/board_audit_20260921'
out.mkdir(parents=True, exist_ok=True)
board = {}
workbook = load_workbook(root / 'docs/RPT205-1-001引脚定义.xlsx', read_only=True, data_only=True)
for sheet in workbook:
    for rownum, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if rownum > 1 and row[4]:
            board[str(row[4])] = dict(net=row[2], standard=row[6], row=rownum)
ddr_board = {}
for row in csv.reader((root / 'docs/RP205-1-001 （DDR4）.txt').read_text(encoding='utf-8-sig').splitlines()):
    if len(row) >= 2:
        ddr_board[row[0]] = row[1]

placed = []
for line in (root / 'RTDS_test_DDR4_64bit.runs/impl_1/TOP_io_placed.rpt').read_text().splitlines():
    fields = [x.strip() for x in line.split('|')]
    if len(fields) > 14 and fields[2] and fields[2] != 'Signal Name':
        placed.append(dict(pin=fields[1], port=fields[2], standard=fields[6],
                           board_net=ddr_board.get(fields[1], board.get(fields[1], {}).get('net', 'UNLISTED')),
                           board_standard=board.get(fields[1], {}).get('standard', ''),
                           board_row=board.get(fields[1], {}).get('row', '')))

def expected_ddr(port):
    m = re.fullmatch(r'c0_ddr4_(dq|dm_dbi_n|dqs_c|dqs_t)\[(\d+)\]', port)
    if m:
        bus, i = m.group(1), int(m.group(2))
        if bus == 'dq':
            return f'DDR4-DQ{"L" if i % 16 < 8 else "U"}{i % 8}-{chr(65+i//16)}'
        half, chip = ('L' if i % 2 == 0 else 'U'), chr(65+i//2)
        return f'DDR4-DM{half}-{chip}' if bus == 'dm_dbi_n' else f'DDR4-DQS{half}-{"N" if bus == "dqs_c" else "P"}-{chip}'
    m = re.fullmatch(r'c0_ddr4_adr\[(\d+)\]', port)
    if m:
        i = int(m.group(1))
        return {14:'DDR4-A14-WEN',15:'DDR4-A15-CASN',16:'DDR4-A16-RASN'}.get(i, f'DDR4-A{i}')
    return {'c0_ddr4_ba[0]':'DDR4-BA0', 'c0_ddr4_ba[1]':'DDR4-BA1',
            'c0_ddr4_bg[0]':'DDR4-BG0','c0_ddr4_ck_t[0]':'DDR4-CLKP',
            'c0_ddr4_ck_c[0]':'DDR4-CLKN','c0_ddr4_cs_n[0]':'DDR4-CS0N',
            'c0_ddr4_cke[0]':'DDR4-CKE','c0_ddr4_odt[0]':'DDR4-ODT0',
            'c0_ddr4_act_n':'DDR4-ACTN','c0_ddr4_reset_n':'DDR4-RESETN',
            'clk_p':'SYSCLK-P','clk_n':'SYSCLK-N'}.get(port)

checked = [p for p in placed if expected_ddr(p['port'])]
errors = [dict(p, expected=expected_ddr(p['port'])) for p in checked if p['board_net'] != expected_ddr(p['port'])]
with (out / 'placed_io_vs_rp.csv').open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(placed[0]))
    writer.writeheader()
    writer.writerows(placed)
result = dict(placed_port_count=len(placed), ddr_and_sysclk_checked=len(checked),
              ddr_and_sysclk_mismatches=errors,
              duplicate_placed_pins=sorted({p['pin'] for p in placed if sum(q['pin']==p['pin'] for q in placed)>1}))
(out / 'pin_check_summary.json').write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
print(json.dumps(result, ensure_ascii=True))
