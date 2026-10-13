"""Make a seven-case, per-figure annotated Word draft from local VCD evidence."""
from pathlib import Path
import json
import re
import shutil
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from wave_figure_notes import FIGURE_NOTES, signal_role


ROOT = Path(__file__).resolve().parent.parent
EVIDENCE = ROOT / "evidence" / "local_wave_sessions"
MARKDOWN = EVIDENCE / "WCFG_ring_axi_regs_观察链与时间刻度.md"
SUPPLEMENT = EVIDENCE / "WCFG_其余四组观察链与时间刻度.md"
FIGURES = EVIDENCE / "word_wave_captures"
OUTPUT = EVIDENCE / "WCFG_ring_axi_regs_观察链与时间刻度_配图_20260928.docx"
OLDER_OUTPUT = EVIDENCE / "WCFG_ring_axi_regs_观察链与时间刻度_配图.docx"


def set_font(style, name, points):
    style.font.name = name
    style.font.size = Pt(points)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|\x60[^\x60]+\x60|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, content):
    for part in TOKEN.split(content):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith(chr(96)) and part.endswith(chr(96)):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, target = part[1:-1].split("](", 1)
            paragraph.add_run(label)
            run = paragraph.add_run(" (" + target + ")")
            run.font.size = Pt(8)
        else:
            paragraph.add_run(part)


def add_table(document, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
    if not rows:
        return
    columns = max(map(len, rows))
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            add_inline(p, value)
            for run in p.runs:
                run.font.size = Pt(7.5 if columns > 4 else 8.5)
                if i == 0:
                    run.bold = True


def add_figure_signal_table(document, figure):
    """Explain every selected trace, including relevant stable controls."""
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    for cell, label in zip(table.rows[0].cells, ("图中信号", "窗口跳变", "设计含义 / 读法")):
        cell.text = label
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for trace in figure["transitions"]:
        role = signal_role(figure["case"], trace["path"])
        if not role:
            raise RuntimeError("Missing signal explanation: " + trace["path"])
        relative = trace["path"].split("/", 2)[2]
        if not relative.startswith(("bus/", "U_DUT/")):
            relative = "top/" + relative
        row = table.add_row().cells
        row[0].text = relative
        row[1].text = str(trace["count"]) if trace["count"] else "0（本窗保持）"
        row[2].text = role
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(7.5)
    widths = (Inches(2.4), Inches(.9), Inches(7.0))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def build():
    manifest = json.loads((FIGURES / "manifest.json").read_text(encoding="utf-8"))
    by_section = {}
    for figure in manifest:
        image = FIGURES / figure["file"]
        if not image.is_file() or image.stat().st_size < 1000:
            raise RuntimeError("Missing or empty screenshot: " + str(image))
        if figure["key"] not in FIGURE_NOTES:
            raise RuntimeError("Missing per-figure design note: " + figure["key"])
        by_section.setdefault((figure["case"], figure["section"]), []).append(figure)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(.62)
    section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.62)
    section.right_margin = Inches(.62)

    set_font(doc.styles["Normal"], "Microsoft YaHei", 9.5)
    set_font(doc.styles["Title"], "Microsoft YaHei", 18)
    set_font(doc.styles["Heading 1"], "Microsoft YaHei", 14)
    set_font(doc.styles["Heading 2"], "Microsoft YaHei", 12)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.core_properties.title = "七组 WCFG 波形：逐图信号与设计观察链（配图草稿）"
    doc.core_properties.subject = "41 个 VCD 实测时间窗口与 RTL / 测试台对照"

    group = None
    number = None
    used = []
    figure_number = 0

    def insert_figures():
        nonlocal number, figure_number
        if group is None or number is None:
            return
        for figure in by_section.get((group, number), []):
            figure_number += 1
            used.append(figure["key"])
            a, b = figure["rangeNs"]
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            run = p.add_run(
                "图 {:02d}  {} 第{}节  {}–{} ns；光标 {} ns。{}".format(
                    figure_number, group, number, a, b, figure["cursorNs"], figure["caption"]
                )
            )
            run.bold = True
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.add_run().add_picture(str(FIGURES / figure["file"]), width=Inches(10.3))
            cause, conclusion = FIGURE_NOTES[figure["key"]]
            p = doc.add_paragraph()
            p.add_run("设计条件：").bold = True
            p.add_run(cause)
            p = doc.add_paragraph()
            p.add_run("波形判断：").bold = True
            p.add_run(conclusion)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run("本图信号逐项解读（跳变次数不等于握手次数）：").bold = True
            add_figure_signal_table(doc, figure)
        number = None

    base = MARKDOWN.read_text(encoding="utf-8-sig").splitlines()
    additional = SUPPLEMENT.read_text(encoding="utf-8-sig").splitlines()
    boundary = next(i for i, line in enumerate(base) if line.startswith("## 证据边界"))
    final_boundary = [
        "## 证据边界和可重复复核",
        "七组的证据由三层构成：VCD 给出选定时间窗内的实际值与事件顺序；RTL 给出状态更新与握手条件；测试台断言、逐拍/逐字节比较和 PASS 日志补足单张截图不能展示的全程结果。各组是模块级仿真，未在同一份 VCD 中覆盖 CPU、PCIe、DDR4 物理器件和 Aurora 线路的完整集成链路。",
        "ring 的 dropped_frames 与 rx_fifo 的 dropped_frames 都按无法存储的 32bit 有效字计；regs 的 ACK 输出与 ring 的 cpu_ack 在各自测试台中分别激励，设计连线可参照本文件的 CPU 确认专题，但不能把两张图当作同一次跨模块握手。axi_tail 证明 CONV_DMA 与行为 RAM 的 WSTRB 按字节写，不能替代板上 DDR 实测。",
        "图中仅选本次观察链需要的信号，七个离线 HTML 各自保留完整信号清单、波形和源码。时间单位均为 ns；若重跑仿真覆盖 VCD，应重新生成截图并复核时间点。复核可用同目录 trace_probe.py 按 ns 提取跳变，也可运行脚本 capture_ring_axi_regs_wave_figures.cjs 和 build_ring_axi_regs_wave_docx.py 重建本文。",
    ]
    lines = base[:boundary] + additional + final_boundary
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith("# "):
            insert_figures()
            doc.add_paragraph("七组 WCFG 波形观察链与时间刻度（逐图详解）", style="Title")
            p = doc.add_paragraph(
                "配图为离线 HTML 对现有 VCD 的后台截图；每张图仅摘出当前观察链需要的信号。"
                "时间轴和黄色光标均以 ns 表示，完整信号仍在对应 HTML 中。"
                "每图下面的表含静态控制，0 次跳变不表示该信号无效；模块级仿真不代替整机验证。"
            )
            p.paragraph_format.space_after = Pt(10)
        elif line.startswith("## "):
            insert_figures()
            match = re.match(r"## (ring|axi|regs|1g|tx|rx_fifo|axi_tail)：", line)
            group = match.group(1) if match else None
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            insert_figures()
            doc.add_paragraph(line[4:], style="Heading 2")
        else:
            match = re.match(r"^\*\*(\d+)\.\s+(.+?)\*\*(.*)$", line)
            if match and group:
                insert_figures()
                number = int(match.group(1))
                p = doc.add_paragraph()
                p.paragraph_format.keep_together = True
                p.add_run(match.group(1) + ". " + match.group(2)).bold = True
                add_inline(p, match.group(3))
            elif line.startswith("- "):
                add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
            else:
                add_inline(doc.add_paragraph(), line)
        i += 1
    insert_figures()

    if len(used) != len(manifest) or set(used) != {x["key"] for x in manifest}:
        raise RuntimeError("Not every screenshot was inserted into the Word document")
    doc.save(OUTPUT)

    # DOCX is a ZIP package. Check only its integrity and image count; do not render Word.
    with zipfile.ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise RuntimeError("DOCX ZIP CRC failed")
        names = archive.namelist()
        if "word/document.xml" not in names:
            raise RuntimeError("DOCX document body missing")
        embedded = [name for name in names if name.startswith("word/media/")]
        if len(embedded) != len(manifest):
            raise RuntimeError("Expected {} embedded figures; found {}".format(len(manifest), len(embedded)))
    older = "not_updated"
    try:
        shutil.copyfile(OUTPUT, OLDER_OUTPUT)
        older = "updated"
    except PermissionError:
        older = "in_use"
    print(json.dumps({"docx": str(OUTPUT), "figures": len(used), "bytes": OUTPUT.stat().st_size,
                      "package_integrity": "PASS", "older_copy": older}, ensure_ascii=False))


if __name__ == "__main__":
    build()
