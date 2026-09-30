from docx import Document

doc = Document(r"E:\PCIE-CPU\.codex_tmp_DDR4最小工程仿真测试报告_V2.0.docx")
for i, p in enumerate(doc.paragraphs[:90]):
    if p.text.strip():
        print(i, repr(p.text[:120]), p.style.name)
print("tables", len(doc.tables), "sections", len(doc.sections), "shapes", len(doc.inline_shapes))
