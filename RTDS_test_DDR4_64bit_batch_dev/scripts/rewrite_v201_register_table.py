from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document


DOCS = Path(r"F:\RTDS_test_DDR4_64bit_batch_dev\docs")
SOURCE = DOCS / "CPU-FPGA PCIe软件寄存器说明V2.01（当前实现）.docx"
BACKUP = DOCS / "CPU-FPGA PCIe软件寄存器说明V2.01（当前实现）_改写前_20260921.docx"


def replace_paragraph_text(paragraph, text):
    """保留段落及首个文本运行的格式，仅替换文字。"""
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in list(paragraph.runs[1:]):
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def set_cell(cell, text):
    """不改变单元格和表格样式，只更新单元格正文。"""
    paragraph = cell.paragraphs[0]
    replace_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


if not BACKUP.exists():
    shutil.copy2(SOURCE, BACKUP)

doc = Document(SOURCE)

# 当前寄存器表：按偏移定位，避免依赖固定行号。
register_rows = {
    row.cells[0].text.strip(): row for row in doc.tables[0].rows[1:]
}
updates = {
    "0x10": ("RC（读后清除）", "外部输入", "返回当前中断状态；本次读地址握手同时产生单周期清除脉冲"),
    "0x14": ("W1P（写1脉冲）/RO", "0x00000000", "写bit0=1产生一个时钟周期的上行读取完成确认；写0无效；读回通常为0"),
    "0x18": ("RO", "0x00000000", "最老待取上行块的DDR字节地址；队列为空时为0"),
    "0x1C": ("RO", "0x00000000", "最老待取上行块的有效字节数；队列为空时为0，非零值为4Byte整数倍"),
    "0x20": ("写1启动/RO忙状态", "0x00000000", "空闲时写bit0=1启动一次下行任务；执行期间读回1；任务结束、参数不合法或链路中止后由硬件清0"),
    "0x2C": ("RO", "当前值/保持至清除", "DDR当前工作条件以及已经发生的接收溢出、响应错误和超时事件，位定义见第4节"),
    "0x30": ("RO", "0x00000000", "已完成且尚未由CPU确认的上行数据块数，范围0～8"),
    "0x34": ("RO", "0x00000000", "当前最早一块待读数据的批次编号；每完成一个非空块递增，32bit自然回绕；队列为空时为0"),
}
for offset, (attr, reset_value, meaning) in updates.items():
    row = register_rows[offset]
    set_cell(row.cells[2], attr)
    set_cell(row.cells[3], reset_value)
    set_cell(row.cells[4], meaning)

# DDR_STATUS表：用直观名称区分当前条件和需要软件清除的历史事件。
status_rows = {row.cells[0].text.strip(): row for row in doc.tables[1].rows[1:]}
status_updates = {
    "0": ("当前状态", "1表示DDR后端允许业务访问"),
    "1": ("当前状态", "1表示MIG初始化和校准完成"),
    "2": ("事件标志（保持至清除）", "DDR AXI读写响应曾出现错误"),
    "3": ("事件标志（保持至清除）", "上行接收或缓存曾发生溢出/丢字事件"),
    "4": ("事件标志（保持至清除）", "DDR AXI或FDMA事务曾发生超时"),
}
for bit, (kind, meaning) in status_updates.items():
    row = status_rows[bit]
    set_cell(row.cells[2], kind)
    set_cell(row.cells[3], meaning)

# RTL依据表只改写不直观的名称。
for row in doc.tables[2].rows[1:]:
    if row.cells[0].text.strip() == "下行窗口、16Byte校验及门铃完成条件":
        set_cell(row.cells[0], "下行窗口、16Byte校验及任务结束条件")

paragraph_replacements = {
    "属性说明：RO为只读，RW为读写，RC为读并触发清除，W1P为写1产生单周期脉冲，W1S为写1置位并由硬件清除。":
        "属性说明：RO为只读，RW为读写，RC为读取后触发清除，W1P为写1产生一个时钟周期的内部脉冲。0x20同时包含写入启动功能和只读忙状态，其行为已在表中直接说明。",
    "读取 0x10 INT_STATE 会产生清除脉冲，同时清除中断通知及上述粘滞状态。DDR_READY和INIT_CALIB_COMPLETE是实时状态，不受该清除脉冲影响。因此处理中断时应先读取并保存 DDR_STATUS，再读取 INT_STATE。":
        "读取 0x10 INT_STATE 会产生清除脉冲，同时清除中断通知及上述事件标志。DDR_READY和INIT_CALIB_COMPLETE反映读取时的当前条件，不受该清除脉冲影响。因此处理中断时应先读取并保存 DDR_STATUS，再读取 INT_STATE。",
    "队头在软件确认前保持稳定；读 INT_STATE只清通知，不释放块。":
        "当前最早待取块的地址、长度和批次编号在软件确认前保持稳定；读 INT_STATE只清除中断通知，不改变待取数据块。",
    "等待H2C DMA完成，并使用驱动提供的顺序保证，确保DDR数据先于门铃可见。":
        "等待H2C DMA完成，并使用驱动提供的顺序保证，确保DDR数据写入完成后才操作下行提交寄存器。",
    "CPU_WR_DONE清0只表示该请求已被硬件消费并结束，可能是正常完成、参数拒绝或链路中止；当前寄存器不能区分这几种结果。":
        "CPU_WR_DONE清0只表示硬件已经结束本次请求，可能是正常完成、参数不合法或链路中止；当前寄存器不能区分这几种结果。",
    "依次读取 FPGA_CPU_START_ADDR、FPGA_CPU_LEN和 UP_BLOCK_SEQ，得到同一队头快照。":
        "依次读取 FPGA_CPU_START_ADDR、FPGA_CPU_LEN和 UP_BLOCK_SEQ，取得同一个最早待取块的地址、长度和批次编号。",
    "写 CPU_REC_STATE.bit0=1确认并释放一个队头块。":
        "写 CPU_REC_STATE.bit0=1，表示最早待取块已经读取完成，允许FPGA后续复用该块。",
    "上行队列采用单消费者模型。读取队头、执行C2H和写确认必须由同一串行流程完成。":
        "同一时刻只允许一条软件流程处理上行数据。读取最早待取块的信息、执行C2H和写完成确认必须按顺序完成。",
    "下行必须先完成H2C数据写入，再写 CPU_WR_DONE门铃。":
        "下行必须先完成H2C数据写入，再向 CPU_WR_DONE.bit0写1启动任务。",
    "DMA完成与MMIO门铃/确认之间必须使用操作系统或驱动提供的内存顺序保证。":
        "DMA完成与MMIO控制寄存器写入之间必须使用操作系统或驱动提供的内存顺序保证。",
}

for paragraph in doc.paragraphs:
    old = paragraph.text
    if old in paragraph_replacements:
        replace_paragraph_text(paragraph, paragraph_replacements[old])

doc.save(SOURCE)
print(f"REGISTER_DOC_REWRITE_PASS output={SOURCE}")
