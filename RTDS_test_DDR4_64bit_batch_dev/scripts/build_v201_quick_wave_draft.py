from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"E:\PCIE-CPU\.codex_tmp_DDR4最小工程仿真测试报告_V2.0.docx")
OUTPUT = ROOT / "docs" / "RTDS_V2.01快速仿真波形确认稿.docx"
EVIDENCE = ROOT / "evidence" / "v201_quick_wave_draft"
FIGURES = EVIDENCE / "figures_no_title_png"
MANIFEST = json.loads((EVIDENCE / "wave_manifest.json").read_text(encoding="utf-8"))

FILES = {
    "1g": (ROOT / "rtl" / "rtds_1g_periodic_tx.v", ROOT / "regression" / "batch_1g" / "tb_rtds_1g.sv"),
    "tx": (ROOT / "external_snapshot" / "fdma_aurora_brg" / "fdma_aurora_brg.srcs" / "sources_1 - loop_validate" / "new" / "fdma_aurora_brg_tx.v", ROOT / "regression" / "batch_1g" / "tb_rtds_tx.sv"),
    "rx_fifo": (ROOT / "rtl" / "rtds_rx_frame_fifo.v", ROOT / "regression" / "batch_1g" / "tb_rtds_rx_fifo.sv"),
    "ring": (ROOT / "rtl" / "rtds_rx_block_writer.v", ROOT / "regression" / "batch_1g" / "tb_rtds_ring.sv"),
    "axi": (ROOT / "RTDS_test_DDR4_64bit.srcs" / "sources_1" / "new" / "CONV_DMA.v", ROOT / "regression" / "batch_1g" / "tb_rtds_axi.sv"),
    "axi_tail": (ROOT / "RTDS_test_DDR4_64bit.srcs" / "sources_1" / "new" / "CONV_DMA.v", ROOT / "regression" / "batch_1g" / "tb_rtds_axi_tail.sv"),
    "regs": (ROOT / "RTDS_test_DDR4_64bit.srcs" / "sources_1" / "new" / "INFO_EXCH.v", ROOT / "regression" / "batch_1g" / "tb_rtds_regs.sv"),
}

DETAILS = {
1: ("配置输入先在250MHz域过滤编号0、11，再把首个合法编号对应的duty与Vin锁存为48bit配置，通过请求/应答握手整体跨入31.25MHz域。锁存完成后configured与configured_id保持，后续合法编号不覆盖。", "新增需求_主机逻辑与验证说明.md中的“首次合法Byte1译码”和V2.01_整体设计与验证架构.md中的“48bit完整握手跨时钟”。", "观察src_clk与user_clk异步运行；cfg_byte_valid提交非法值时configured不变，合法值到达后decode_valid、cfg_req/cfg_ack依次动作，最终configured和configured_id稳定。"),
2: ("1G域以63拍为发送机会。状态机用FIRST、LAST两拍发送6Byte，小端duty在前、Vin在后；第二拍REM=1并结束帧。", "新增需求_主机逻辑与验证说明.md规定6Byte载荷；整体架构文档规定31.25MHz下63拍，即2.016us。", "比较相邻tx_sof的时间刻度；首拍tx_sof=1、末拍tx_eof=1，tx_data与tx_rem组合后形成完整6Byte。"),
3: ("valid一旦拉高便保持载荷到ready握手；忙时周期事件只增加skipped_periods。channel_up掉线清理半帧，恢复后从新的周期机会发送。", "V2.01整体架构的“忙时跳过、不补发”和“链路恢复”规则。", "tx_ready为0时valid/data/SOF/EOF/REM保持；skipped_periods递增。channel_up下降时发送停止，恢复后重新出现合法两拍帧。"),
4: ("下行桥按16Byte一帧向FDMA请求两个64bit字，地址从0开始；跨域FIFO输出两拍LocalLink，首拍SOF、次拍EOF。", "PCIe交互协议_V2.01.md规定下行地址16Byte对齐、长度为16Byte整数倍。", "raddr=0后出现两个rvalid/rready握手；用户侧两次valid/ready依次携带SOF和EOF，数据序号连续。"),
5: ("300帧批量读取时，FIFO可编程满阈值为504，给完整读突发保留空间；用户侧长反压后按序排空，溢出与下溢标志保持0。", "整体架构文档的批量下行与反压策略；本轮综合修复采用512-max(BURST_WORDS,8)。", "limit=300，peak_words与fifo_wr_cnt进入近满区；ready恢复后valid握手继续，fifo_overflow/fifo_underflow不置位。"),
6: ("10G接收端把每个有效64bit拍写入异步FIFO，存储域按顺序取出；SOF/EOF仅随线路出现，不作为固定帧长判定。", "新增需求说明和整体架构均规定上行采用有效数据流模型。", "user_clk与mem_clk同时显示；rx_data的输入顺序在frame_data输出中保持，SOF/EOF变化不阻塞传递。"),
7: ("FIFO满时不覆盖旧数据，无法写入的一个64bit拍按一个低32bit有效字计入饱和丢失计数；释放空间后继续接收。", "PCIe交互协议_V2.01.md把0x38定义为32bit有效数据字丢失数。", "fifo_full后继续rx_valid会使dropped_frames增加；frame_ready打开后frame_valid被消费，length_error保持0。"),
8: ("块写入器只取frame_data低32bit。两个字拼成一个64bit写拍；窗口封存时若剩一个字，使用WSTRB=0F写低半。", "新增需求说明规定高32bit永久无效；整体架构规定紧凑拼接和奇数尾拍。", "输入高32bit变化不影响fdma_wdata中的有效序列；half_valid表示暂存单字，尾写fdma_wstrb=0F。"),
9: ("周期边界只提出封存请求，必须等在途FDMA写响应成功后才发布块；队头地址、长度和序号在cpu_ack前稳定。", "PCIe交互协议_V2.01.md定义0x18/0x1C保持至CPU确认，读通知不释放数据。", "seal_pending先出现，fdma_wdone后ready_blocks和irq才变化；cpu_ack前head_addr/head_len/head_seq保持。"),
10:("八个256KiB块循环使用，地址为0x100000+k×0x40000。测试连续发布并确认24块，覆盖三轮地址复用。", "新增需求说明的100ms×8块地址规划。", "head_seq单调递增，head_addr按八个基址循环；第三轮复用时前一轮已经由cpu_ack释放。"),
11:("ready_blocks达到8时保留最老队头，后续输入被reject_word拒绝并按字累计；CPU确认后队列继续前进。", "协议规定队列满不得覆盖未读块。", "满队期间head_addr/head_len/head_seq不变；额外两个字使dropped_frames增加2，随后cpu_ack逐块释放。"),
12:("写响应未返回前不发布；fdma_werror使本次有效字计入丢失且不形成可读空洞，取消错误后后续块可发布。", "整体架构的“成功响应后计长和发布”规则。", "inject_error期间fdma_wdone与fdma_werror同时出现，write_failure和丢失计数响应但ready_blocks不提前增加；恢复后新块发布。"),
13:("CONV_DMA把AW、W、B作为独立握手通道；地址扩展覆盖低4MiB，八个上行块基址应写入不同存储位置。", "V2.01整体架构规定DDR映射扩展并消除旧64KiB别名。", "AWVALID或WVALID在对应READY为0时保持；AWADDR依次覆盖0x100000至0x2C0000，BRESP完成后fdma_wdone出现。"),
14:("520拍传输从0x180FF8开始，CONV_DMA按4KiB边界和AXI突发长度拆分，写后按相同顺序读回。", "整体架构和测试清单要求跨4KiB传输不得形成非法突发。", "观察0x180FF8附近AWADDR/AWLEN以及后续地址事务；每个事务在边界内结束，最终写响应和读回完成。"),
15:("fdma_wstrb直接贯穿CONV_DMA到M_AXI_WSTRB，RAM模型逐字节更新。先FF整拍写入，再0F只更新低32bit。", "整体架构规定奇数尾拍WSTRB=0F。", "波形显示FDMA与AXI两侧WSTRB均为0F；回读M_AXI_RDATA保留原高32bit并呈现新的低32bit。"),
16:("0x20写1锁存下行地址和长度并产生单周期start；忙期间重复提交不重新启动，done后门铃自动清零。", "PCIe交互协议_V2.01.md对0x20自动清零门铃的定义。", "第一次合法写使start和starts计数变化；后续忙时写不增加计数，done后可再次产生新start。"),
17:("0x14写1且低字节WSTRB有效时产生单周期ack并自动读回0；状态寄存器直接呈现队列数、序号和丢字数。", "PCIe交互协议_V2.01.md对0x14、0x30、0x34、0x38的定义。", "合法写产生ack，WSTRB=E不产生ack；读取时rdata与RX_READY_BLOCKS、RX_HEAD_SEQ、RX_DROPPED_FRAMES一致。"),
}

CODE_PATTERNS = {
1: (("decode_valid", "cfg_req", "configured_id"), ("submit(0)", "submit(8'(id))", "ONE_G/INVALID")),
2: (("PERIOD_CYCLES", "period_tick", "tx_rem"), ("ONE_G/CADENCE", "observed_payload", "ONE_G/LAST")),
3: (("skipped_periods", "channel_up", "tx_valid"), ("SKIP_NOT_EXERCISED", "tx_ready <= 0", "channel_up <= 0")),
4: (("FRAME_BYTES", "fdma_raddr", "tx_sof"), ("run_case(1", "TX/DATA_OR_BOUNDARY", "TX/ADDRESS")),
5: (("TX_PROG_FULL_THRESH", "fifo_almost_full", "fifo_overflow"), ("run_case(300", "peak_words<500", "FIFO_OVERFLOW")),
6: (("xpm_fifo_async", "fifo_write", "fifo_read"), ("user_clk", "mem_clk", "TEST PASS")),
7: (("drop_bin", "fifo_full", "dropped_frames"), ("fifo_full", "dropped_frames", "length_error")),
8: (("half_valid", "pair_start", "fdma_wstrb"), ("words=n%2+2", "check_head", "send_word")),
9: (("seal_pending", "write_success", "push_block"), ("STREAM/EARLY_PUBLISH", "cpu_ack", "wait_count")),
10: (("write_block", "read_block", "BASE_ADDR"), ("for(int n=0;n<24", "seq%8", "ack()")),
11: (("reject_word", "queued", "drop_sum"), ("wait_count(n-23)", "send_word(999)", "STREAM/DROP")),
12: (("write_failure", "fdma_werror", "dropped_frames"), ("inject_error=1", "ERROR_RESPONSE", "inject_error=0")),
13: (("aw_pending", "w_pending", "M_AXI_AWVALID"), ("WRITE_BEGIN", "'h100000", "allow_aw")),
14: (("write_len_minus_one", "M_AXI_AWLEN", "M_AXI_ARLEN"), ("180ff8", "520", "read_words")),
15: (("fdma_wstrb", "M_AXI_WSTRB", "M_AXI_WDATA"), ("2ffff8", "8'h0f", "upper32")),
16: (("CPU_WR_DONE", "r_RD_START", "CPU_WR_DONE_pos"), ("starts", "command_addr", "done")),
17: (("CPU_REC_STATE", "RX_READY_BLOCKS", "S_AXI_WSTRB"), ("acks", "wstrb", "RX_DROPPED_FRAMES")),
}


def set_font(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


def line_of(path: Path, needles):
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    for needle in needles:
        for i, line in enumerate(lines, 1):
            if needle in line:
                return i
    return 1


def line_refs(path: Path, needles):
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    found=[]
    for needle in needles:
        for i,line in enumerate(lines,1):
            if needle.lower() in line.lower():
                found.append(f"{path.relative_to(ROOT)}:{i}({needle})")
                break
    return "；".join(found) if found else f"{path.relative_to(ROOT)}:1"


doc = Document(TEMPLATE)
h1_ppr = deepcopy(doc.paragraphs[34]._p.pPr)
h2_ppr = deepcopy(doc.paragraphs[35]._p.pPr)
body_ppr = deepcopy(doc.paragraphs[36]._p.pPr)
caption_ppr = deepcopy(doc.paragraphs[77]._p.pPr)

doc.paragraphs[5].text = "RTDS V2.01 快速仿真波形确认稿"
for run in doc.paragraphs[5].runs:
    set_font(run, 22, True)
for table in doc.tables[:3]:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                text = p.text.replace("2026-9-18", "2026-9-20").replace("26.9.18", "26.9.20")
                if "AXI BIST仿真测试报告V2.0" in text:
                    text = text.replace("AXI BIST仿真测试报告V2.0", "RTDS V2.01快速仿真波形确认稿")
                p.text = text
                for r in p.runs:
                    set_font(r, 9)

# 清除旧正文，保留封面、审批、修订、目录和节属性。
body = doc._element.body
start = doc.paragraphs[34]._p
removing = False
for child in list(body):
    if child is start:
        removing = True
    if removing and child.tag != qn("w:sectPr"):
        body.remove(child)

settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")


def apply_ppr(p, source):
    old = p._p.pPr
    if old is not None:
        p._p.remove(old)
    p._p.insert(0, deepcopy(source))


def add_para(text="", kind="body"):
    p = doc.add_paragraph()
    source = h1_ppr if kind == "h1" else h2_ppr if kind == "h2" else caption_ppr if kind == "caption" else body_ppr
    apply_ppr(p, source)
    if kind == "caption":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 14 if kind == "h1" else 12 if kind == "h2" else 10.5, kind in {"h1", "h2"})
    return p


def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, 8.5, row is table.rows[0])
    add_para("")


add_para("1 文档目的与证据边界", "h1")
add_para("本确认稿用当前工程七组模块级行为仿真的真实VCD、WDB、WCFG和日志建立“设计实现→设计方案依据→RTL/测试台代码依据→波形直接证据”链路。正文波形均由VCD事件重绘，明确属于“VCD波形重绘图”，并非Vivado原生截图。")
add_para("七组日志均包含TEST PASS，退出码为0。XSim日志中的单词error仅来自Tcl脚本文本回显（如save_error变量），不代表运行期错误。结论不覆盖Aurora物理建链、MIG校准、真实PCIe Root Port闭环、综合实现时序或上板测试。")

add_para("2 快速仿真与波形清单", "h1")
rows=[]
for i,item in enumerate(MANIFEST,1):
    rows.append((i,item["title"],item["case"],f'{item["start_ps"]/1e6:.3f}～{item["end_ps"]/1e6:.3f} us'))
add_table(("编号","波形/用例标题","仿真组","VCD实际窗口"),rows,(1.1,8.8,2.0,3.2))

add_para("3 设计、代码与波形证据", "h1")
for i,item in enumerate(MANIFEST,1):
    title=item["title"]
    case=item["case"]
    design,plan,wave=DETAILS[i]
    rtl,tb=FILES[case]
    rtl_ref=line_refs(rtl,CODE_PATTERNS[i][0])
    tb_ref=line_refs(tb,CODE_PATTERNS[i][1])
    add_para(f"3.{i} {title}", "h2")
    add_para("设计实现："+design)
    add_para("设计方案证据："+plan)
    add_para(f"代码证据：{rtl_ref}；{tb_ref}。测试台使用独立激励、握手监视和$fatal断言，并以TEST PASS作为完整退出条件。")
    add_para("波形证据与理解方法："+wave+f'。本图窗口来自实际VCD跳变，范围为{item["start_ps"]/1e6:.3f}～{item["end_ps"]/1e6:.3f} us。')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent=Cm(-1.3); p.paragraph_format.right_indent=Cm(-1.3)
    p.add_run().add_picture(str(FIGURES/item["figure"]),width=Cm(17.2))
    add_para(f"图{i} {title}（VCD波形重绘图）", "caption")
    add_para("结论与边界：该图与对应断言共同证明本节列出的模块级接口行为；不据此宣称物理链路、真实DDR或整机闭环已经验收。")

add_para("4 日志、原始数据与复现入口", "h1")
add_para("七组原始证据位于 evidence\\v201_quick_wave_draft\\raw，每组包含VCD、WDB、WCFG和xsim日志。wave_manifest.json记录17张图对应的实际时间窗与信号列表；figures_no_title_png用于正文，figures_with_title_png用于独立审阅。")
add_para("复现脚本为 scripts\\render_v201_quick_waves.py 和 scripts\\build_v201_quick_wave_draft.py。WDB可通过 regression\\batch_1g\\open_module_full_wave.cmd 在Vivado GUI中查看全部已记录信号。")

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUTPUT)
print(f"DOCX_BUILD_PASS {OUTPUT}")
