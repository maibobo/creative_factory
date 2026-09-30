from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "CPU-FPGA_PCIe软件寄存器说明_V2.01_当前实现.md"
TARGET = ROOT / "docs" / "CPU-FPGA PCIe软件寄存器说明V2.01（当前实现）.docx"


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.replace("  ", " ").strip()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["No Spacing"]
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(9)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for r_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for c_index in range(width):
            value = clean_inline(row_values[c_index]) if c_index < len(row_values) else ""
            cells[c_index].text = value
            for paragraph in cells[c_index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    if r_index == 0:
                        run.bold = True
        if r_index == 0:
            set_repeat_table_header(table.rows[0])
            for cell in cells:
                set_cell_shading(cell, "D9EAF7")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True

    title = clean_inline(lines[0][2:])
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("CPU-FPGA PCIe 软件接口实现说明\n").bold = True
    sub.add_run("2026-09-21\n依据当前生产工程RTL整理")
    doc.add_page_break()

    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []
    first_heading_skipped = False
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                if table_rows:
                    add_table(doc, table_rows)
                    table_rows = []
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                continue
            table_rows.append(cells)
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []
        if not line.strip():
            continue
        if line.startswith("# ") and not first_heading_skipped:
            first_heading_skipped = True
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            doc.add_heading(clean_inline(heading.group(2)), level=level)
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.add_run(clean_inline(numbered.group(2)))
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(line[2:]))
            continue
        p = doc.add_paragraph()
        p.add_run(clean_inline(line))

    if table_rows:
        add_table(doc, table_rows)
    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title = "CPU-FPGA PCIe软件寄存器说明V2.01（当前实现）"
    doc.core_properties.subject = "当前XDMA AXI-Lite寄存器与DDR上下行操作说明"
    doc.core_properties.author = "Codex，依据当前生产工程RTL整理"
    doc.save(TARGET)

    with zipfile.ZipFile(TARGET, "r") as archive:
        bad = archive.testzip()
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        missing = required.difference(archive.namelist())
    if bad or missing:
        raise RuntimeError(f"DOCX structure invalid: bad={bad!r}, missing={sorted(missing)}")
    check = Document(TARGET)
    if len(check.paragraphs) < 30 or len(check.tables) < 3:
        raise RuntimeError("DOCX content unexpectedly incomplete")
    print(f"DOCX_OK paragraphs={len(check.paragraphs)} tables={len(check.tables)} path={TARGET}")


if __name__ == "__main__":
    build()


