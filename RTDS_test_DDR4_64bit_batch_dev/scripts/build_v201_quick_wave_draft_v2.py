from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"E:\PCIE-CPU\.codex_tmp_DDR4最小工程仿真测试报告_V2.0.docx")
OUTPUT = ROOT / "docs" / "RTDS_V2.01快速仿真波形确认稿.docx"
EVIDENCE = ROOT / "evidence" / "v201_quick_wave_draft_v2"
FIGURES = EVIDENCE / "figures_no_title_png"
MANIFEST = json.loads((EVIDENCE / "wave_manifest.json").read_text(encoding="utf-8"))


def set_font(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")


doc = Document(TEMPLATE)
h1_ppr = deepcopy(doc.paragraphs[34]._p.pPr)
h2_ppr = deepcopy(doc.paragraphs[35]._p.pPr)
body_ppr = deepcopy(doc.paragraphs[36]._p.pPr)
caption_ppr = deepcopy(doc.paragraphs[77]._p.pPr)
doc.paragraphs[5].text = "RTDS V2.01 快速仿真波形确认稿"
for run in doc.paragraphs[5].runs:
    set_font(run, 22, True)

for table in doc.tables[:3]:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.text = p.text.replace("2026-9-20", "2026-9-21").replace("26.9.20", "26.9.21")
                p.text = p.text.replace("AXI BIST仿真测试报告V2.0", "RTDS V2.01快速仿真确认稿深化版")
                for r in p.runs:
                    set_font(r, 9)

body = doc._element.body
start = doc.paragraphs[34]._p
removing = False
for child in list(body):
    if child is start:
        removing = True
    if removing and child.tag != qn("w:sectPr"):
        body.remove(child)

settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")


def apply_ppr(p, source):
    old = p._p.pPr
    if old is not None:
        p._p.remove(old)
    p._p.insert(0, deepcopy(source))


def add_para(text="", kind="body"):
    p = doc.add_paragraph()
    source = h1_ppr if kind == "h1" else h2_ppr if kind == "h2" else caption_ppr if kind == "caption" else body_ppr
    apply_ppr(p, source)
    if kind == "caption":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 14 if kind == "h1" else 12 if kind == "h2" else 10.5, kind in {"h1", "h2"})
    return p


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = str(value)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for ridx, row in enumerate(table.rows):
        for i, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, 8.5, ridx == 0)
    add_para("")
    return table


def add_figure(filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(-1.3)
    p.paragraph_format.right_indent = Cm(-1.3)
    p.add_run().add_picture(str(FIGURES / filename), width=Cm(17.2))
    add_para(caption, "caption")


def figs(prefix):
    return [m for m in MANIFEST if m["file"].startswith(prefix)]


add_para("1 文档目的", "h1")
add_para("本文档说明RTDS V2.01七组模块级快速仿真的设计原理、测试原因、测试方法和波形证据。每个功能点先说明数据通路和控制条件，再解释状态机、关键信号与软件可见结果，最后按信号发生顺序阅读波形。")
add_para("正文波形直接来自本次仿真数据库。长时间没有状态变化的区间被拆开，关键握手按局部时间轴放大。多时钟用例按时钟域分区，每个分区的第一行是该域时钟。总线稳定期间以带状轨迹显示，并在整个保持区间标出完整数值。")
add_para("七组日志均包含TEST PASS，分别覆盖1G配置与周期发送、16Byte下行、10G接收跨时钟、上行块写入、AXI访问、部分字节写和CPU寄存器交互。")

add_para("2 验证架构与信号阅读方法", "h1")
add_table(("仿真组","实际设计模块","测试重点"),[
    ("1g","RTDS_1G_PERIODIC_TX","首次配置、48bit跨时钟、2.016µs发送、反压和链路恢复"),
    ("tx","fdma_aurora_brg_tx","16Byte帧、异步FIFO、300帧积压和排空"),
    ("rx_fifo","RTDS_RX_FRAME_FIFO","10G输入跨时钟、满状态与有效字损失计数"),
    ("ring","RTDS_RX_BLOCK_WRITER","低32bit拼接、DDR写响应、八块循环和异常处理"),
    ("axi/axi_tail","CONV_DMA与AXI RAM","高地址、4KiB拆分、AW/W独立握手和WSTRB"),
    ("regs","INFO_EXCH","下行提交、CPU确认、字节使能和状态寄存器"),
],(2.2,5.0,8.0))
add_para("二值信号的高、低轨迹分别表示1和0。总线的上下平行线表示数值在该区间保持不变，交叉线只表示总线从一个值切换到另一个值。交叉线本身不是错误或未知值。地址、数据和字节使能使用十六进制；长度、状态、序号和计数器使用十进制。")

add_para("3 详细设计与仿真证据", "h1")

add_para("3.1 首次合法配置译码与跨时钟锁存", "h2")
add_para("验证目的：CPU下行帧的Byte1编号只允许1～10。主机上电后只接受第一次合法编号，随后周期发送的duty和Vin必须始终来自这一份配置。配置由250MHz控制域产生，却要在31.25MHz的1G用户域使用，因此不能把48bit参数逐位同步。")
add_para("设计过程：cfg_byte_valid表示配置字节已经在下行读取接口完成握手。组合译码产生decoded_duty、decoded_vin和decode_valid。首次合法事件把{Vin,duty}写入cfg_hold并置cfg_req。跨时钟邮箱把保持不变的48bit数据和请求一起送入1G域；req_sync有效后，目的域一次性写入cfg_user并置cfg_ack。确认返回到源域后configured置1。cfg_req在本次上电周期保持，因此后续编号不能改写cfg_hold和configured_id。")
add_table(("信号","来源","用途"),[
    ("cfg_hold","250MHz域首次合法配置","跨时钟期间保持完整48bit参数"),
    ("cfg_req/req_sync","配置源域/跨时钟单元","表示有一份稳定配置等待接收"),
    ("cfg_user/cfg_ack","31.25MHz域","保存1G发送使用的参数并返回确认"),
    ("configured/configured_id","250MHz域","供调试和系统确认首次配置已完成"),
],(3.6,4.3,7.3))
add_para("如何阅读波形：先看250MHz区，非法编号不会形成decode_valid和cfg_req；合法编号到达后cfg_hold、cfg_req和configured_id稳定。再看31.25MHz区，req_sync到达后cfg_user一次性取得完整参数，cfg_ack返回后configured置位。两个时钟边沿没有固定相位关系，这正是该测试必须覆盖的原因。")
add_figure("01_config_cdc.png","图1 首次合法配置从250MHz域完整传入31.25MHz域")

add_para("3.2 1G载荷格式与2.016µs发送机会", "h2")
add_para("验证目的：6Byte参数通过32bit接口发送，需要两次成功握手。这里验证的不是“两拍本身”，而是6Byte字节顺序、首尾标记、第二拍有效字节数，以及连续发送机会是否按63个31.25MHz周期出现。")
add_table(("状态值","状态名称","作用与转移"),[
    ("0","IDLE","没有在途数据；period_tick到达后进入FIRST"),
    ("1","FIRST","输出前4Byte并置tx_sof；tx_ready为1后进入LAST"),
    ("2","LAST","输出剩余2Byte并置tx_eof、tx_rem=1；握手后回IDLE"),
],(2.2,3.0,10.0))
add_para("图2a中，FIRST状态发送duty低字节、duty高字节和Vin的低两字节；LAST状态发送Vin剩余两字节。tx_valid拉高后，如果tx_ready为0，状态和载荷都必须保持。图2b比较连续两个tx_sof，period_count从0计到62后产生period_tick，因此发送机会间隔是63×32ns=2.016µs。")
add_figure("02a_1g_payload.png","图2a 6Byte参数在32bit接口上的两次握手")
add_figure("02b_1g_cadence.png","图2b 连续发送机会相隔63个1G用户时钟周期")

add_para("3.3 1G反压、周期跳过与链路恢复", "h2")
add_para("为什么需要跳过：当FIRST或LAST已经有效但tx_ready为0时，当前数据不能撤销。若此时新的2.016µs发送机会继续排队，长时间反压会形成无法控制的任务积压。本设计只保留正在发送的一帧，新到达的机会增加skipped_periods，不建立补发队列。该计数用于调试链路吞吐，不参与载荷生成。")
add_para("链路下降时，channel_up使状态机立即回IDLE，丢弃未完成的半帧，但cfg_user和cfg_ack不清除。链路恢复后arm_count等待4个用户时钟，让复位和同步状态稳定；随后从新的period_tick开始发送。掉线期间错过的任务不补发。")
add_para("图3a先观察tx_ready下降，tx_curr_st停留在当前发送状态，tx_valid、tx_data、SOF/EOF和REM保持；期间period_tick再次到达时skipped_periods增加。图3b左侧显示channel_up下降和状态回IDLE，右侧显示channel_up恢复、arm_count计满以及随后连续多个完整帧重新出现。")
add_figure("03a_1g_backpressure.png","图3a 反压期间保持当前数据并记录错过的周期机会")
add_figure("03b_1g_link_recovery.png","图3b 链路重新可用后恢复周期发送")

add_para("3.4 16Byte下行帧边界、数据顺序与地址步进", "h2")
add_para("验证目的：V2.01把每个下行配置帧改为16Byte。FDMA和10G业务接口均为64bit，因此一帧应包含两个被接受的数据字。测试重点是读取长度、地址步进、数据顺序和帧边界，而不是把“两拍”作为独立性能指标。")
add_table(("状态机","状态","含义"),[
    ("FDMA读取","RD_IDLE/RD_REQ/RD_DATA/RD_WAIT","等待任务、提出读请求、接收两个数据字、等待本次读结束"),
    ("10G发送","TX_IDLE/TX_SOF/TX_DATA","等待FIFO数据、发送首字、发送后续字并在第二字置EOF"),
],(3.0,5.0,7.2))
add_para("波形上半区中，raddr=0后rareq提出一次16Byte读取，rvalid与rready完成两次握手。下半区中，同样两个数据依次出现在tx_data；第一个握手带SOF，第二个握手带EOF。received连续增加且数据序号不跳变，说明跨域FIFO没有漏拍、重复拍或调换顺序。下一帧的raddr应增加16Byte。")
add_figure("04_downlink_frame.png","图4 一帧16Byte数据从FDMA读取到10G用户接口的完整过程")

add_para("3.5 300帧下行与发送FIFO安全余量", "h2")
add_para("为什么设置504：发送FIFO物理深度为512个64bit字。一次FDMA读取开始后，上游可能继续返回完整突发，因此不能等到FIFO完全满才停止申请。当前16Byte帧每次读取2字，设计至少保留8字安全余量，阈值为512-max(2,8)=504。")
add_para("测试台先让10G接收端ready保持0，同时连续提交300帧，使FIFO占用升高。图5a中fifo_wr_cnt和peak_words达到504，fifo_almost_full阻止新的FDMA读事务，fifo_overflow保持0。图5b中ready恢复后，valid与ready连续握手，received递增，FIFO占用下降；fifo_underflow保持0，证明停止与恢复没有破坏顺序。")
add_figure("05a_tx_fifo_fill.png","图5a 接收端停顿时FIFO占用达到504")
add_figure("05b_tx_fifo_drain.png","图5b 接收端恢复后FIFO继续按序发送")

add_para("3.6 10G接收数据跨时钟传递", "h2")
add_para("验证目的：10G用户接口工作在约156.25MHz，DDR写入与块管理工作在250MHz。异步FIFO必须保证每个64bit有效拍按原顺序到达存储域。线路SOF和EOF仍可变化，但上行业务层不依赖固定帧长。")
add_para("波形按两个时钟域分区。接收域中，rx_valid有效时完整rx_data进入跨时钟通路，rx_ready保持1。存储域中，FIFO非空使frame_valid有效，frame_ready允许时读指针推进，frame_data按相同顺序输出，received同步增加。SOF/EOF的组合变化没有改变数据进入FIFO的条件。")
add_figure("06_rx_fifo_cdc.png","图6 64bit接收数据从156.25MHz域按序传入250MHz域")

add_para("3.7 FIFO满状态、数据损失统计与恢复", "h2")
add_para("这个测试验证异常拥塞保护，不代表正常运行应频繁出现。真实设计默认FIFO深度为2048；快速仿真缩为16，以便在短时间内构造满状态。可能触发满状态的实际原因包括DDR或AXI长时间没有响应、八个块长期未被CPU取走，或输入平均速率持续超过存储通路处理能力。")
add_para("10G接收接口不能根据DDR状态暂停远端线路，因此rx_ready固定为1。FIFO可写时，fifo_write把数据写入；full_user或写复位忙时，当前有效拍不能保存，drop_bin增加1。这里每个64bit线路拍只有低32bit是业务数据，所以计数单位是一个32bit有效字。drop_bin通过计数跨时钟单元传到250MHz域，形成软件最终读取的dropped_frames，因此目的域数值会比源域晚若干同步拍。")
add_para("图7a应从fifo_full已经有效后的下一次rx_valid开始阅读：rx_ready仍保持1，但drop_bin从0变1；dropped_frames经过跨时钟同步后也变为1。图7b中frame_ready拉高后，frame_valid与frame_ready连续握手，received增加，fifo_full解除；随后追加的数据再次在存储域被读取。软件读取0x38发现计数变化后，应把相应数据段标记为不完整，并由应用决定报警或重新同步，硬件无法恢复已经丢失的数据。")
add_figure("07a_rx_fifo_full_drop.png","图7a FIFO满时拒绝一次写入并增加源域损失计数")
add_figure("07b_rx_fifo_release.png","图7b 损失计数传到存储域并在排空后恢复接收")

add_para("3.8 低32bit数据在DDR写入前的紧凑拼接", "h2")
add_para("数据位置：frame_data来自接收异步FIFO，只有[31:0]有效。RTDS_RX_BLOCK_WRITER位于该FIFO与CONV_DMA之间，先在250MHz域整理数据，再通过fdma_waddr、fdma_wdata和fdma_wstrb发起DDR写入。")
add_para("第一个32bit字写入half_data并置half_valid。第二个字到达时pair_start有效，active_transfer锁存{后到字,先到字}、WSTRB=FF和有效字节数8。若采集段结束时只有一个字，tail_start把half_data放在低32bit，WSTRB设为0F，有效字节数为4。直到fdma_wdone成功返回，used_bytes才分别增加8或4。")
add_para("图8a从两个frame_valid/frame_ready握手开始，随后观察half_valid、pair_start和完整fdma_wdata。图8b显示half_data在采集段结束后形成0F尾写，并继续显示地址、写请求、数据握手、最终响应和used_bytes变化，因此可以确认该操作确实发生在DDR写入之前。")
add_figure("08a_pair_pack.png","图8a 两个低32bit有效字组成一个64bit DDR写数据")
add_figure("08b_tail_pack.png","图8b 单个剩余字仅写DDR低四字节")

add_para("3.9 100ms采集段结束、DDR响应与CPU待取队列", "h2")
add_para("seal_pending的含义：time_count到达PERIOD_CYCLES-1时产生tick，表示当前100ms采集段已经到结束点。模块从这一拍开始不再向当前块追加新字，并置seal_pending，等待half_data和已开始的DDR事务处理完。它不是DDR成功信号，也不是CPU通知信号。")
add_table(("状态值","状态","接口行为"),[
    ("0","IDLE","可以接受新字或准备结束当前采集段"),
    ("1","REQ","fdma_wareq有效，等待CONV_DMA接受请求并置busy"),
    ("2","DATA","fdma_wvalid有效，等待fdma_waccept接收数据"),
    ("3","RESPONSE","等待最终fdma_wdone；fdma_werror决定本次是否成功"),
],(2.0,3.0,10.2))
add_para("当状态机回到IDLE、seal_pending仍有效且half_valid已经清除时，seal_done表示当前采集段的所有DDR事务都已结束。只有used_bytes非0才产生push_block，把地址、有效长度和序号写入待取队列。queued从0变1时irq通知CPU。CPU完成C2H后产生cpu_ack，pop_block使read_block推进并释放当前队头。")
add_para("图9a按tick、seal_pending、tail_start、状态机和fdma_wdone的顺序阅读。图9b按write_success、seal_done、push_block、queued、irq、head_*、cpu_ack和pop_block阅读。必须等待B响应，是因为W数据被接口接收并不等于DDR写事务已经成功。")
add_figure("09a_segment_end.png","图9a 当前100ms采集段结束并等待最后一次DDR写响应")
add_figure("09b_queue_and_ack.png","图9b 数据块进入CPU待取队列并在C2H完成后释放")

add_para("3.10 八个256KiB物理块与三轮地址复用", "h2")
add_para("256KiB表示每个物理块可使用的最大地址空间，不表示每100ms一定产生256KiB数据。CPU实际搬运长度由head_len给出。八个物理块的基址为0x100000、0x140000、0x180000、0x1C0000、0x200000、0x240000、0x280000和0x2C0000。")
add_para("测试连续形成24个小数据块，并在每个块检查完成后立即由CPU模型确认。24=8×3，因此能够验证三轮地址复用。write_block和read_block都只在0～7之间循环；head_seq与publish_seq不随地址回绕，用于软件检查顺序。第三轮再次使用某个地址时，该地址上一轮的数据已经确认释放。")
add_para("三张图分别显示序号0～7、8～15和16～23。阅读时同时观察head_addr和head_seq：地址每八块重复一次，序号始终递增。queued在该测试段只在0和1之间变化，说明地址复用前没有遗留未取数据。")
add_figure("10a_ring_round1.png","图10a 八个物理块的第一轮地址使用")
add_figure("10b_ring_round2.png","图10b 八个物理块的第二轮地址使用")
add_figure("10c_ring_round3.png","图10c 八个物理块的第三轮地址使用")

add_para("3.11 八块均待取时的保护和软件可见结果", "h2")
add_para("reject_word是块写入器内部的单拍判断。frame_valid与frame_ready已经形成输入事件，但enable为0、queued等于8或当前块已达到256KiB时，该字不能写入DDR，reject_word有效。这样可以保证尚未由CPU确认的数据不会被覆盖。")
add_para("dropped_frames是兼容旧接口保留的信号名称，V2.01中实际统计32bit有效字。软件看不到reject_word脉冲，只能通过0x38读取累计值。图11a中queued达到8后，head_addr和head_seq保持，额外两个输入分别产生reject_word，计数增加2。图11b中cpu_ack产生pop_block，queued下降、read_block和队头推进，frame_ready继续允许后续数据处理。")
add_para("软件影响：0x38增加表示上行序列存在不可恢复的数据缺口。CPU仍可读取队列中已有的完整数据块，但应把相关时间段标记为不完整。确认释放队头不会自动清除0x38。")
add_figure("11a_queue_full.png","图11a 八块均待取时保持最老队头并记录新增数据损失")
add_figure("11b_queue_resume.png","图11b CPU确认后队列推进并恢复接收能力")

add_para("3.12 DDR写响应延迟、错误响应和后续恢复", "h2")
add_para("测试目的：上游数据握手完成后，AXI从端仍可能延迟返回B响应，也可能返回SLVERR或DECERR。设计不能在响应到达前把数据计入有效长度，更不能向CPU提供尚未确认成功的数据。")
add_para("response_delay和inject_error都是测试台信号。response_delay=300模拟B响应长时间延迟；inject_error=1模拟错误响应。真实系统中的对应原因可能是AXI互连译码失败、DDR控制器访问错误或下游响应异常。")
add_para("图12a中，fdma_waccept以后状态机停留在RESPONSE，fdma_wdone未到达，ready_blocks不增加。图12b中，fdma_wdone与fdma_werror同时出现，write_failure有效；used_bytes不增加，dropped_frames按active_transfer中的有效字数增加。下一次成功事务继续使用尚未确认的位置，因此DDR有效区域中不形成空洞；成功后才更新长度和队头。")
add_figure("12a_delayed_response.png","图12a 写响应延迟期间保持等待且不增加待取块数")
add_figure("12b_error_recovery.png","图12b 写错误不计入有效长度，后续成功事务继续使用该位置")

add_para("3.13 上行DDR区域的高地址访问", "h2")
add_para("这里的高地址是相对旧仿真模型的64KiB索引而言。V2.01上行区域从0x100000延伸到0x2FFFFF，如果RAM模型或互连仍只使用低16位地址，不同物理块会落到同一个低地址，造成数据互相覆盖。")
add_para("图13a把八个块基址分别作为M_AXI_AWADDR发起写事务。AWVALID/AWREADY与WVALID/WREADY是两个独立通道，允许任意先后完成；二者完成后才进入B响应阶段。图13b用相同八个地址发起读事务，M_AXI_RDATA与各地址预先写入的内容逐一对应，证明地址高位参与了存储索引。")
add_figure("13a_high_address_write.png","图13a 八个上行块基址形成不同AXI写地址")
add_figure("13b_high_address_readback.png","图13b 八个上行块地址分别读回各自数据")

add_para("3.14 520拍请求的4KiB边界拆分", "h2")
add_para("设计规格：AXI4的INCR突发不能跨越4KiB地址边界，单个突发最多256拍。CONV_DMA接收的FDMA请求可以大于一个AXI突发，因此必须按“当前4KiB页剩余拍数、256拍上限、请求剩余拍数”三者的最小值拆分。对齐的长请求应正常完成，不应因为需要跨页而报错。")
add_para("示例地址0x180FF8的低12位是0xFF8，距离0x181000边界只有8Byte。数据宽度为8Byte，所以当前页只剩1拍。520拍写请求应分为：0x180FF8处1拍，AWLEN=0；0x181000处256拍，AWLEN=255；0x181800处256拍，AWLEN=255；0x182000处7拍，AWLEN=6。每个突发的最后一个有效数据字产生WLAST，B响应后地址再推进。")
add_para("图14a的四个局部分别显示上述AWADDR和AWLEN，任何一个突发都没有越过自身的4KiB页。图14b显示读方向使用相同拆分规则，ARADDR和ARLEN对应四段，RDATA连续返回且每段末尾出现RLAST。若FDMA地址不是8Byte对齐，CONV_DMA会结束该请求并报告错误；这与合法长请求的自动拆分不同。")
add_figure("14a_4k_write_split.png","图14a 520拍写请求拆分为1、256、256和7拍")
add_figure("14b_4k_read_split.png","图14b 读请求按相同地址边界拆分并连续返回")

add_para("3.15 WSTRB=FF和0F的部分字节写", "h2")
add_para("WSTRB的每一位控制M_AXI_WDATA中的一个字节。FF表示8Byte全部写入，0F表示仅低4Byte写入。上行数据可能只剩一个32bit有效字，因此必须使用0F，防止无效的高32bit覆盖DDR中原有内容。")
add_para("测试先用FF把一个已知64bit值写入同一地址，建立可检查的高32bit基准；再用0F写入新的低32bit；最后读回。只有采用这一顺序，才能证明高32bit确实保持，而不是碰巧读到0。")
add_para("图15a显示fdma_wstrb和M_AXI_WSTRB均为FF。图15b显示两侧均为0F，M_AXI_WDATA虽然包含64bit，但RAM模型只更新低四个字节。图15c的M_AXI_RDATA由原高32bit和新低32bit组成，证明WSTRB从FDMA到AXI RAM没有丢失。")
add_figure("15a_wstrb_ff.png","图15a 使用FF写入完整64bit基准值")
add_figure("15b_wstrb_0f.png","图15b 使用0F只更新低32bit")
add_figure("15c_wstrb_readback.png","图15c 读回值保留原高32bit并包含新低32bit")

add_para("3.16 V2.01下行提交、忙状态保持与硬件清零", "h2")
add_para("接口含义：软件先完成H2C写DDR，再把下行地址和长度写入0x24、0x28，最后向0x20 bit0写1。CPU_WR_DONE是0x20的内部状态，不表示PCIe数据已经自动发送。CPU_WR_DONE_pos检测0→1，生成单周期r_RD_START，适配器据此开始读取DDR。")
add_para("CPU_WR_DONE为1期间表示本次请求仍在处理。重复向0x20写1不会产生新的上升沿，地址和长度寄存器也保持本次任务的值。适配器完成最后一帧，或检查后拒绝非法请求时，产生CPU_WR_DONE_EVENT，INFO_EXCH据此把CPU_WR_DONE清0。清0表示FPGA侧本次请求已经结束，不表示远端子机已经执行配置。")
add_para("图16a显示AW/W写0x20后CPU_WR_DONE置1，CPU_WR_DONE_d晚一拍跟随，CPU_WR_DONE_pos和start只保持一个时钟，starts只增加1。图16b左侧显示done使CPU_WR_DONE清0并在读回路径得到0；右侧显示随后一次新提交再次产生start。")
add_figure("16a_submit_start.png","图16a 软件写0x20产生一次下行启动脉冲")
add_figure("16b_submit_clear_restart.png","图16b 硬件完成事件清零状态并允许下一次提交")

add_para("3.17 CPU完成确认、WSTRB过滤与状态寄存器", "h2")
add_para("0x14用于CPU完成C2H后的确认。CPU_REC_STATE不是保持到软件再次清零的寄存器，而是单周期命令：只有写地址为0x14、WDATA bit0为1且WSTRB bit0为1时才产生ack。WSTRB=E的二进制是1110，低字节没有写使能，因此不能改变位于低字节的bit0，也不能释放队头。")
add_table(("偏移","硬件输入","软件含义"),[
    ("0x30","RX_READY_BLOCKS","当前可以读取的完整数据块数量，范围0～8"),
    ("0x34","RX_HEAD_SEQ","当前最老待取块的连续序号，用于检查漏取或重复取"),
    ("0x38","RX_DROPPED_FRAMES","累计损失的32bit有效字数，名称为兼容旧接口保留"),
],(2.0,4.2,9.0))
add_para("图17a中，两次WSTRB=F的0x14写操作分别形成ack，acks从0增加到2；WSTRB=E的写操作没有形成第三个ack。图17b中，AXI-Lite主机依次读取0x30、0x34和0x38。araddr选择寄存器，rvalid表示rdata有效；rdata分别等于对应硬件输入值1、7和9，证明地址译码与CPU可见状态一致。")
add_figure("17a_ack_wstrb.png","图17a 低字节有效时形成确认脉冲，WSTRB=E时被过滤")
add_figure("17b_status_reads.png","图17b CPU依次读取待取块数、队头序号和有效字损失计数")

add_para("4 原始数据与复核入口", "h1")
add_para("七组原始日志、WDB、WCFG和波形数据库位于evidence\\v201_quick_wave_draft\\raw。深化版图片和清单位于evidence\\v201_quick_wave_draft_v2。wave_manifest.json记录每张子图使用的时钟域和信号。")
add_para("使用regression\\batch_1g\\open_module_full_wave.cmd可以在Vivado GUI中打开任一WDB并查看全部记录信号。本文档只说明已执行的模块级行为仿真，不把物理链路、DDR校准、完整PCIe Root Port或上板状态写成已验证结果。")

# 清理模板及正文中不采用的旧术语与异常标点。
replacements = {
    "VCD波形重绘图":"仿真波形",
    "结论与边界":"验证说明",
    "封存":"结束当前采集段",
    "发布":"加入待取队列",
    "消费":"处理",
    "。。":"。",
    "&#x20;":"",
}
for p in doc.paragraphs:
    for old,new in replacements.items():
        if old in p.text:
            full=p.text.replace(old,new)
            if p.runs:
                p.runs[0].text=full
                for r in p.runs[1:]: r.text=""
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for old,new in replacements.items():
                    if old in p.text:
                        full=p.text.replace(old,new)
                        if p.runs:
                            p.runs[0].text=full
                            for r in p.runs[1:]: r.text=""

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUTPUT)
print(f"DOCX_V2_BUILD_PASS figures={len(doc.inline_shapes)} output={OUTPUT}")
