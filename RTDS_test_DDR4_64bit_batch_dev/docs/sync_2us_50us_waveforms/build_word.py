from pathlib import Path
import sys
import zipfile

sys.path.insert(0, str(Path.home() / 'AppData/Local/Temp/sync_docx_deps_20260930'))
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r'C:\project\RTDS\DDR\RTDS_test_DDR4_64bit_batch_dev')
wave_dir = root / 'docs' / 'sync_2us_50us_waveforms'
out = root / 'docs' / '主机同源同步实现与波形记录_20260930.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.3)
sec.right_margin = Cm(2.3)

styles = doc.styles
for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    st = styles[style_name]
    st.font.name = 'Microsoft YaHei'
    st.font.size = Pt(10.5 if style_name == 'Normal' else (18 if style_name == 'Title' else 13))
    rpr = st.element.get_or_add_rPr()
    ea = rpr.rFonts
    if ea is None:
        ea = OxmlElement('w:rFonts')
        rpr.insert(0, ea)
    ea.set(qn('w:eastAsia'), 'Microsoft YaHei')

title = doc.add_heading('主机 FPGA 同源同步实现与波形记录', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('工程：RTDS_test_DDR4_64bit_batch_dev    日期：2026-09-30')
doc.add_paragraph('范围：主机 FPGA 的 2 µs Aurora 1G 同步帧、50 µs RK3588 GPIO，以及 CPU/子机接口交接。')

doc.add_heading('1. 已实施的主机逻辑', 1)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
for c, s in zip(table.rows[0].cells, ['项目','实现','位置']): c.text = s
rows = [
    ('同源节拍', '板级100 MHz经Clock Wizard产生250 MHz；每500拍翻转一次2 µs事件，每12500拍输出一个50 µs GPIO周期。高低各6250拍；上升沿与本周期首个2 µs事件重合。', 'rtl/rtds_sync_timebase.v'),
    ('1G发送', '事件翻转跨入31.25 MHz用户域；首次合法配置锁存后从下一事件开始发送。反压或发送中错过的事件计数但不排队；在途帧按握手完成。', 'rtl/rtds_1g_periodic_tx.v'),
    ('CPU GPIO', 'cpu_sync_50us输出接RPT205 TEST3/AM11，LVCMOS33；RK3588接收端已确认支持3.3 V。复位时为低，1G未建链时仍持续运行。', 'TOP.v；DDR4_PIN.xdc'),
    ('10G', '保留现行两拍64bit下行转发；主机不截取10G数据的低2Byte。', '现有PCIe与10G通路'),
]
for row in rows:
    cells = table.add_row().cells
    for cell, s in zip(cells,row): cell.text = s

doc.add_heading('2. 1G 帧字节与时序', 1)
doc.add_paragraph('每帧四拍、16个有效字节；每拍 tx_rem=3，首拍SOF、末拍EOF。CPU首个64bit读拍的Byte1（fdma_rdata[15:8]）是1G配置编号1～10，Byte0不参与1G配置。')
table = doc.add_table(rows=1, cols=3)
table.style='Table Grid'
for c,s in zip(table.rows[0].cells,['拍','字节','内容']): c.text=s
for row in [
    ('0','Byte0–3','duty低、高字节；Vin低两字节'),
    ('1','Byte4–7','Vin高两字节；Byte6=0xF0；Byte7=0x00'),
    ('2','Byte8–11','0x00，预留'),
    ('3','Byte12–15','0x00，预留'),
]:
    cells=table.add_row().cells
    for cell,s in zip(cells,row): cell.text=s
doc.add_paragraph('2 µs是250 MHz源事件的精确间隔。1G帧首拍经过跨域和31.25 MHz采样后，正常相邻首拍间隔量化为1984或2016 ns；反压、断链时可能缺帧。本版不含序号或补报。')

doc.add_heading('3. 已执行检查', 1)
for s in [
    '定向节拍仿真通过：连续两个50 µs周期，每周期25个2 µs事件；GPIO高低各6250个250 MHz时钟拍。',
    '定向1G仿真通过：10组配置共60帧，校验四拍字节序、0xF0标识、SOF/EOF/REM、反压保持、跳过计数和链路重建。',
    '主工程Vivado 2020.2 RTL elaborate通过，cpu_sync_50us端口与AM11/LVCMOS33约束生效；0 error、0 critical warning。16条普通warning位于既有PCIe、DDR/IP文件。',
    '依照当前要求，未运行主工程综合、实现、时序分析或联板测量。',
]: doc.add_paragraph(s,style='List Bullet')

figures = [
    ('timebase_overview.png','图1  同源节拍：50 µs GPIO与每2 µs源事件'),
    ('timebase_zoom.png','图2  上升沿附近的2 µs事件与250 MHz时钟'),
    ('one_g_frame.png','图3  1G四拍同步帧与Byte6=0xF0'),
    ('one_g_backpressure.png','图4  反压时保持帧数据并跳过事件'),
]
present = [(wave_dir/name,caption) for name,caption in figures if (wave_dir/name).exists()]
doc.add_heading('4. 后台 VCD 波形记录', 1)
if present:
    doc.add_paragraph('下列图由后台VCD信号变化数据绘制；图像及VCD文件保存在同目录。图中表示仿真接口时序，不代表板上测量。')
    for path,caption in present:
        doc.add_paragraph(caption)
        doc.add_picture(str(path),width=Cm(16.0))
else:
    doc.add_paragraph('未取得兼容的真实VCD，波形截图待手动仿真或授权的后台VCD采集后补入。现有XSim WDB原始数据仍保留在regression/v201_runs。')

doc.add_heading('5. 软件与子机协同', 1)
for s in [
    'RK3588将对应3.3 V GPIO配置为输入，以上升沿对齐50 µs计算节拍；下降沿不触发计算。记录边沿时间和软件超期次数。',
    'CPU先用既有16Byte H2C流程提交Byte1=1～10的首个有效配置；该配置锁定到主机复位。GPIO运行与1G建链、结果就绪没有同义关系。',
    'CPU取结果继续按UP_READY_BLOCK_CNT、块地址/长度、C2H、CPU_REC_STATE流程处理约100 ms上行块；本版没有每50 µs的PCIe读写或同步寄存器。',
    '子机收到完整16Byte且Byte6=0xF0的帧后锁存参数；跨到Buck的25 MHz域，在该域生成一个时钟周期即40 ns的触发脉冲。这部分由子机/Buck负责人实现。',
    '硬件需确认TEST3/AM11到RK3588输入脚的实际连线；联板应分别实测CPU GPIO周期和占空比、1G帧到达间隔、子机40 ns脉冲及多子机到达偏差。',
]: doc.add_paragraph(s,style='List Bullet')

doc.add_heading('6. 交付文件', 1)
for s in [
    'rtl/rtds_sync_timebase.v；rtl/rtds_1g_periodic_tx.v；TOP.v；当前工程实际使用的DDR4_PIN.xdc与XPR源文件列表。',
    'regression/batch_1g/tb_rtds_sync.sv、tb_rtds_1g.sv、run_modules.tcl；docs/PCIe交互协议_V2.01.md。',
    '代码改前备份：docs/sync_2us_50us_prechange_20260930。',
]: doc.add_paragraph(s,style='List Bullet')

doc.save(str(out))
with zipfile.ZipFile(out) as z:
    assert z.testzip() is None
    names=set(z.namelist())
    assert {'[Content_Types].xml','word/document.xml','word/_rels/document.xml.rels'} <= names
    assert len([n for n in names if n.startswith('word/media/')]) == len(present)
print(f'WORD_OK path={out} figures={len(present)} bytes={out.stat().st_size}')
