from pathlib import Path
import json,re,hashlib,collections,zipfile,xml.etree.ElementTree as ET,sys,math
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
from docx import Document
from docx.shared import Pt
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path('F:/RTDS_test_DDR4_64bit_batch_dev')
OLD=Path('E:/PCIE-CPU/outputs/RTDS_spec_test_refined_20261010')
OUT=ROOT/'docs/RTDS_functional_20261011';OUT.mkdir(exist_ok=True)
C=json.loads((OLD/'content.json').read_text('utf-8'))
INV=json.loads((OLD/'port_inventory.json').read_text('utf-8'))['modules']
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
# Reparse local copies; do not substitute a remote source when a local file is missing.
parser=(OLD/'scripts/inventory.py').read_text('utf-8').split('def collect():')[0]
parser=parser.replace("str(path.relative_to(OUT/'_sources')).replace('\\\\','/')","str(path)")
ns={'__file__':str(OLD/'scripts/inventory.py')};exec(parser,ns)
source_audit=[];port_changes=[];inactive_modules=[]
for m in INV:
    rel=m['path'];p=ROOT/rel.removeprefix('current/')
    if rel.startswith('vendor/'):
        matches=list((ROOT/'RTDS_test_DDR4_64bit.srcs').rglob(Path(rel).name))
        assert matches,rel
        p=matches[0]
    assert p.exists(),p
    new=next(x for x in ns['modules'](p) if x['module']==m['module'])
    a={(x['name'],x['direction'],x['width']) for x in m['ports']};b={(x['name'],x['direction'],x['width']) for x in new['ports']}
    if a!=b:port_changes.append({'module':m['module'],'removed':sorted(a-b),'added':sorted(b-a)})
    source_audit.append({'module':m['module'],'path':str(p),'sha256':sha(p),'same_as_1010':sha(p)==sha(OLD/'_sources'/rel)})
    m['ports']=new['ports']
    m['local_path']=str(p)
    if m['module']=='RTDS_SYNC_TIMEBASE':inactive_modules.append(m)
INV=[m for m in INV if m not in inactive_modules]

specs=[list(r) for r in C['specs'] if int(r[0][2:])<47]
cats={'下行发送':'下行 / 命令与发送','上行接收与拼接':'上行 / 接收与拼接','块管理':'上行 / 块生命周期','1G 配置与发送':'1G / 配置与发送','地址与存储':'存储 / 事务与地址','寄存器控制':'控制 / 寄存器','异常恢复':'异常 / 检测与恢复','同源节拍':'节拍 / 周期与相位','时钟与复位':'系统 / 时钟与复位','模块输入输出':'系统 / 指示'}
for r in specs:
    r[1]=cats.get(r[1],r[1])
    r[4]=r[4].replace('需专项验证，现有冒烟仅覆盖部分断链。','')
    r[4]=r[4].replace('不能用PERIOD_CYCLES=200的缩时结果替代；','')
    if int(r[0][2:])>=32:
        r[2]+='（32bit）'
        r[3]=r[3].replace('，32bit小端','')+' '+r[4]
        r[4]=''
    if r[0]=='S-15':
        r[3]='未就绪不发起下行访问；在途任务被中断后结束命令，恢复就绪后应接收新任务并排除旧数据和旧完成事件。'
        r[4]='设计目标；TC-F11已发现恢复缺陷，当前实现尚未满足。'
    if r[0]=='S-24':r[4]='解除后从计数0开始。'
    if r[0]=='S-06':r[4]+='【本机差异】当前启用RTL仍是两拍6 Byte，末拍REM=1；本条为目标规格，尚未落到本机活动工程。'
    if r[0] in ('S-22','S-23','S-24'):r[4]+='【本机差异】timebase文件存在但未被当前XPR启用，TOP未连接同源事件/GPIO；本条保留为目标规格。'
    if r[0]=='S-21':r[4]+='【本机差异】当前周期来自31.25MHz用户域63拍，即2016ns，非250MHz同源2µs事件。'

# Refine existing broad requirements into independently observable functionality.
extra=[
('下行 / 命令与发送','单任务参数快照','任务接受后锁定本批地址与长度，忙期间后续寄存器写入不改变在途任务。','', 'S-14'),
('下行 / 命令与发送','按帧递增读地址','将合法批量按16 Byte帧读取，下一帧地址在上一帧基础上加16。','每帧2个64bit读拍。','S-04'),
('下行 / 命令与发送','反压期间保持边界','未被接受的有效拍保持数据、SOF、EOF和REM，恢复后从同一拍继续。','','S-04'),
('下行 / 命令与发送','完成事件与业务执行分离','批量结束事件清命令忙；命令拒绝和断链也可结束命令，忙清零不代表对端执行成功。','','S-14'),
('上行 / 接收与拼接','帧边界不筛除有效字','每个有效接收拍贡献一个低32bit字；SOF、EOF和REM不决定当前有效字是否进入拼接。','','S-07'),
('上行 / 接收与拼接','奇数尾字收尾','周期收尾时单独写入尚未配对的低32bit字，不把未使用高半计入长度。','WSTRB=0x0F；有效长度增加4 Byte。','S-10'),
('上行 / 块生命周期','空周期不发布','一个周期没有成功写入有效字时，不生成空的可读块或消耗可读序号。','','S-25'),
('上行 / 块生命周期','在途写响应阻止发布','周期事件只发起收尾；在途数据接受后仍要等待成功写响应，之后才能发布非空块。','','S-10'),
('上行 / 块生命周期','满队列所有权保护','8块均未确认时不覆盖任何未读块；不能存储的新字计入丢失。','','S-09'),
('上行 / 块生命周期','读清与释放独立','清通知只消除通知状态；直到CPU完成数据读取并确认，队头仍归CPU所有。','','S-11'),
('上行 / 块生命周期','队头元数据一致','队头地址、长度和序号共同指向同一可读块，确认后整体转到下一块。','','S-11'),
('存储 / 事务与地址','AXI写通道独立接受','AW与W可在不同周期接受；B响应结束该写事务，不能用W接受替代写成功。','','S-13'),
('存储 / 事务与地址','4KiB与最大突发拆分','跨页请求拆成多个不跨4KiB的INCR突发，突发间地址连续。','每突发最多256拍；每拍8 Byte。','S-13'),
('1G / 配置与发送','非法配置保持未启动','首次合法配置到来前，编号0及大于10的编号不能建立配置或启动周期帧。','','S-05'),
('1G / 配置与发送','首次配置不被覆盖','首个合法配置确认后，即使每个下行帧继续产生Byte1事件，发送参数仍保持首次配置。','','S-05'),
('1G / 配置与发送','跨域发送完整参数','用户域在配置传递完成后使用成组的48bit参数，duty和Vin不能来自不同配置。','duty16bit；Vin32bit。','S-05'),
('1G / 配置与发送','忙时跳过不补发','已配置且链路恢复等待结束后，源事件遇发送忙或ready为0则跳过，不在恢复后追补积压周期。','skipped_periods为32bit自然回绕计数。','S-21'),
('1G / 配置与发送','断链中止与配置保持','用户复位或断链中止当前帧，保留已锁定配置；恢复等待后由新源事件启动首拍。','主机复位清除配置。','S-21'),
('节拍 / 周期与相位','GPIO独立持续输出','1G反压、帧忙或链路断开不暂停同源GPIO和2µs源事件。','','S-23'),
('异常 / 检测与恢复','响应错误不记成功字节','非OKAY写响应标记错误，对应字不加入块有效长度，并累计丢失字数。','','S-10'),
('异常 / 检测与恢复','无进展超时诊断','在途DDR/FDMA事务长期无进展时置超时状态，软件可通过0x2C观察并经0x10读清。','生产DDR TIMEOUT_CYCLES=65535；专项模型使用256。','S-26'),
('异常 / 检测与恢复','溢出诊断与累计计数分离','读清诊断不清0x38累计丢字；累计值按32bit字统计并饱和，主机复位清零。','','S-12'),
]
split_from={}
for i,(cat,name,req,lim,parent) in enumerate(extra,875):
    sid=f'S-{i}';specs.append([sid,cat,name,req,lim]);split_from[sid]=parent
    if parent in ('S-22','S-23','S-24'):
        specs[-1][4]='目标规格；本机当前TOP尚未接入同源节拍。'

# Aggregate common buses, retaining every exact identifier, direction and evaluated width.
def busgroup(n):
    n=n.lower()
    if re.search(r'axi|^mem_(aw|ar|w|r|b)',n):return 'AXI总线'
    if 'fdma' in n and 'clk' not in n and 'rst' not in n and 'timeout' not in n:return 'FDMA总线'
    if n.startswith('c0_ddr4'):return 'DDR4物理接口'
    return '控制与数据接口'
meaning={'AXI总线':'存储或寄存器读写；AW/W/B为写事务，AR/R为读事务，valid与ready共同成立时接受该拍。',
'FDMA总线':'读写请求、字地址、拍数、数据握手及事务结束；接受数据与响应成功分别表示。',
'DDR4物理接口':'DDR4地址命令、时钟、数据与双向选通。',
'控制与数据接口':'本模块的时钟复位、业务数据、控制及状态。'}
module_meanings={
'CONV_DMA':'FDMA请求转换为AXI突发；完成及错误由写响应产生。',
'INFO_EXCH':'AXI-Lite寄存器读写、下行启动、读清通知与一次性读完确认。',
'RTDS_1G_PERIODIC_TX':'接收Byte1配置，输出周期帧、配置状态和跳过计数。当前活动RTL为两拍6 Byte；四拍目标见S-06。',
'RTDS_RX_BLOCK_WRITER':'拼接低32bit接收字、写DDR并维护可读块队列。',
'RTDS_RX_FRAME_FIFO':'跨接收时钟与存储时钟缓存数据，输出有效数据和丢字状态。',
'RTDS_SYNC_TIMEBASE':'250MHz同源2µs事件及50µs GPIO。',
'RTDS_FPGA_CFG_CDC':'成组传递配置参数及配置确认。','RTDS_FPGA_COUNT_CDC':'跨域传递累计计数快照。',
'RTDS_FPGA_RESET_SYNC':'目的时钟域复位同步。','RTDS_FPGA_1G_INIT':'生成1G初始化时钟及GT/协议核复位顺序。',
'LED':'周期LED指示。','TEMP_RD':'SYSMON温度原始码采样。','CLK_GEN':'板载时钟输入、生成时钟及锁定状态。',
'pcie_fdma_aurora_brg_adapter':'下行命令执行、上行存储管理、配置字节提取和异常状态。',
'fdma_aurora_brg_tx':'DDR读数据缓冲与10G两拍帧发送。','rtds_ddr4_backend':'AXI存储访问、DDR就绪和错误诊断。',
'TOP':'板级PCIe、DDR4、Aurora及GPIO连接。','PCIe':'PCIe与DDR/FDMA业务连接及状态汇总。',
'aurora_power_on_reset_seq':'上电GT与Aurora核复位时序。'}
merged={};port_to_group={};groups=[]
for m in INV:
    buckets=collections.defaultdict(list)
    for p in m['ports']:buckets[busgroup(p['name'])].append(p)
    for group,ports in buckets.items():
        sid=min((C['port_spec_ids'][m['module']+'.'+p['name']] for p in ports),key=lambda x:int(x[2:]))
        lines=[m['module']+' / '+group]
        for (direction,width),ps in __import__('itertools').groupby(sorted(ports,key=lambda p:(p['direction'],p['width'])),key=lambda p:(p['direction'],p['width'])):
            lines.append(f'{direction} {width}bit：'+', '.join(p['name'] for p in ps))
        req=meaning[group] if group!='控制与数据接口' else module_meanings.get(m['module'], '协议核封装的业务接口、链路状态与时钟复位。' if 'aurora' in m['module'] else '封装层引出同名业务控制、数据及状态。')
        limit='CONV_DMA：fdma_wready输入表示源数据有效；fdma_wvalid输出表示该拍已接受。' if group=='FDMA总线' and 'CONV_DMA' in m['module'] else ''
        specs.append([sid,'接口 / '+group,'\n'.join(lines),req,limit])
        oldids=[]
        for p in ports:
            name=m['module']+'.'+p['name'];old=C['port_spec_ids'][name]
            merged[old]=sid;port_to_group[name]=sid;oldids.append(old)
        groups.append({'id':sid,'module':m['module'],'group':group,'old_ids':oldids,'ports':ports})

# Verification points describe partitions and observations, rather than verdicts.
decomp={
1:('固定常量分为版本、构建标识和单块容量；区分真实寄存器与模型BAR读取。','观察读返回字段及只读属性；区分块容量和下行窗口长度。'),
2:('按长度为零、非16倍数、地址不对齐、地址低于窗口、末端越界、长度超限分解拒绝路径。','关联命令忙结束、length_error与FDMA读请求；关注拒绝是否发生在任何存储访问之前。'),
3:('以最小帧分解首拍和末拍，再关联主机载荷字节序、DDR读拍与10G发送拍。','观察数据传递和SOF/EOF/REM对应关系，防止边界标记提前或滞后。'),
4:('区分首读拍等待、首读拍接受、第二拍接受及下一帧首拍。','观察cfg_byte_valid只随帧首实际接受产生，cfg_byte来自该拍Byte1，计数与帧数对应。'),
5:('分解中等批量、满窗口批量以及上下行并行运行。','观察逐帧地址递增、累计拍数和批结束事件，识别中间帧丢失或跨方向串扰。'),
6:('分别阻塞读返回与LocalLink接受，包含帧首、帧尾及FIFO积压后恢复。','观察停顿前后同一有效拍是否保持，恢复后的数据顺序和帧边界是否连续。'),
7:('逐个合法ID隔离复位；将duty档位和Vin档位交叉组合全部覆盖。','观察configured_id、完整48bit配置和实际四拍输出的一致性。'),
8:('区分尚无配置时非法输入、首次合法输入、锁定后再次合法或非法输入。','观察配置建立和保持的时刻；后续Byte1事件是否改变已锁定参数。'),
9:('分解四拍载荷、定时启动、拍中反压和断链中止；帧格式与周期接受分别观察。','观察字节排布、边界、保留字段、SOF间隔及skipped_periods变化。'),
10:('以高32bit变化而低32bit可预测的接收拍覆盖偶数字拼接；再分解跨帧边界的连续字流。','观察前后低32bit字在64bit写拍中的次序，以及高半输入是否污染业务载荷。'),
11:('分解读侧停止、FIFO达到满、额外有效字到达、排空及重新输入。','观察旧数据保留、额外字丢失计数和恢复后的顺序；缩深度FWFT容量与生产深度分开解释。'),
12:('跨越第7→8及第15→16块，分解物理地址回绕和逻辑序号递增。','观察基址、序号、长度和内容共同对应，避免不同轮次块地址别名。'),
13:('8块均未确认时继续输入，再逐块确认释放。','观察队头与未读内容是否保持、丢字计数如何增长以及空位恢复后的接收。'),
14:('区分数据接受、延迟成功响应和错误响应；延迟跨周期与普通延迟分别覆盖。','观察有效长度与可读数量相对于响应的变化顺序，失败写是否误计成功。'),
15:('分别覆盖孤立尾字和奇数块最后一字，对尾字地址预置高半旧值。','观察低4字节更新、高4字节保留和有效长度；区分发送数据值与写字节使能。'),
16:('块已可读但CPU尚未确认时读取通知寄存器，再读取载荷及元数据。','观察中断通知清除与块所有权保持，读清动作不推进队头。'),
17:('分解一次确认、多次独立确认和无byte0有效写；队列从多块递减到空。','观察ACK脉冲数、队头推进次数、可读数和寄存器读回零。'),
18:('覆盖累计值接近最大、首次到达最大、继续丢字及主机复位。','观察计数饱和与清零；现有常规小计数检查不能覆盖该边界。'),
19:('AW、W和B独立停顿，末拍另设停顿；读通道现有定向读回与独立AR/R随机延迟分开。','观察AW/W保持、事务完成与BRESP；现有用例没有独立AR/R随机延迟覆盖。'),
20:('分解8个块地址无别名、页末第一拍、最大256拍突发和末段余拍。','观察拆分后的地址、AWLEN和连续读回数据；区分低32bit块地址与非零高地址扩展。'),
21:('空闲启动→忙时重复启动和改参→完成事件→再次启动；再覆盖W1P与字节使能。','观察启动脉冲、锁定地址/长度、忙清除及确认脉冲，防止命令重复执行。'),
22:('短批量与超过FIFO深度的300帧分别施压；从长阻塞切换到周期性接受。','观察FIFO占用峰值和继续读DDR的门控关系，以及溢出/读空状态。'),
23:('区分主机全局复位、1G用户域复位、业务链路中断；配置与在途任务分别观察。','观察配置是否按复位类别清除或保持，帧从何处重新开始；无复位恢复缺陷单列跟踪。'),
24:('区分真实Root Port/XDMA、行为后端与物理板卡；历史小事务和当前工程重跑分开。','观察链路、BAR访问、DMA、MSI及数据闭环；历史日志和门禁脚本结论分别记录。'),
25:('生产参数下区分连续周期事件、空周期、已完成偶数字、跨边界在途响应和奇数尾字。','同时观察tick时间与ready_blocks首次变化；可读事件不得被当成周期事件。'),
27:('分解板上DDR校准、主机DMA、真实GT链路以及GPIO测量。','观察硬件采样数据和外部测量记录；仿真模型不覆盖物理链路质量。'),
28:('从解除复位后的首次事件开始，连续观察两个GPIO周期内的源事件；运行中重新复位。','观察toggle的相邻翻转间隔及复位输出，不把toggle完整高低周期误当2µs。'),
29:('分解GPIO高电平、低电平、完整周期与首次边沿相位。','观察GPIO与源事件的同源关系以及每周期事件计数。'),
30:('真实adapter每帧Byte1提取接入真实1G；逐ID重复非法→合法→后续配置序列。','观察端到端配置脉冲、配置锁定、LUT参数与四拍载荷，而非只观察中间cfg_byte。'),
31:('合法集合分为最小长度、最大长度和窗口最后一帧；与P-02拒绝集合对照。','观察读取末地址、帧数、拍数和完成次数，关注合法末端是否被错误拒绝。'),
32:('分别在任务前和读事务在途时撤销DDR、PCIe、Aurora就绪；恢复后换地址和模式提交新任务。','同时观察0x20与CONV_DMA读busy、旧返回数据和新任务；忙清零不能单独证明恢复。'),
33:('使用真实INFO_EXCH累计5块，长度依次1～5字；每块先读清通知、再读数据、最后确认。','观察5→0的队列变化和60 Byte有效数据归属，确认前后队头元数据保持/更新。'),
34:('分别覆盖未校准、就绪、读/写SLVERR、FIFO真实溢出、响应超时和读清错误。','观察0x2C各位与物理来源对应，校准位和累计丢字不因清错被混清。'),
35:('温度原始码与busy分别变化，覆盖采样更新前后及复位。','观察0x0C字段和流水时刻；数值为原始码，不直接当摄氏温度。'),
36:('指示计数从复位开始经过半周期与完整周期。','观察LED输出翻转与计数，不将LED作为链路或DDR完成状态。')}
points=[]
for r in C['points']:
    n=int(r[0][2:]);scope,obs=decomp[n]
    ss=r[5].split('、');ss+= [s for s,parent in split_from.items() if parent in ss]
    points.append([r[0],cats.get(r[1],r[1]),r[2],scope,obs,'、'.join(ss),r[6]])
append_points=[
('1G / 配置与发送','反压跳过与重新接受','发送忙和ready=0分别跨过源事件；恢复后观察是否追补。','观察skipped_periods与帧启动的关系，未配置和断链期间另作负向覆盖。','S-21','TC-M01'),
('1G / 配置与发送','帧中断后重新从首拍启动','断链发生在四拍帧中，重新就绪并经过恢复等待。','观察配置保持、旧帧中止和新帧SOF；不等同于整个DDR读事务恢复。','S-21','TC-M01'),
('上行 / 块生命周期','空周期不生成块','周期到期前没有任何有效字及成功写事务。','观察可读数、发布序号与写事务数保持，随后非空块仍从初始基址发布。','S-25','TC-M09'),
('上行 / 块生命周期','尾字响应跨周期','仅有一个未配对字，周期到期触发尾写后阻塞响应。','观察0F尾写、等待响应状态与可读事件的先后。','S-10、S-25','TC-M09'),
('异常 / 检测与恢复','读错误与写错误分路','分别在模型地址边界触发R SLVERR和B SLVERR。','观察响应错误状态；写失败与丢字增加、块未发布之间的关系。','S-10、S-26、S-43','TC-F11'),
('异常 / 检测与恢复','生产FIFO溢出与清错','阻塞B响应令真实2048深度FIFO积压，再持续送字。','观察溢出、超时和丢字后排空清错；清诊断与清累计数分开。','S-08、S-12、S-43','TC-F11'),
('存储 / 事务与地址','AR/R独立停顿覆盖缺口','需要将AR接受延迟与R返回/接受延迟独立组合，并覆盖末拍。','观察读请求保持和R数据顺序；当前M03以定向读回为主，本项待补专项。','S-13','待补：尚无独立AR/R停顿用例'),
('节拍 / 周期与相位','反压断链不影响GPIO','在1G持续反压及链路切换时同时连续观察GPIO和源toggle。','观察节拍是否继续；M08只验证无业务扰动下的时间基准，本项待补组合场景。','S-22、S-23','待补：尚无组合场景用例'),
('控制 / 寄存器','空队列确认','队列为空时写0x14，再形成新可读块。','观察空确认是否误消耗未来块、改变计数或序号；当前五块读取未覆盖空确认。','S-11、S-37','待补：尚无用例'),
('1G / 配置与发送','跳过计数自然回绕','将跳过计数推进至32bit边界再产生忙事件。','观察自然回绕，区别上行丢字饱和行为；当前M01只检查计数增加。','S-21','待补：尚无用例')]
for n,(cat,name,scope,obs,ss,cc) in enumerate(append_points,37):points.append([f'P-{n:02d}',cat,name,scope,obs,ss,cc])

cases=[list(r) for r in C['cases']]
bycase={r[0]:r for r in cases}
tbmap={f'TC-M{i:02d}':name for i,name in enumerate(['1g','ring','axi','axi_tail','regs','rx_fifo','tx'],1)}
tbmap['TC-M08']='sync'
case_sources={}
for cid,name in tbmap.items():
    paths=[ROOT/f'regression/batch_1g/tb_rtds_{name}.sv']
    iface=ROOT/f'regression/batch_1g/rtds_{name}_test_if.sv'
    if iface.exists():paths.append(iface)
    case_sources[cid]=paths
newnames={'TC-F09':'byte1_1g','TC-F10':'tx_boundaries','TC-F11':'ready_recovery','TC-F12':'regs_multiblock','TC-M09':'period_100ms'}
for cid,name in newnames.items():
    case_sources[cid]=[ROOT/f'regression/new_cases_20261010/tb_{name}.sv']
    if cid.startswith('TC-F'):case_sources[cid]+=[ROOT/'regression/new_cases_20261010/rtds_real_harness.sv',ROOT/'regression/batch_1g/rtds_regs_test_if.sv']
    else:case_sources[cid]+=[ROOT/'regression/batch_1g/rtds_ring_test_if.sv']
fastfiles=[ROOT/'regression/sim_xdma_trythrough'/x for x in ['tb_xdma_rtds_trythrough.sv','rp_api.svh','rtds_test_pkg.sv','xdma_rp_backend.sv','aurora_ll_peer_bfm.sv','rtds_scoreboard.sv','rtds_assertions.sv']]
for i in range(1,9):case_sources[f'TC-F{i:02d}']=fastfiles

common_real='共享H环境：clk=250MHz、uclk=156.25MHz、gclk=31.25MHz；regbus驱动真实INFO_EXCH；U_ADAPTER、U_DMA、U_DDR和U_1G为真实RTL，DDR_SIM_RAM/CPU预装与读回为模型。RX_PERIOD_CYCLES=2000（8µs），FDMA和DDR超时=256拍。\n控制变量：pcie_up/channel_up/lane_up为链路就绪，ddr_reset_n控制DDR就绪；hold_r/hold_b阻塞模型响应，fault_read/fault_write把模型访问地址换为0x00800000触发SLVERR；tx_ready控制10G接受。\n数据与检查：expected_tx_base/expected_tx_len为本批地址和字节数，expected_seed为确定性数据基数（不是随机种子），expected_byte1为每帧配置编号；tx_pattern(offset)={seed+offset+1,seed+offset}，偶数offset把[15:8]替换为Byte1。read_requests/read_beats/tx_beats/tx_frames/cfg_pulses/done_events/acks分别累计读请求、读接受拍、发送拍、发送帧、配置脉冲、结束事件、确认脉冲；check_tx使能逐拍比较。cpu_read_mode/cpu_araddr/cpu_arvalid/cpu_rready驱动单拍64bit模型C2H；cpu_rdata/rresp/rlast为返回。\n握手：FDMA数据接受后继续等done/error；regbus按AW/W→B、AR→R驱动，CPU写0x24/0x28后写0x20=1。reset()保持100拍并等待200拍稳定；wait_end()最多100000次读0x20，结束后等100拍。'
details={
'TC-M01':('src_clk半周期2ns，user_clk半周期16ns；expected_duty/expected_voltage是每ID期望参数，observed_payload/expected_payload为48bit观察/参考值；first_word/second_word保存头两拍，beat_index定位四拍；completed_frames/previous_frames记录帧累计与阶段起点；check_cadence控制周期检查，last_sof_time保存前帧时间；previous_stall、held_payload、partial_frame追踪停顿和未完帧。id遍历1～10，target是wait_frames目标；无随机变量。',
'每个ID先复位24个用户拍；submit(0)、submit(11)后等待160拍观察未配置；submit(id)后wait_frames到新增2帧，核对LUT和四拍；再submit(10)观察首次锁定不覆盖。ready拉低210拍后检查skipped_periods增加；在帧首已接受后再拉低90拍检查当前拍保持；恢复后检查继续发送。链路断开100拍，再恢复并检查配置保留和新帧。周期检查启用时记录SOF间隔1984ns或2016ns；全局看门狗2ms。'),
'TC-M02':('clk=250MHz，PERIOD_CYCLES=200（0.8µs缩时）；memory为按字节寻址的独立关联数组；response_delay=3/300为模型写响应延迟，inject_error控制失败响应；address/data/strb保存接受事务，n为块序号，j为块中字索引，words=n%2+2；send_word(value)把$urandom放高32bit，低32bit放可预测value，随机高半用于检查是否被忽略；未显式设置随机种子。',
'n=0～23每块送2/3字，低字=n*16+j；等待可读后按基址100000+(seq%8)*40000、长度words*4和逐字内容比较，再ACK。n=24～31积压8块，每块1字；多送999、1000，检查丢2字及队头未覆盖，再逐块确认。response_delay=300时送2000/2001，等220拍观察尚不可读，再等成功响应检查seq32的8字节。inject_error=1时送3000/3001，等240拍检查无可读块、累计丢4字；恢复后送4000检查seq33的4字节及0F尾写。'),
'TC-M03':('M_AXI_ACLK=250MHz；cycle_count驱动确定性门控：allow_aw=(cycle_count%11)>7；allow_w=(cycle_count%5)>1且末拍满足%17>13；allow_b=(cycle_count%31)==0。ram_*是门控前后的RAM握手/响应；inject_error使BRESP=SLVERR。address/count/seed/bad分别为任务字节地址、64bit拍数、数据标识和注错使能；beat为拍序；block_id=0～7；seed是载荷高32bit，不是随机种子。本例无随机激励。',
'复位10拍后先向100000+block_id*40000的8块各写16拍{block_id,beat}，再逐块读回；从180FF8写并读520拍{123,beat}，拆分1+256+256+7拍，断言AW不跨4KiB。向200000写16拍{456,beat}并bad=1；整个过程检查AW和W在阻塞时保持、W尚在接受时done不得提前、B到达后error与bad一致。最后检查读写busy均释放；看门狗2ms。'),
'TC-M04':('共享M03确定性AW/W/B门控；fdma_wstrb首次FF随后0F；address=0x2FFFF8、count=1；seed首次12345678、再次DEADBEEF；write_words同时把seed放入两半。',
'复位10拍；先用FF写1234567812345678；改0F再写DEADBEEFDEADBEEF；按64bit读回并逐拍比较12345678DEADBEEF，检查M_AXI_WSTRB透传0F；过程中执行AW/W保持与4KiB边界断言；看门狗2ms。'),
'TC-M05':('clk=250MHz；starts/acks累计start和ack单拍；wr(address,value,strobe=F)为AXI-Lite写任务，check(address,value)读回比较；regbus.command_addr/command_len为锁参观察值；done为结束事件输入；不使用随机。',
'复位8拍，读版本/构建/容量及诊断常量；写24=8000、28=16、20=1。忙时再次启动并试写9000/32，检查starts=1且参数仍为8000/16；注入done单拍使忙清零，再启动使starts=2。连续两次写14=1，每次读回0；再以不含byte0的strobe写1，检查acks仍为2。'),
'TC-M06':('user_clk半周期3.2ns，mem_clk半周期2ns；FIFO_DEPTH=16；expected[$]保存全部64bit输入参考，received累计读拍，value={$urandom,n}；随机高半用于确认FIFO整拍保存，未显式设置种子；stored决定该拍是否加入参考队列；SOF=(n%3==0)、EOF=(n%5==0)、REM=7。',
'复位16个用户拍再稳定30拍；停止读侧后依次送n=0～16并入参考队列，等待20个存储拍观察满。再送n=100但不入队，等待20拍检查丢1字；打开ready逐拍弹出参考并比对直到received=17；恢复后送99并比到18，检查length_error未置位；看门狗100µs。'),
'TC-M07':('fdma_clk半周期2ns，user_clk半周期3.2ns；FRAME_BYTES=16、MEM_WINDOW_BYTES=32768。received为接受总拍数，peak_words记录dbg_fdma_bus[35:26]最大占用；monitor_enable屏蔽复位窗口；was_stalled/held_payload保存停顿拍。frames/hold_cycles配置批量和初始阻塞；frame_id/beat定位FDMA生产；cycle驱动ready=(cycle%7)<4；payload(word_index)={A5C32001,word_index}。无随机。',
'run_case(1,10)和run_case(300,3000)分别复位20个FDMA拍、解除后稳定20拍，再启链路。produce每帧等待rareq核对raddr=frame_id*16，发两拍直到rready接受；consume初始阻塞指定周期后按4/7拍接受。逐拍比较payload、首SOF末EOF和REM7；停顿期间比较held_payload；结束核对received=2*frames和溢出/读空状态，长批量peak_words至少500；日志峰值为2/504。'),
'TC-M08':('clk_250m半周期2ns；rst_n低有效；i为逐周期索引0～24999；prev_toggle保存上拍源事件值；expected_gpio=(i%12500)<6250，expected_toggle=((i/500)%2)==0；events_in_50us累计GPIO周期中的翻转次数；无随机。',
'保持复位5个下降沿，检查两输出为0；解除后每个上升沿延迟1ns采样，逐拍比较GPIO与toggle参考式。每12500拍重新计事件，周期末检查25次，连续运行25000拍覆盖两个GPIO周期；最后下降沿重新置复位并延迟1ns检查输出清零；看门狗120µs。'),
'TC-F09':('duty/voltage为当前ID参考参数；expected_cfg=(voltage*1048576<<16)|duty；expected_data为当前四拍之一；beat=0～3、frames累计完成帧；base保存每ID起始帧数，id遍历1～10，more_frames(target)最多10000个1G拍。未引入随机。',
'每ID独立H.reset()。通过H.valid_tx依次发送8000/16/seed100000/id0和8010/16/seed200000/id11，等200个gclk确认未配置且帧数未增加；8020/16/seed300000/id为首次合法配置，等base+2帧，比较configured_id和cfg_user；8030/16/seed400000发送不同合法ID，等base+4帧确认锁定。每拍比较expected_data、REM3、SOF只在beat0、EOF只在beat3；第二拍检查F0/00，后两拍全零。'),
'TC-F10':('reject(addr,len)参数为拒绝请求；每次清读/请求/发送/结束计数。合法请求由H.valid_tx(addr,len,seed,id)同时核对数据及计数；全部为确定性边界，无随机。',
'reject依次输入(8000,0)、(8000,15)、(8000,17)、(8001,16)、(7FF0,16)、(FFF0,32)、(8000,32784)。每项先读10清诊断并等100拍，启动后等结束，检查length_error、无读请求/读拍/发送拍、done_events=1及20读0。合法三项为(8000,16,110000,1)、(8000,32768,220000,2)、(FFF0,16,330000,3)；检查每帧读请求地址递增16、读发送len/8拍、len/16帧与配置脉冲、恰1结束事件。'),
'TC-F11':('kind=0/1/2选择Aurora、PCIe、DDR中断，length分别16/64/32；value保存真实寄存器读值；failed_checks累计失败而不中断后续检查；fresh_check使能恢复数据比较，fresh_beats/fresh_errors记录新任务接受拍及旧数据污染。故障由H.hold_r/hold_b、fault_read/fault_write及就绪信号在模型边界施加，无随机。',
'H.reset(0)检查2C=0，no_access()提交8000/16且无读访问；DDR起来并读清后2C=3；分别撤PCIe、Aurora就绪重做无访问。每kind先prepare_tx(8000,length,550000,1)，hold_r阻塞直到fdma_rbusy且b_rvalid；中断200拍检查20=0，恢复并等500拍。若旧busy仍在则记录失败，提交9000/32的新seed660000+kind*4096/id2，观察2000拍并比较4个新拍；清理复位仅用于继续测试，不算恢复通过。再独立注入R SLVERR检查2C=7、读10后3；B SLVERR送111/222，等2500拍检查无块和丢2字；hold_r延迟400拍检查2C=19；hold_b阻塞后送2200字使真实FIFO溢出，检查2C溢出/超时位，释放后等50000拍排空并读清。最后A000/16新任务核对；failed_checks为0才可判通过。'),
'TC-F12':('n=0～4为块序号，w为块内字索引；value保存寄存器值，data保存64bit C2H返回；输入低32bit=1000+n*16+w，块字数n+1；wait_count(count)最多5000拍；无随机。',
'复位后逐块送1～5个字并等待可读数n+1，累计5块；每块先检查18=100000+n*40000、1C=(n+1)*4、34=n。读10并等100拍后重读这些元数据和30，确认未释放。按w+=2调用cpu_read64，分别比较有效低/高字，奇数尾只比较低字；所有有用字比完才写14=1并读回0，等待count=4-n，再核对下一队头。最后头地址/长度/序号为0，acks=5、dropped_words=0，共比60 Byte。'),
'TC-M09':('clk半周期2ns，PERIOD_CYCLES=25000000；period_tick观察U_DUT.tick；ticks/previous_tick记录事件数/前事件时间；writes为成功响应数、publications为已检查可读块数；hold_response将写响应延后；addr/data/strb保存实际写事务；seq/len是期望队头参数，v为输入32bit字。stored_words[0:5]已声明但未使用，不参与记分。无随机、无force计数或时间缩放。',
'复位8拍并enable；第1tick后等30拍，检查空块不发布且writes=0。送11111111、22222222，数据接受/响应后仍不可提前可读；第2tick后检查地址100000、长8、seq0，再ACK。延迟99999000ns，在下个周期边界前hold_response=1并送33333333/44444444，第3tick后仍无可读且状态等待响应；释放后检查140000/8/seq1再ACK。送55555555，第4tick时阻塞尾响应，观察180000的0F尾写且不可读；释放后检查180000/4/seq2再ACK。每相邻tick检查$time差100000000ns；最终4tick、3次成功写、3个发布、丢字0；看门狗410ms。')}

def decls(p):
    """Keep actual declarations including task arguments; comments supply original context."""
    s=p.read_text('utf-8-sig',errors='replace')
    out=[]
    for line in s.splitlines():
        t=line.strip()
        if re.match(r'(?:logic|reg|wire|int(?:eger)?|time|bit|byte|localparam|parameter)\b',t) and ';' in t:
            t=t.split(';',1)[0]+';'
            if len(t)<700:out.append(t)
        elif re.match(r'(?:task|function)\s+automatic\b',t):out.append(t)
    return list(dict.fromkeys(out))

case_catalog={}
for cid,paths in case_sources.items():
    catalog='\n'.join(p.name+'：\n'+'\n'.join(decls(p)) for p in paths)
    case_catalog[cid]=catalog
    r=bycase[cid]
    refs='、'.join(p[0] for p in points if cid in p[6])
    if cid in details:
        prep,steps=details[cid]
        r[3]=('测试台：'+paths[0].name+'\n'+(common_real+'\n' if cid in newnames and cid.startswith('TC-F') else '')+prep)
        r[4]='覆盖：'+refs+'\n'+'\n'.join(f'{i+1}. {x.strip()}。' for i,x in enumerate(steps.rstrip('。').split('。')) if x.strip())
    else:
        r[3]+='\n共同运行：tb_xdma_rtds_trythrough，CARD_ADDR_HI=0；fdma_clk=250MHz、user_clk=156.25MHz，RX_PERIOD_CYCLES=2500（10µs）。F01～F08是同一顺序测试的阶段。\n变量：frame_index为上行序号；bar_value为BAR读值；compare_mismatches为主机字节比较错误；config_sample_count/config_word_index检查配置提取数量与首拍位置；test_error_count汇总backend_error_count、scoreboard_error_count、assertion_error_count；tx_total_frames/tx_stall_cycles为输出统计；card_address/length_bytes/first_frame是start_tx_command任务的卡地址、字节数和模式帧号。\n随机：BFM与后端用$urandom_range(0,3)控制独立停顿，RX高半用$urandom检验有效低字提取；模式seed为确定性数据标识，源码未显式记录随机种子。'
        r[4]+='\n执行共性：start_tx_command先arm_tx(first_frame,length_bytes/16)，再写24/28/20；wait_cpu_tx_complete和wait_tx_complete各有200000拍边界。按握手递增记分，不在valid单独为1时计数。'
    # Exact declaration index is included in Excel, not hidden solely in an annex.
    r[3]+='\n声明索引（与上述含义对应；共享总线按通道阅读）：\n'+catalog
    assert len(r[3])<32767,(cid,len(r[3]))
bycase['TC-F08'][4]=bycase['TC-F08'][4].replace('cfg_sample_count','config_sample_count')
bycase['TC-F08'][5]=bycase['TC-F08'][5].replace('cfg_sample_count','config_sample_count').replace('stalls=1512','历史该次stalls=1512（随机结果，不作为固定通过阈值）')
for cid,name in newnames.items():
    bycase[cid][2]='模块生产参数' if cid=='TC-M09' else '行为模型集成'
    bycase[cid][7]=('失败：6项恢复检查｜' if cid=='TC-F11' else '通过｜')+'2026-10-10｜'+bycase[cid][2]
bycase['TC-F11'][5]='通过要求：任务前未就绪无访问；中断后命令结束且旧读事务排空；无全局复位即可发送新模式，4新拍、1完成事件，无旧数据。实际未满足：Aurora/PCIe旧busy残留且新任务前2拍旧数据；DDR恢复后新任务busy不结束、无数据。其余状态/清错检查不改变本例失败结论。'
bycase['TC-F09'][5]='10个合法编号逐个锁定；0/11不启动；每ID至少4个四拍帧；duty/Vin与LUT一致，Byte6=F0、Byte7～15=0、REM=3；不同后续ID不覆盖。'
bycase['TC-F10'][5]='7个非法请求均无DDR读访问、无发送拍、结束事件1次；合法16/32768/16 Byte分别1/2048/1帧、2/4096/2拍；每帧1个配置事件，帧读地址步长16，末端FFF0+16=10000。'
bycase['TC-F12'][5]='五块地址100000/140000/180000/1C0000/200000，长度4/8/12/16/20 Byte，序号0～4；有效载荷共60 Byte；读清不移动队头，每次ACK只释放1块，数量5→0且ACK恰5次。'
bycase['TC-M09'][5]='4次tick、3个相邻100ms间隔；空周期不发布；依次发布100000/8、140000/8、180000/4；写数据2222222211111111、4444444433333333、0000000055555555，STRB为FF/FF/0F；三次可读均晚于成功响应；丢字0。归档实测总仿真400.000394ms。'
for r in cases:
    if r[0]!='TC-R01':r[7]='归档'+r[7]+'；本机当前XPR未重跑'
    if r[0] in ('TC-M01','TC-F09'):r[7]+='；活动1G版本不符'

# Read local evidence logs, and preserve provenance rather than pretending to rerun.
execution=json.loads((ROOT/'evidence/execution_summary_20261010.json').read_text('utf-8-sig'))
evidence=[]
for e in execution:
    name=e['case'];run=e['run']
    paths=list((ROOT/'verification_20261010'/run).rglob(Path(e['log'].replace('\\','/')).name))
    assert paths,(name,e['log'])
    p=paths[0];txt=p.read_text('utf-8-sig',errors='replace')
    tags=[l for l in txt.splitlines() if re.search(r'TEST PASS|TEST FAIL|CHECK FAIL|RTDS_RESULT|RTDS_TRYTHROUGH_PASS|BOARD_CDC.*PASS',l)]
    assert tags,(name,p)
    evidence.append({**e,'local_log':str(p),'sha256':sha(p),'conclusions':tags})

# Structural/reference/content checks only; no appearance or rendering loop.
assert len(specs)<len(C['specs'])*0.4
assert len(port_to_group)==sum(len(m['ports']) for m in INV)
for rows in (specs,points,cases):assert len({r[0] for r in rows})==len(rows)
sids={r[0] for r in specs};cids=set(bycase)
for p in points:
    assert set(p[5].split('、'))<=sids,p
    assert p[6].startswith('待补') or set(p[6].split('、'))<=cids,p

wb=Workbook();wb.remove(wb.active)
headers=[['编号','功能类别','名称（接口含方向/位宽）','设计要求 / 含义','参数与限制'],['编号','功能类别','验证关注点','覆盖场景与分解','观察内容','对应规格编号','对应用例编号'],['编号','名称','验证层级','准备与输入（变量 / 含义 / 激励）','执行过程','预期结果 / 通过判据','优先级','执行结果']]
for title,head,rows,widths in zip(['设计规格','测试点','测试用例'],headers,[specs,points,cases],[[13,25,65,80,52],[13,25,34,75,75,34,34],[14,36,23,110,110,85,11,44]]):
    ws=wb.create_sheet(title);ws.append(head)
    for r in rows:ws.append(r)
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.zoomScale=85
    for j,w in enumerate(widths,1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(j)].width=w
    for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='425B76');cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
    for i in range(2,ws.max_row+1):
        row=ws[i]
        for c in row:
            c.font=Font(name='微软雅黑',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
            if i%2==0:c.fill=PatternFill('solid',fgColor='F0F4F8')
        ws.row_dimensions[i].height=min(409,max(42,15*max(sum(max(1,math.ceil(len(line)*1.5/widths[j])) for line in str(c.value or '').splitlines()) for j,c in enumerate(row))))
        if title=='测试用例':
            color='FCE4D6' if '失败' in str(row[-1].value) else 'E2F0D9' if str(row[-1].value).startswith('通过') else 'FFF2CC'
            row[-1].fill=PatternFill('solid',fgColor=color)
            row[3].comment=Comment('长用例的完整分段阅读版见同目录Word第2部分；本单元格保留完整内容，可在编辑栏展开。','Codex')
    ws.row_dimensions[1].height=34
xlsx=OUT/'RTDS_功能规格_测试点_详细用例_20261011.xlsx';wb.save(xlsx)

doc=Document();doc.styles['Normal'].font.name='Microsoft YaHei';doc.styles['Normal'].font.size=Pt(10)
doc.add_heading('RTDS 规格与测试文档核查报告',0);doc.add_paragraph('功能重组版 · 2026-10-11 · 本机F盘基线')
doc.add_heading('1. 核查结论',1)
doc.add_paragraph(f'以 {ROOT} 的本机源码、测试台和归档日志为依据。设计规格从{len(C["specs"])}行减少到{len(specs)}行，减少{(1-len(specs)/len(C["specs"]))*100:.1f}%；其中功能/寄存器规格{len(specs)-len(groups)}行，接口组{len(groups)}行，保留本机活动范围{len(INV)}个模块的{len(port_to_group)}个端口名、方向和当前参数位宽。测试点{len(points)}项，用例{len(cases)}项。')
doc.add_paragraph('规格先按业务功能展开，接口仅作为末尾压缩清单。测试点改为场景分解与观察内容，不设置通过判据；通过判据留在用例。长用例在Excel保留完整变量及执行说明，本报告另提供连续段落阅读版。未对文档进行外观核查或PDF/PNG渲染。')
doc.add_paragraph('旧文件保持不变；本次不修改业务RTL、不开发测试台、不运行Vivado。2026-10-10结果来自拷贝到本机的已有执行证据，不是本次重新仿真。此前“新增五类均未执行”的状态已经过时：F09/F10/F12/M09通过，F11失败。')
doc.add_paragraph('本机存在版本混杂：XPR启用的1G RTL为PERIOD_CYCLES=63、两拍6 Byte、REM=3/1；TOP没有同源事件或GPIO连接。四拍16 Byte、Byte6=F0及同源2µs/50µs保留为目标规格，并在Excel限制栏明确尚未落实到本机活动工程。timebase文件虽存在，当前XPR未列入；其模块测试属于单独测试台证据，不能证明TOP集成。所有通过/失败均标记归档结果，本机当前XPR未重跑。')
doc.add_paragraph('文件内容比对：'+str(sum(x['same_as_1010'] for x in source_audit))+'/'+str(len(source_audit))+'个模块所在文件与旧快照SHA256一致（同文件多模块分别列项）。端口差异：'+json.dumps(port_changes,ensure_ascii=False)+'。旧828端口不直接视为本机现状；当前表按本机声明重建，未启用timebase单列差异。')
doc.add_paragraph('编号迁移：S-01～S-16、S-19～S-46继续保留；S-17/S-18维持旧版退役，不重新占用。S-875～S-896细分原功能规则。828条接口规格合并时选组内最早旧编号作为新编号，其余编号并入该组，详见附表。P-01～P-25、P-27～P-36保留并重写，P-26仍退役；新增P-37～P-46补充独立覆盖观察和可见缺口。22个TC编号保持不变。')
doc.add_heading('2. 验证记录',1)
doc.add_paragraph('验证层级：M类主要为真实单模块加行为激励/响应；F01～F08为模型BAR/DDR/Root Port API加adapter，不能证明真实INFO_EXCH、XDMA或GT；F09～F12使用真实INFO_EXCH/adapter/CONV_DMA/DDR后端及行为RAM，CPU预装和C2H主机仍为模型；M09为真实块写模块加FDMA响应模型，生产时钟与周期参数；R01仅保留历史真实Root Port证据。')
for e in evidence:
    cid=e.get('test_case') or next((cid for cid,n in tbmap.items() if n==e['case']), 'TC-F01～TC-F08' if e['case']=='fast_base' else '设计检查：board_cdc')
    doc.add_heading(f'{cid} / {e["case"]}',2)
    doc.add_paragraph(f'日期：{e["started"]}；结论：'+('通过' if e['passed'] else '失败')+'；运行归档：'+e['run'])
    doc.add_paragraph('本机日志：'+e['local_log']+'\nSHA256：'+e['sha256'])
    doc.add_paragraph('\n'.join(e['conclusions']))
doc.add_heading('用例逐项阅读版：变量、输入与执行',2)
doc.add_paragraph('AXI声明读法：AW=写地址/属性，W=写数据/字节使能/末拍，B=写响应，AR=读地址/属性，R=读数据/响应/末拍；ADDR为字节地址，LEN为拍数减一，SIZE为每拍字节数的log2，VALID/READY为该通道握手，ID为事务标识。FDMA的addr为字节地址、size为64bit拍数、req为请求、busy为在途、done/error为响应完成；CONV_DMA的wready输入实际上是源有效，wvalid输出是接受脉冲。clk/rst为时钟/复位，head/ready/drop为队头/队列/丢字状态。以下精确声明索引保留变量原名，辅助任务参数在声明中列出；未使用的声明不计为覆盖。')
for r in cases:
    doc.add_heading(r[0]+' '+r[1],3)
    for label,text in [('验证层级',r[2]),('准备与输入',r[3].split('\n声明索引')[0]),('执行过程',r[4]),('预期结果 / 通过判据',r[5]),('执行结果',r[7])]:
        doc.add_paragraph(label+'：'+text)
    if r[0] in case_sources:doc.add_paragraph('源码：'+'；'.join(str(x) for x in case_sources[r[0]]))
doc.add_heading('3. 设计检查',1)
doc.add_paragraph(f'接口检查：本机文件重新解析方向及当前实例参数后的位宽，当前活动范围{len(INV)}模块/{len(port_to_group)}端口。AXI、FDMA、DDR4接口成组展示，inout仍保留。厂商IP仅列启用封装层端口，不展开内部实现。旧版特有的tick_toggle_2us和TOP同源GPIO连接不伪装成本机已实现；差异见第一部分。当前接口组不代替板级连接、约束和实现后时序签核。')
doc.add_paragraph('CDC/时钟复位：已有CFG/COUNT/RESET同步封装；业务250MHz、10G用户156.25MHz、1G用户31.25MHz，异步FIFO承担数据跨域。相同频率不能推导同相。2026-10-10当前配置board_cdc冒烟通过，仅覆盖测试台选定通路；完整TOP CDC尚未通过，不能用冒烟替代结构CDC与约束收敛。本次未重跑CDC工具。')
doc.add_paragraph('约束与复位时序沿用已有审计边界：1G初始化时钟由100MHz四分频；GT/协议核按启动计数先后释放。物理时钟、IO约束及真实MIG UI关系应由实现报告确认。本次不把模型时钟行为当成真实时序收敛证据。')
doc.add_paragraph('语义核查重点：接口接受≠DDR写响应成功≠块可读≠CPU读完确认；0x14为W1P，0x10读清不释放块；0x20清忙包含拒绝/中止；0x2C为ready、calib、响应错误、RX溢出、超时位。目标版本1G四拍16Byte、Byte6=F0，2µs为toggle翻转间隔、50µs为GPIO完整周期；这些目标与本机旧活动1G版本的差异已单列。100ms周期事件与块可读时刻分别判定。')
doc.add_heading('4. 问题与待补项',1)
doc.add_paragraph('版本优先问题：需后续统一本机活动RTL、XPR和四拍测试台基线，再重新编译和运行。当前新测试台传入tick_toggle_2us，而活动1G模块没有该端口；不能直接将归档PASS迁移到当前根目录。此次仅生成文档，未擅自将归档源码覆盖当前工程。')
doc.add_paragraph('高优先级：TC-F11存在6项失败，覆盖三种中断路径。Aurora/PCIe中断后旧CONV_DMA读busy残留，后续32Byte新任务前两拍旧数据；DDR中断后新任务busy不结束且无数据。测试台失败后全局复位属于清理，不能记为恢复通过。需后续修复取消事务、旧返回排空及DDR中断状态处理，然后无复位重跑。')
doc.add_paragraph('非零CARD_ADDR_HI历史扩展失败仍保留：当前PCIe.v将64bit地址低32bit接入DDR后端，生产配置CARD_ADDR_HI=0。扩展失败不作为本次新增必补用例，也不把本次低地址结果解释为支持>4GiB。')
doc.add_paragraph('真实Root Port：TC-R01为历史16Byte事务记录，曾出现功能PASS但512Byte脚本门禁FAIL；当前版本证据不足。GT、真实DDR物理校准与板卡DMA尚无本轮验证。完整TOP CDC尚未通过。随机测试源码没有明确固化种子，归档stall计数只作为该次观测，不作为所有重跑的固定期望。')
for p in points:
    if p[6].startswith('待补'):doc.add_paragraph(p[0]+' '+p[2]+'：'+p[3]+' '+p[4])
doc.add_heading('附表：接口编号合并',1)
t=doc.add_table(rows=1, cols=3);t.style='Light Shading Accent 1'
for c,s in zip(t.rows[0].cells,['新编号','模块 / 接口组','并入旧编号']):c.text=s
for g in groups:
    cs=t.add_row().cells
    for c,s in zip(cs,[g['id'],g['module']+' / '+g['group'],'、'.join(g['old_ids'])]):c.text=s
doc.add_heading('附表：功能细分来源',1)
for s,parent in split_from.items():doc.add_paragraph(s+' ← '+parent+'：'+next(r[2] for r in specs if r[0]==s))
doc.add_heading('附录：源码、日志、文档与编号对应',1)
for a in source_audit:
    ids='、'.join(g['id'] for g in groups if g['module']==a['module'])
    doc.add_paragraph(ids+' / '+a['module']+'\n'+a['path']+'\nSHA256 '+a['sha256']+'；与10/10快照'+('一致' if a['same_as_1010'] else '存在差异'))
seen=set()
for cid,paths in case_sources.items():
    for p in paths:
        if p in seen:continue
        seen.add(p);related='、'.join(k for k,v in case_sources.items() if p in v)
        doc.add_paragraph(related+'\n'+str(p)+'\nSHA256 '+sha(p))
for p in [OLD/'content.json',OLD/'port_inventory.json',ROOT/'evidence/execution_summary_20261010.json',ROOT/'evidence/local_wave_sessions/new_cases_report_20261010.md']:
    doc.add_paragraph('规格/追溯及结果来源：'+str(p))
for p in Path('E:/codex_history_md/PCIE-CPU').glob('2026-10-10*.md'):
    doc.add_paragraph('历史背景（不替代当前源码和日志）：'+str(p))
doc.add_heading('附录：测试台声明索引',1)
seen=set()
for paths in case_sources.values():
    for p in paths:
        if p in seen:continue
        seen.add(p);doc.add_heading(p.name,2)
        for line in decls(p):doc.add_paragraph(line)
docx=OUT/'RTDS_功能规格与验证核查报告_20261011.docx';doc.save(docx)
with zipfile.ZipFile(docx) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):ET.fromstring(z.read(name))
payload={'specs':specs,'points':points,'cases':cases,'port_to_group':port_to_group,'source_audit':source_audit,'port_changes':port_changes,'evidence':evidence,'groups':groups,'split_from':split_from}
(OUT/'content_and_evidence.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps({'xlsx':str(xlsx),'docx':str(docx),'specs':len(specs),'functional':len(specs)-len(groups),'interface_groups':len(groups),'reduction_percent':round((1-len(specs)/len(C['specs']))*100,1),'points':len(points),'cases':len(cases),'ports':len(port_to_group),'source_same':sum(a['same_as_1010'] for a in source_audit)},ensure_ascii=False))
